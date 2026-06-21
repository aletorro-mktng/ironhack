"""NiceGUI user interface for Tell Tales Ink."""

from __future__ import annotations

import asyncio
import base64
import json
import mimetypes
import os
import re
import sys
import zipfile
from datetime import datetime, timedelta
from html import escape
from pathlib import Path

from nicegui import app, ui
from pydub import AudioSegment

SRC_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))
os.chdir(PROJECT_ROOT)

import chapter_reader
from content_pipeline import create_generation_prompt, run_pipeline, save_output
from context_filter import select_relevant_context
from draft_store import (
    archive_saved_draft,
    delete_saved_draft,
    get_saved_draft,
    list_saved_drafts,
    read_saved_draft_content,
    render_saved_drafts_markdown,
    save_draft,
    update_saved_draft,
)
from elevenlabs_integration import list_voices, strip_production_tags, synthesize_speech
from llm_integration import DEFAULT_MODEL, generate_text, generate_text_with_image
from image_generation import generate_external_visual
from prompt_templates import list_supported_content_types
from character_assets import character_asset_note, resolve_character_portrait_asset
from quote_graphics import (
    QUOTE_GRAPHIC_FORMATS,
    QUOTE_GRAPHIC_THEMES,
    quote_graphic_format_options,
    render_quote_cards,
)
from selection_options import (
    AUDIENCE_OPTIONS,
    BLOG_LENGTH_OPTIONS,
    CHARACTER_TAGS,
    CONSTRAINT_OPTIONS,
    ELEVENLABS_MODEL_OPTIONS,
    INSTAGRAM_FORMAT_OPTIONS,
    INSTAGRAM_HASHTAG_OPTIONS,
    LINKEDIN_CONTENT_FORMAT_OPTIONS,
    NEWSLETTER_OBJECTIVE_OPTIONS,
    NEWSLETTER_STRUCTURE_OPTIONS,
    PLATFORM_OPTIONS,
    PODCAST_DESTINATION_OPTIONS,
    PODCAST_FORMAT_OPTIONS,
    PODCAST_LENGTH_OPTIONS,
    PODCAST_TONE_OPTIONS,
    PRESS_RELEASE_TIMING_OPTIONS,
    QUOTE_BOOK_OPTIONS,
    QUOTE_MOOD_TAGS,
    SOCIAL_OBJECTIVES,
    YOUTUBE_CONTENT_FORMAT_OPTIONS,
)


APP_TITLE = "Tell Tales Ink"
DASHBOARD_ASSET_DIR = PROJECT_ROOT / "assets" / "dashboard"
PRESS_PROFILE_PATH = PROJECT_ROOT / "outputs" / "press_profiles.json"
PRESS_RELEASE_DESTINATION = "PR Distribution Services"

# The platform/destination each content type should default to in the Generator.
PLATFORM_BY_CONTENT_TYPE = {
    "instagram_caption": "Instagram",
    "linkedin_content": "LinkedIn",
    "youtube_content": "YouTube",
    "newsletter_blurb": "newsletter",
    "blog_post": "blog",
    "quote_post": "Instagram",
    "review_pull_quote": "Instagram",
    "character_spotlight": "Instagram",
    "press_release": PRESS_RELEASE_DESTINATION,
}
# Content types whose platform is fully implied — hide the dropdown to reduce clutter.
PLATFORM_IMPLIED_CONTENT_TYPES = {"linkedin_content", "youtube_content", "newsletter_blurb", "press_release"}
# Content types where Social objectives / Audience are irrelevant to the deliverable.
HIDE_SOCIAL_AUDIENCE_CONTENT_TYPES = {"press_release", "review_pull_quote"}
# Full set of asset types the generator supports (generation is generic, so any of
# these still works if passed in). The Press Release dropdown offers only the
# companion subset below; the others remain available to other workflows.
PR_ASSET_OPTIONS = [
    "Media kit",
    "Pitch email",
    "Author bio",
    "Book description",
    "Boilerplate",
    "Fact sheet",
    "Cover image brief",
    "Press kit checklist",
]

# Companion assets offered alongside a press release (focused subset).
PRESS_RELEASE_COMPANION_ASSETS = [
    "Media pitch email",
    "Boilerplate",
    "Fact sheet",
    "Cover image brief",
    "Social announcement posts",
]

# All publishing deliverables Tell Tales Ink can produce (display labels).
PUBLISHING_ASSET_TYPES = [
    "Press release",
    "Media kit",
    "Media pitch email",
    "Author bio",
    "Book description",
    "Boilerplate",
    "Fact sheet",
    "Cover image brief",
    "Social announcement posts",
    "Press kit checklist",
]

# Slug -> human label for publishing content types.
PUBLISHING_CONTENT_LABELS = {
    "press_release": "Press release",
    "media_kit": "Media kit",
    "media_pitch_email": "Media pitch email",
    "author_bio": "Author bio",
    "book_description": "Book description",
    "boilerplate": "Boilerplate",
    "fact_sheet": "Fact sheet",
    "cover_image_brief": "Cover image brief",
    "social_announcement_posts": "Social announcement posts",
    "press_kit_checklist": "Press kit checklist",
}


def _format_eta(seconds: float) -> str:
    if seconds <= 0 or seconds == float("inf"):
        return "calculating"
    seconds = int(round(seconds))
    minutes, remaining = divmod(seconds, 60)
    if minutes:
        return f"{minutes}m {remaining:02d}s"
    return f"{remaining}s"


async def set_generation_progress(
    *,
    progress,
    label,
    status=None,
    started_at: datetime,
    completed_steps: int,
    total_steps: int,
    completed_assets: int,
    total_assets: int,
    message: str,
) -> None:
    total_steps = max(1, int(total_steps or 1))
    total_assets = max(1, int(total_assets or 1))
    completed_steps = max(0, min(int(completed_steps), total_steps))
    completed_assets = max(0, min(int(completed_assets), total_assets))
    ratio = completed_steps / total_steps
    elapsed = max(0.1, (datetime.now() - started_at).total_seconds())
    eta = _format_eta((elapsed / ratio) - elapsed if ratio > 0 else 0)
    progress.visible = True
    progress.value = ratio
    pct = int(round(ratio * 100))
    detail = f"{pct}% complete · ETA {eta} · Assets {completed_assets}/{total_assets} · {message}"
    if hasattr(label, "set_text"):
        label.set_text(detail)
    else:
        label.value = detail
    if status is not None:
        status.value = message
    try:
        ui.run_javascript(f'document.title = "[{pct}%] {APP_TITLE}";')
    except Exception:
        pass
    await asyncio.sleep(0)


async def finish_generation_progress(
    *,
    progress,
    label,
    status=None,
    total_assets: int = 1,
    message: str = "Complete",
) -> None:
    progress.visible = True
    progress.value = 1
    detail = f"100% complete · ETA 0s · Assets {max(1, total_assets)}/{max(1, total_assets)} · {message}"
    if hasattr(label, "set_text"):
        label.set_text(detail)
    else:
        label.value = detail
    if status is not None:
        status.value = message
    try:
        ui.run_javascript(f'document.title = "{APP_TITLE}";')
    except Exception:
        pass
    await asyncio.sleep(0)


CHATGPT_COMPARISON_MODEL_OPTIONS = list(dict.fromkeys([
    os.getenv("CHATGPT_COMPARISON_MODEL", DEFAULT_MODEL),
    DEFAULT_MODEL,
    "gpt-5.4-mini",
    "gpt-5.4",
    "gpt-5.4-instant",
]))

CAMPAIGN_FORMATS = [
    ("podcast", "Podcast"),
    ("instagram_caption", "Instagram"),
    ("youtube_content", "YouTube"),
    ("linkedin_content", "LinkedIn"),
    ("blog_post", "Blog Post"),
    ("newsletter_blurb", "Newsletter"),
    ("quote_post", "Quote Post"),
    ("review_pull_quote", "Review Pull Quote"),
    ("character_spotlight", "Character Spotlight"),
    ("press_release", "Press Release"),
]

CAMPAIGN_FORMAT_LABELS = {key: label for key, label in CAMPAIGN_FORMATS}
# Human-readable labels for the Generator content type dropdown (UX-02).
CONTENT_TYPE_LABELS = {
    "instagram_caption": "Instagram Caption",
    "youtube_content": "YouTube",
    "linkedin_content": "LinkedIn",
    "blog_post": "Blog Post",
    "newsletter_blurb": "Newsletter Blurb",
    "character_spotlight": "Character Spotlight",
    "review_pull_quote": "Review / Pull Quote",
    "quote_post": "Quote Post",
    "press_release": "Press Release",
    "podcast": "Podcast",
    "media_kit": "Media Kit",
    "media_pitch_email": "Media Pitch Email",
    "author_bio": "Author Bio",
    "book_description": "Book Description",
    "boilerplate": "Boilerplate",
    "fact_sheet": "Fact Sheet",
    "cover_image_brief": "Cover Image Brief",
    "social_announcement_posts": "Social Announcement Posts",
    "press_kit_checklist": "Press Kit Checklist",
}


def content_type_label(content_type: str) -> str:
    return CONTENT_TYPE_LABELS.get(content_type, content_type.replace("_", " ").title())


def set_field_status(label, state: str, message: str) -> None:
    """Set a status label's text and a state class (loading/success/error/idle)."""
    prefix = {"loading": "⏳ ", "success": "✓ ", "error": "⚠ "}.get(state, "")
    label.set_text(f"{prefix}{message}")
    label.classes(remove="mce-status-loading mce-status-success mce-status-error")
    if state in {"loading", "success", "error"}:
        label.classes(add=f"mce-status-{state}")
CAMPAIGN_QUANTITY_OPTIONS = ["1", "2", "3", "4", "5"]
CAMPAIGN_DURATION_OPTIONS = ["1 week", "2 weeks", "1 month", "3 months"]
_CAMPAIGN_DURATION_DAYS = {"1 week": 7, "2 weeks": 14, "1 month": 30, "3 months": 90}
_CAMPAIGN_CADENCE_POSTS_PER_DAY = {
    "daily": 1.0,
    "every 2 days": 0.5,
    "twice a week": 2 / 7,
    "weekly": 1 / 7,
}


def campaign_post_quantity(cadence, duration) -> int:
    """Posts to generate per content type for a posting ``cadence`` across a campaign
    ``duration`` (cadence x duration). E.g. every 2 days over 1 month (30 days) -> 15.
    Returns 1 when either input is unrecognized."""
    rate = _CAMPAIGN_CADENCE_POSTS_PER_DAY.get(str(cadence or "").strip().lower())
    days = _CAMPAIGN_DURATION_DAYS.get(str(duration or "").strip())
    if not rate or not days:
        return 1
    return max(1, round(days * rate))
CAMPAIGN_STYLE_OPTIONS = {
    "podcast": ["cinematic monologue", "host + guest interview", "lore deep dive", "news-style segment"],
    "instagram_caption": ["punchy caption", "reel script", "carousel copy", "story sequence"],
    "youtube_content": ["SEO title set", "description", "community post", "talking points"],
    "linkedin_content": ["authority post", "article", "reflection post", "announcement"],
    "blog_post": ["editorial essay", "listicle", "reader guide", "how-to guide", "SEO explainer"],
    "newsletter_blurb": ["announcement", "personal note", "digest", "behind-the-scenes"],
    "quote_post": ["character quote", "theme quote", "spoiler-free quote", "dark academia quote"],
    "review_pull_quote": ["critic pull quote", "reader praise", "short testimonial", "media kit blurb"],
    "character_spotlight": ["profile", "arc analysis", "spoiler-safe intro", "fan-facing deep dive"],
    "press_release": ["FOR IMMEDIATE RELEASE", "media advisory", "award announcement", "launch announcement"],
}

BLOG_FORMAT_OPTIONS = [
    "character deep-dive",
    "listicle",
    "interview",
    "best of",
    "news",
    "photo/gallery slideshow",
    "comparison",
    "op-ed",
    "curated content",
    "critics review roundup",
    "polls / survey / quiz",
]
BLOG_STRUCTURE_CHECKLIST_OPTIONS = [
    "strong hook",
    "canon context",
    "spoiler warning",
    "character stakes",
    "reader takeaway",
    "image/visual break",
    "quote pullout",
    "CTA",
    "SEO FAQ",
]
CHARACTER_SPOTLIGHT_PLATFORM_FORMATS = [
    "Instagram post",
    "Instagram reel",
    "Instagram story",
    "Instagram carousel",
    "YouTube cover",
    "YouTube post",
    "LinkedIn post square",
    "LinkedIn post horizontal",
    "Blog cover",
    "Blog square",
    "Blog horizontal",
]
CHARACTER_IMAGE_MODE_OPTIONS = [
    "use uploaded image",
    "use character portrait",
    "text-only highlight image",
]
QUOTE_IMAGE_STYLE_OPTIONS = list(QUOTE_GRAPHIC_THEMES.keys())

# Distinct content angles cycled across the multiple posts generated for one content
# type, so a set of e.g. 15 Instagram posts doesn't feel repetitive. Each carries a
# content focus and a matching visual treatment (used to vary images per post too).
CAMPAIGN_POST_ANGLES = [
    {"name": "Teaser / intrigue", "content": "Open a curiosity loop and tease the premise spoiler-free; make them need to know more.", "visual": "moody and minimal, lots of negative space, single mysterious focal point"},
    {"name": "Character spotlight", "content": "Center one character — their voice, wound, or impossible choice.", "visual": "intimate portrait-forward composition, dramatic single-subject framing"},
    {"name": "Quote pull", "content": "Build around one striking line from the world; let the language carry it.", "visual": "bold typographic quote-card layout, strong type hierarchy"},
    {"name": "Behind the story", "content": "Share the making-of: craft, inspiration, world-building, or author intent.", "visual": "textured archival / notebook aesthetic, warm and human"},
    {"name": "Reader social proof", "content": "Frame around the reading experience and reactions (never invent fake reviewer quotes).", "visual": "clean testimonial / review-card layout, credible and editorial"},
    {"name": "Book aesthetic / mood", "content": "Lead with atmosphere and tone — the feeling of the world more than the plot.", "visual": "atmospheric cinematic color, strong mood lighting, rich texture"},
    {"name": "Release urgency / CTA", "content": "Drive action — availability, timing, and where to get it.", "visual": "high-contrast announcement energy, bold focal point and clear hierarchy"},
    {"name": "Theme exploration", "content": "Explore a core theme (revenge, grief, justice, identity) and why it resonates.", "visual": "symbolic conceptual imagery tied to the theme"},
    {"name": "World / setting", "content": "Spotlight the setting and its rules, beauty, or dangers.", "visual": "wide environmental establishing shot, strong sense of place"},
    {"name": "Conflict / stakes", "content": "Surface the central conflict and what is at risk if it goes wrong.", "visual": "tense composition, dramatic shadow and contrast"},
]


def campaign_post_angle(asset_number: int) -> dict:
    """Pick the content/visual angle for a given post in a multi-post set (cycles)."""
    return CAMPAIGN_POST_ANGLES[(max(1, int(asset_number)) - 1) % len(CAMPAIGN_POST_ANGLES)]


def campaign_set_angle_names(count: int) -> list:
    """Ordered, de-duplicated angle names spanning a set of ``count`` posts."""
    names = [campaign_post_angle(i)["name"] for i in range(1, max(1, int(count)) + 1)]
    return list(dict.fromkeys(names))


def campaign_visual_theme_for_post(base_style: str, asset_number: int) -> str:
    """Rotate the visual theme per post (starting from the user's chosen style) so
    images in a set vary in palette/composition instead of all looking the same."""
    themes = QUOTE_IMAGE_STYLE_OPTIONS
    if not themes:
        return base_style or "Gothic"
    base = base_style if base_style in themes else "Gothic"
    start = themes.index(base) if base in themes else 0
    return themes[(start + max(1, int(asset_number)) - 1) % len(themes)]


# Platforms offered by the Chapter Promos tab -> the generation content type used.
CHAPTER_PROMO_PLATFORMS = {
    "Instagram": "instagram_caption",
    "Blog": "blog_post",
    "LinkedIn": "linkedin_content",
    "YouTube": "youtube_content",
}

# Performance cues / emotional tags inserted into a podcast script. These survive the
# pre-TTS strip (see NATURAL_DELIVERY_TAGS) so ElevenLabs (v3) can act on them.
PODCAST_PERFORMANCE_CUES = [
    "[laughs]", "[chuckles]", "[sighs]", "[exhales]", "[gasps]", "[groans]", "[scoffs]",
    "[whispers]", "[softly]", "[quietly]", "[lowers voice]", "[under breath]",
    "[excited]", "[enthusiastic]", "[nervous]", "[anxious]", "[somber]", "[serious]",
    "[sarcastic]", "[dryly]", "[mocking]", "[playful]", "[warmly]", "[thoughtful]",
    "[hesitant]", "[cautious]", "[tired]", "[clears throat]", "[deep breath]",
    "[emphatic]", "[trailing off]", "[beat]", "[short pause]", "[long pause]",
]

# Stock transitions, sound effects and music cues. These are producer notes — stripped
# before TTS (they are not spoken) but kept in the script for the human editor.
PODCAST_SOUND_EFFECTS = [
    "[TRANSITION: hard cut]", "[TRANSITION: crossfade]", "[TRANSITION: whoosh]",
    "[TRANSITION: riser]", "[TRANSITION: stinger]", "[TRANSITION: glitch]",
    "[TRANSITION: time jump]", "[TRANSITION: scene change]",
    "[SFX: heartbeat]", "[SFX: footsteps]", "[SFX: door creak]", "[SFX: phone buzz]",
    "[SFX: camera shutter]", "[SFX: thunder]", "[SFX: rain]", "[SFX: wind howl]",
    "[SFX: clock ticking]", "[SFX: paper rustle]", "[SFX: glass shatter]",
    "[SFX: distant scream]", "[SFX: vinyl crackle]", "[SFX: static burst]", "[SFX: tense drone]",
    "[INTRO MUSIC: low cinematic strings, fade under host]",
    "[OUTRO MUSIC: warm resolving theme, 12s]",
    "[MUSIC BED: tense underscore]", "[STINGER: dramatic hit]",
]
INSTAGRAM_VISUAL_FORMAT_OPTIONS = [
    "Instagram Post (4:5)",
    "Instagram Reel (9:16)",
    "Instagram Story (9:16)",
    "Square Post (1:1)",
]
BLOG_VISUAL_FORMAT_OPTIONS = [
    "Blog Cover (16:9)",
    "Blog Square (1:1)",
    "Blog Horizontal (1.91:1)",
    "LinkedIn Article Horizontal (1.91:1)",
]
BLOG_GALLERY_VISUAL_FORMAT_OPTIONS = [
    "Blog Cover (16:9)",
    "Blog Square (1:1)",
    "Blog Horizontal (1.91:1)",
    "Instagram Post (4:5)",
    "Square Post (1:1)",
]

PODCAST_PREVIEW_MAX_SEGMENTS = 3
PODCAST_TTS_MAX_CHARS = 2400
PODCAST_OUTPUT_FORMAT = "mp3_44100_128"
IMAGE_UPLOAD_DIR = PROJECT_ROOT / "outputs" / "uploads"
ALLOWED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def list_existing_images(limit: int = 60) -> dict[str, str]:
    """Return {path: label} of recent uploaded/generated/brand images for the asset picker (BUG-IMG-01)."""
    roots = [
        PROJECT_ROOT / "outputs" / "uploads",
        PROJECT_ROOT / "outputs" / "generated_visuals",
        PROJECT_ROOT / "assets",
    ]
    files = []
    for root in roots:
        if root.exists():
            for path in root.rglob("*"):
                if path.is_file() and path.suffix.lower() in ALLOWED_IMAGE_EXTENSIONS:
                    files.append(path)
    try:
        files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    except OSError:
        pass
    options: dict[str, str] = {}
    used_labels: set[str] = set()
    for path in files[:limit]:
        label = path.name
        if label in used_labels:
            label = f"{path.parent.name}/{path.name}"
        used_labels.add(label)
        options[str(path)] = label
    return options
IMAGE_AWARE_CONTENT_TYPES = {"instagram_caption", "youtube_content", "linkedin_content", "character_spotlight", "blog_post"}


def normalize_selected(value) -> str:
    if value is None:
        return "Not specified"
    if isinstance(value, str):
        cleaned = value.strip()
        return cleaned or "Not specified"
    if isinstance(value, (list, tuple, set)):
        items = [str(item).strip() for item in value if str(item).strip()]
        if not items:
            return "Not specified"
        return ", ".join(dict.fromkeys(items))
    cleaned = str(value).strip()
    return cleaned or "Not specified"


def strip_markdown(value) -> str:
    """Return plain text with common Markdown emphasis/markup removed."""
    text = normalize_selected(value)
    if text == "Not specified":
        return text
    text = re.sub(r"`{1,3}([^`]*)`{1,3}", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", r"\1", text)
    text = re.sub(r"(?<!_)_(?!_)(.+?)_(?!_)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s{0,3}>\s?", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s{0,3}[-*+]\s+", "• ", text, flags=re.MULTILINE)
    return text.strip()


def markdown_to_html(value) -> str:
    """Escape text, then render a safe subset of inline Markdown to HTML.

    Newlines are left intact for containers that use ``white-space:pre-wrap``.
    """
    text = normalize_selected(value)
    out = escape(text)
    out = re.sub(r"`([^`]+)`", r"<code>\1</code>", out)
    out = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)
    out = re.sub(r"__(.+?)__", r"<strong>\1</strong>", out)
    out = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", r"<em>\1</em>", out)
    out = re.sub(r"(?<!_)_(?!_)(.+?)_(?!_)", r"<em>\1</em>", out)
    out = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2">\1</a>', out)
    # Thematic break: a standalone line of --- / *** / ___ becomes a divider (not literal text).
    out = re.sub(
        r"^\s{0,3}(?:-{3,}|\*{3,}|_{3,})\s*$",
        '<hr style="border:none;border-top:1px solid rgba(0,0,0,.14);margin:10px 0;">',
        out,
        flags=re.MULTILINE,
    )
    out = re.sub(r"^\s{0,3}#{1,6}\s*(.+)$", r"<strong>\1</strong>", out, flags=re.MULTILINE)
    out = re.sub(r"^\s{0,3}[-*+]\s+(.+)$", r"• \1", out, flags=re.MULTILINE)
    return out


def load_press_profiles() -> dict:
    if not PRESS_PROFILE_PATH.exists():
        return {}
    try:
        return json.loads(PRESS_PROFILE_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_press_profiles(profiles: dict) -> None:
    PRESS_PROFILE_PATH.parent.mkdir(parents=True, exist_ok=True)
    PRESS_PROFILE_PATH.write_text(json.dumps(profiles, indent=2, sort_keys=True), encoding="utf-8")


def press_profile_options() -> list[str]:
    return sorted(load_press_profiles().keys())


def first_nonempty_line(text: str) -> str:
    for line in str(text or "").splitlines():
        cleaned = line.strip().lstrip("#").strip()
        if cleaned:
            return cleaned
    return ""


def brief_line(label: str, value) -> str:
    return f"{label}: {normalize_selected(value)}"


def safe_upload_filename(name: str) -> str:
    stem = Path(name or "uploaded_image").stem
    suffix = Path(name or "").suffix.lower()
    if suffix not in ALLOWED_IMAGE_EXTENSIONS:
        suffix = ".png"
    stem = re.sub(r"[^\w\s-]", "", stem, flags=re.UNICODE)
    stem = re.sub(r"[-\s]+", "_", stem).strip("_").lower() or "uploaded_image"
    return f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{stem[:70]}{suffix}"


def generate_press_release_assets(
    *,
    selected_assets,
    topic: str,
    draft: str,
    contact: dict[str, str],
    news_angle: str,
) -> dict[str, object]:
    assets = [asset for asset in (selected_assets or []) if str(asset).strip()]
    if not assets:
        return {}
    output_dir = PROJECT_ROOT / "outputs" / "press_release_assets" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: list[str] = []
    errors: list[str] = []
    rendered: list[dict[str, str]] = []
    for asset in assets:
        prompt = f"""Create this press-release support asset for Tell Tales Ink.

Asset: {asset}
Topic: {topic}
News angle: {news_angle}
Media contact: {json.dumps(contact, ensure_ascii=False)}

Press release draft:
{draft}

Return a polished, usable asset. If the selected asset is a checklist, use checkboxes. If it is a pitch email, include subject line and body. If it is a media kit, include sections and reusable copy."""
        try:
            content = generate_text(prompt)
            path = output_dir / f"{slugify_filename(asset, 'press_asset')}.md"
            path.write_text(content, encoding="utf-8")
            paths.append(str(path))
            rendered.append({"asset": str(asset), "content": content, "path": str(path)})
        except Exception as exc:
            errors.append(f"{asset}: {exc}")
    zip_path = output_dir / "press_release_assets.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in paths:
            file_path = Path(path)
            if file_path.exists():
                archive.write(file_path, arcname=file_path.name)
    result: dict[str, object] = {"paths": paths, "zip_path": str(zip_path) if paths else "", "assets": rendered}
    if errors:
        result["error"] = "; ".join(errors)
    return result


def image_data_url(image_path: str | Path | None) -> str:
    if not image_path:
        return ""
    path = Path(str(image_path))
    if not path.exists():
        return ""
    mime_type = mimetypes.guess_type(path.name)[0] or "image/png"
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def social_image_instructions(content_type: str) -> str:
    if content_type == "instagram_caption":
        return """
IMAGE-AWARE OUTPUT REQUIREMENTS:
- Look carefully at the uploaded image and write content that reflects what is visibly present.
- Generate: 1) Instagram-ready caption, 2) Reel/Post/Story/Carousel text depending on selected formats, 3) SEO-friendly image title, 4) image description / alt text, 5) targeted hashtags.
- Use bookish, horror, YA thriller, dark academia, and Bookstagram audience language only when it fits the image and brief.
- Do not claim character identity, awards, reviews, or plot events unless supported by the brief or knowledge context.
"""
    if content_type == "youtube_content":
        return """
IMAGE-AWARE OUTPUT REQUIREMENTS:
- Treat the uploaded image as the visual basis for a YouTube thumbnail, community post, or video asset.
- Generate: 1) SEO title options, 2) YouTube description, 3) thumbnail/image title, 4) image description / alt text, 5) relevant hashtags/tags.
- Make the title clickable without becoming clickbait. Keep claims grounded in the brief and knowledge context.
"""
    if content_type == "linkedin_content":
        return """
IMAGE-AWARE OUTPUT REQUIREMENTS:
- Treat the uploaded image as a professional post/article visual.
- Generate: 1) LinkedIn post or article copy based on selected format, 2) professional title, 3) image description / alt text, 4) meta description, 5) restrained hashtags.
- Keep the tone credible, polished, and author-brand appropriate.
"""
    if content_type == "character_spotlight":
        return """
IMAGE-AWARE OUTPUT REQUIREMENTS:
- Use the uploaded image or selected character portrait as the visual basis when provided.
- Generate platform-native character spotlight text for the selected destination format.
- For Instagram, YouTube, and LinkedIn post formats, keep on-image and caption text short, punchy, and scannable.
- For blog destination formats, the description may be longer and more analytical.
- Include a concise visual headline/highlight that can fit on a generated image.
"""
    if content_type == "blog_post":
        return """
IMAGE-AWARE OUTPUT REQUIREMENTS:
- If an uploaded or character image is provided, use it as a blog visual and generate image title, alt title, caption, and meta description.
- Suggest 1-3 basic companion visuals such as quote-text graphics, character cards, comparison cards, or simple graph/chart ideas when useful.
- Keep blog visuals grounded in the selected book, character, and knowledge-base context.
"""
    return ""


def build_chatgpt_baseline_prompt(content_type: str, structured_brief: str) -> str:
    """Build a fair baseline prompt without Tell Tales Ink knowledge-base context."""
    return f"""You are ChatGPT responding in a fresh, general-purpose chat.

Create the requested {content_type} using only the user's topic and selections below.

Important comparison rules:
- Do not use the Tell Tales Ink knowledge bases, brand playbooks, templates, quote libraries, or hidden project context.
- Do not mention this is a comparison.
- Match the requested content type, selected formats, quantity, style, platform, audience, constraints, and CTA as closely as possible.
- If the brief asks for verifiable quotes, reviews, awards, or manuscript details but does not provide exact evidence, avoid inventing them.
- Produce polished, publication-ready copy.

User topic and selections:
{structured_brief}
"""


def comparison_side_by_side_html(mythos_content: str = "", chatgpt_content: str = "", model: str = "") -> str:
    """Render Tell Tales Ink and ChatGPT outputs side by side for human judging."""
    if not mythos_content and not chatgpt_content:
        return """
        <div style="border:1px dashed rgba(17,17,20,.18);border-radius:20px;padding:18px;color:#6b6b6b;background:#fff;">
            Generate a Tell Tales Ink draft first, then create the ChatGPT baseline here.
        </div>
        """

    safe_mythos = escape(mythos_content or "No Tell Tales Ink draft generated yet.")
    safe_chatgpt = escape(chatgpt_content or "No ChatGPT baseline generated yet.")
    safe_model = escape(model or "Selected ChatGPT model")
    return f"""
    <div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:14px;">
        <article style="border:1px solid rgba(142,31,47,.18);border-radius:18px;background:#fff;padding:14px;">
            <div style="font-weight:800;color:#8E1F2F;margin-bottom:4px;">Tell Tales Ink</div>
            <div style="font-size:12px;color:#6b6b6b;margin-bottom:10px;">Knowledge base + templates + brand workflow</div>
            <pre style="white-space:pre-wrap;font-family:inherit;font-size:13px;line-height:1.5;margin:0;">{safe_mythos}</pre>
        </article>
        <article style="border:1px solid rgba(17,17,20,.12);border-radius:18px;background:#fff;padding:14px;">
            <div style="font-weight:800;color:#111114;margin-bottom:4px;">ChatGPT</div>
            <div style="font-size:12px;color:#6b6b6b;margin-bottom:10px;">{safe_model}</div>
            <pre style="white-space:pre-wrap;font-family:inherit;font-size:13px;line-height:1.5;margin:0;">{safe_chatgpt}</pre>
        </article>
    </div>
    """


async def generate_chatgpt_comparison(
    *,
    content_type: str,
    structured_brief: str,
    mythos_content: str,
    model: str,
    comparison_status,
    comparison_view,
    chatgpt_output,
    chatgpt_prompt_path,
    chatgpt_draft_path,
) -> None:
    """Generate and display a fresh ChatGPT baseline next to Tell Tales Ink."""
    if not mythos_content or not mythos_content.strip():
        comparison_status.value = "Generate a Tell Tales Ink draft first."
        ui.notify(comparison_status.value, type="warning")
        return
    if not structured_brief or not structured_brief.strip():
        comparison_status.value = "No structured brief captured yet. Generate a Tell Tales Ink draft first."
        ui.notify(comparison_status.value, type="warning")
        return

    comparison_status.value = "Generating ChatGPT baseline..."
    prompt = build_chatgpt_baseline_prompt(content_type, structured_brief)
    selected_model = model or DEFAULT_MODEL
    try:
        baseline = await asyncio.to_thread(generate_text, prompt, selected_model)
        model_slug = re.sub(r"[^A-Za-z0-9_.-]+", "_", selected_model).strip("_") or "chatgpt"
        saved_prompt = save_output(prompt, "chatgpt_comparison", f"{content_type}_{model_slug}_prompt")
        saved_draft = save_output(baseline, "chatgpt_comparison", f"{content_type}_{model_slug}_draft")
    except Exception as exc:
        comparison_status.value = f"ChatGPT comparison failed: {exc}"
        ui.notify(comparison_status.value, type="negative")
        return

    chatgpt_output.value = baseline
    chatgpt_prompt_path.value = str(saved_prompt)
    chatgpt_draft_path.value = str(saved_draft)
    comparison_view.content = comparison_side_by_side_html(mythos_content, baseline, selected_model)
    comparison_status.value = "Comparison ready. Pick the stronger result below."
    ui.notify("Tell Tales Ink vs ChatGPT comparison ready.", type="positive")


def save_comparison_preference(
    *,
    preference: str,
    notes: str,
    content_type: str,
    structured_brief: str,
    mythos_content: str,
    chatgpt_model: str,
    chatgpt_content: str,
    chatgpt_draft_path: str,
    chatgpt_prompt_path: str,
    comparison_save_status,
    comparison_vote_path,
) -> None:
    """Save the human judge's comparison decision."""
    if not preference:
        comparison_save_status.value = "Choose Tell Tales Ink, ChatGPT, or Tie before saving."
        ui.notify(comparison_save_status.value, type="warning")
        return

    report = f"""# Tell Tales Ink vs ChatGPT Human Preference

- Content type: {content_type or "Not specified"}
- Preference: {preference}
- ChatGPT model: {chatgpt_model or DEFAULT_MODEL}
- ChatGPT draft path: {chatgpt_draft_path or "Not saved"}
- ChatGPT prompt path: {chatgpt_prompt_path or "Not saved"}

## Judge Notes

{notes or "_No notes provided._"}

## Structured Brief

```text
{structured_brief or "No structured brief captured."}
```

## Tell Tales Ink Output

{mythos_content or "_No Tell Tales Ink output captured._"}

## ChatGPT Output

{chatgpt_content or "_No ChatGPT output captured._"}
"""
    saved_path = save_output(report, "mythos_vs_chatgpt", "human_preference")
    comparison_vote_path.value = str(saved_path)
    comparison_save_status.value = "Comparison judgment saved."
    ui.notify("Comparison judgment saved.", type="positive")


def build_brief(
    content_type: str,
    topic: str,
    related_book: str,
    platform: str,
    social_objectives,
    audience,
    cta: str,
    constraints,
    quote_book: str,
    quote_moods,
    character_tags,
    podcast_format: str,
    podcast_speakers: str,
    podcast_roles: str,
    podcast_tone,
    podcast_length: str,
    elevenlabs_model: str,
    elevenlabs_voice_ids: str,
    podcast_show_title: str = "",
    podcast_episode_title: str = "",
    podcast_episode_number: str = "",
    blog_length: str = "",
    blog_format: str = "",
    blog_sections: str = "",
    blog_structure_options=None,
    blog_seo_keywords: str = "",
    blog_image_mode: str = "",
    blog_character: str = "",
    instagram_formats=None,
    instagram_hashtags: str = "",
    instagram_hook: str = "",
    instagram_image_content_type: str = "",
    instagram_carousel_slides=5,
    cs_character: str = "",
    cs_focus: str = "",
    cs_platform_format: str = "",
    cs_image_mode: str = "",
    nl_subject: str = "",
    nl_preview: str = "",
    nl_structure: str = "",
    pr_timing: str = "",
    pr_embargo_date: str = "",
    pr_city: str = "",
    pr_state: str = "",
    pr_release_date: str = "",
    pr_contact_name: str = "",
    pr_contact_title: str = "",
    pr_organization: str = "",
    pr_contact_email: str = "",
    pr_contact_phone: str = "",
    pr_website: str = "",
    pr_news_angle: str = "",
    pr_release_goal: str = "",
    pr_primary_announcement: str = "",
    pr_supporting_proof: str = "",
    pr_quote_source: str = "",
    pr_target_media: str = "",
    pr_required_assets: str = "",
    review_book: str = "",
    review_source: str = "",
    review_quote: str = "",
    review_mode: str = "",
    review_attribution: str = "",
    review_promo_line: str = "",
    review_graphic_formats=None,
    review_visual_style: str = "",
    review_quote_mode: str = "",
    review_manual_quote: str = "",
    review_image_type: str = "",
    review_instagram_formats=None,
    review_card_style: str = "",
    uploaded_image_path: str = "",
) -> str:
    lines = [
        brief_line("Topic", topic),
        brief_line("Related book/source", related_book),
        brief_line("Content type", content_type),
        brief_line("CTA", cta),
        brief_line("Constraints", constraints),
    ]
    if content_type != "press_release":
        platform_lines = [brief_line("Platform", platform)]
        if content_type not in HIDE_SOCIAL_AUDIENCE_CONTENT_TYPES:
            platform_lines += [
                brief_line("Social objectives", social_objectives),
                brief_line("Audience", audience),
            ]
        lines[3:3] = platform_lines

    if content_type in IMAGE_AWARE_CONTENT_TYPES:
        lines.extend(
            [
                brief_line("Uploaded image path", uploaded_image_path),
                "Instruction: If an uploaded image is provided, make the content respond directly to the visible image. If topic text is blank, use the image as the primary creative brief.",
            ]
        )

    if content_type == "instagram_caption":
        lines.extend(
            [
                brief_line("Instagram post formats", instagram_formats),
                brief_line("Instagram baseline hashtag strategy", instagram_hashtags),
                brief_line("Instagram opening hook (opening line / tone only)", instagram_hook),
                brief_line("Image content type", instagram_image_content_type),
            ]
        )
        lines.extend(
            instagram_generation_instructions(
                instagram_formats,
                carousel_slides=instagram_carousel_slides,
                image_content_type=instagram_image_content_type,
            )
        )

    if content_type == "quote_post":
        lines.extend(
            [
                brief_line("Requested quote book/source", quote_book),
                brief_line("Requested quote mood/category tags", quote_moods),
                brief_line("Requested character tags", character_tags),
                brief_line("Available character visual assets", character_asset_note(character_tags, book=quote_book)),
            ]
        )

    if content_type == "podcast":
        lines.extend(
            [
                brief_line("Podcast format", podcast_format),
                brief_line("Podcast speaker count", podcast_speakers),
                brief_line("Podcast speaker roles/names", podcast_roles),
                brief_line("Podcast tone", podcast_tone),
                brief_line("Podcast target length", podcast_length),
                brief_line("ElevenLabs model", elevenlabs_model),
                brief_line("ElevenLabs voice IDs", elevenlabs_voice_ids),
                brief_line("Podcast show title", podcast_show_title),
                brief_line("Podcast episode title", podcast_episode_title),
                brief_line("Podcast episode number", podcast_episode_number),
            ]
        )

    if content_type == "blog_post":
        lines.extend(
            [
                brief_line("Blog target length", blog_length),
                brief_line("Blog post format", blog_format),
                brief_line("Blog desired sections/outline", blog_sections),
                brief_line("Blog structure checklist", blog_structure_options),
                brief_line("Blog SEO keywords", blog_seo_keywords),
                brief_line("Blog image mode", blog_image_mode),
                brief_line("Blog selected character portrait", blog_character),
            ]
        )
        if str(blog_format or "").strip().lower() == "photo/gallery slideshow":
            lines.append(
                "Instruction: Generate a photo/gallery slideshow plan with multiple image concepts. Include image title, image caption, alt title, SEO description, and visual direction for each slide/image."
            )

    if content_type == "character_spotlight":
        lines.extend(
            [
                brief_line("Character spotlight subject", cs_character),
                brief_line("Character spotlight angle/focus", cs_focus),
                brief_line("Character spotlight platform format", cs_platform_format),
                brief_line("Character spotlight image mode", cs_image_mode),
                brief_line("Available character visual assets", character_asset_note(cs_character, book=related_book)),
            ]
        )

    if content_type == "newsletter_blurb":
        lines.extend(
            [
                brief_line("Newsletter subject line", nl_subject),
                brief_line("Newsletter preview text", nl_preview),
                brief_line("Newsletter structure", nl_structure),
            ]
        )

    if content_type == "review_pull_quote":
        lines.extend(
            [
                brief_line("Review book/source", review_book),
                brief_line("Selected review source", review_source),
                brief_line("Grounded review quote", review_quote),
                brief_line("Pull quote style", review_mode),
                brief_line("Quote source mode", review_quote_mode),
                brief_line("Attribution", review_attribution),
                brief_line("Brand promo line", review_promo_line),
                brief_line("Instagram formats", review_instagram_formats),
                brief_line("Pull quote image destination formats", review_graphic_formats),
                brief_line("Pull quote image color style", review_visual_style),
            ]
        )
        lines.extend(
            review_generation_instructions(
                review_instagram_formats,
                image_type=review_image_type,
                card_style=review_card_style,
                quote_mode=review_quote_mode,
                manual_quote=review_manual_quote,
            )
        )

    if content_type == "press_release":
        lines.extend(
            [
                brief_line("Press release timing", pr_timing),
                brief_line("Press release distribution", PRESS_RELEASE_DESTINATION),
                brief_line("Press release embargo date", pr_embargo_date),
                brief_line("Press release dateline city", pr_city),
                brief_line("Press release dateline state/region", pr_state),
                brief_line("Press release date", pr_release_date),
                brief_line("Press release media contact name", pr_contact_name),
                brief_line("Press release media contact title", pr_contact_title),
                brief_line("Press release organization/imprint", pr_organization),
                brief_line("Press release media contact email", pr_contact_email),
                brief_line("Press release media contact phone", pr_contact_phone),
                brief_line("Press release website", pr_website),
                brief_line("Press release news angle", pr_news_angle),
                brief_line("Press release supporting proof", pr_supporting_proof),
                brief_line("Press release quote source", pr_quote_source),
                brief_line("Press release target media", pr_target_media),
                brief_line("Companion assets to generate alongside the press release", pr_required_assets),
            ]
        )

    return "\n".join(lines)


def run_pipeline_with_image(content_type: str, topic: str, image_path: str | Path) -> dict:
    """Run the Tell Tales Ink pipeline with markdown context plus an uploaded image."""

    filtered_context = select_relevant_context(
        content_type=content_type,
        topic=topic,
    )
    filtered_context_path = save_output(
        content=filtered_context,
        content_type=content_type,
        label="filtered_context",
    )

    generation_prompt = create_generation_prompt(
        content_type=content_type,
        topic=topic,
        filtered_context=filtered_context,
    )
    generation_prompt = "\n\n".join(
        [
            generation_prompt,
            social_image_instructions(content_type),
            "Return organized, publication-ready output with clear labels for title, description, caption/post copy, hashtags, and image alt text where relevant.",
        ]
    )
    prompt_path = save_output(
        content=generation_prompt,
        content_type=content_type,
        label="generation_prompt",
    )

    generated_content = generate_text_with_image(generation_prompt, image_path)
    draft_path = save_output(
        content=generated_content,
        content_type=content_type,
        label="draft",
    )

    return {
        "filtered_context_path": filtered_context_path,
        "prompt_path": prompt_path,
        "draft_path": draft_path,
        "generated_content": generated_content,
    }


def platform_choices(content_type: str) -> list[str]:
    return list(PODCAST_DESTINATION_OPTIONS if content_type == "podcast" else PLATFORM_OPTIONS)


def style_choices(content_type: str) -> list[str]:
    return CAMPAIGN_STYLE_OPTIONS.get(content_type, ["default"])


def default_style(content_type: str) -> str:
    choices = style_choices(content_type)
    return choices[0] if choices else "default"


def build_campaign_asset_brief(
    *,
    content_type: str,
    asset_number: int,
    asset_count: int,
    campaign_topic: str,
    campaign_book: str,
    angle_name: str = "",
    angle_guidance: str = "",
    sibling_angles=None,
    campaign_tone: str,
    campaign_objectives,
    campaign_audience,
    campaign_cta: str,
    campaign_constraints,
    style: str,
    quantity: str,
    instagram_formats=None,
    instagram_hook="",
    youtube_formats=None,
    linkedin_formats=None,
    blog_length: str = "",
    blog_format: str = "",
    blog_sections: str = "",
    blog_structure_options=None,
    blog_seo_keywords: str = "",
    blog_image_mode: str = "",
    podcast_format: str = "",
    podcast_length: str = "",
    podcast_tone: str = "",
    podcast_speakers: str = "",
    podcast_roles: str = "",
    quote_book: str = "",
    quote_moods=None,
    quote_characters=None,
    quote_graphic_formats=None,
    quote_graphic_theme: str = "",
    pull_quote_graphic_formats=None,
    pull_quote_graphic_theme: str = "",
    press_timing: str = "",
    press_news_angle: str = "",
    press_release_goal: str = "",
    newsletter_structure: str = "",
    character_subject: str = "",
    character_format: str = "",
    character_image_mode: str = "",
) -> str:
    """Build a deterministic per-asset campaign brief from user selections."""
    lines = [
        "Campaign Mode Asset Brief",
        brief_line("Campaign topic", campaign_topic),
        brief_line("Campaign related book/source", campaign_book),
        brief_line("Selected campaign output", content_type),
        brief_line("Asset number", f"{asset_number} of {asset_count}"),
        brief_line("Requested quantity for this content type", quantity),
        brief_line("Style / variant", style),
        brief_line("Overall campaign tone", campaign_tone),
        brief_line("Campaign objectives", campaign_objectives),
        brief_line("Campaign audience", campaign_audience),
        brief_line("Campaign CTA", campaign_cta),
        brief_line("Campaign constraints", campaign_constraints),
        "Instruction: This asset is part of a coordinated campaign. Make it native to its platform, but keep the same core message, audience promise, and brand voice as the campaign.",
    ]

    if angle_name:
        lines.append(brief_line("Content angle for THIS post", f"{angle_name} — {angle_guidance}"))
        if sibling_angles:
            lines.append(brief_line("Distinct angles used across the full set", ", ".join(sibling_angles)))
        lines.append(
            "Instruction: This is one of several posts in the same campaign set, each with a DIFFERENT angle. "
            "Commit fully to the angle above and make this post visibly distinct from the others — different opening line, "
            "hook, structure, imagery, and emphasis. Do NOT reuse the framing, phrasing, or call-to-action wording of the sibling angles. "
            "Across the full set the posts must feel varied and non-repetitive."
        )

    if content_type == "instagram_caption":
        lines.extend(
            [
                brief_line("Instagram selected post formats", instagram_formats),
                brief_line("Instagram opening hook (opening line / tone only)", instagram_hook),
                brief_line("Hashtag strategy", "Use audience-relevant book community hashtags only when requested."),
            ]
        )
    elif content_type == "youtube_content":
        lines.append(brief_line("YouTube selected deliverables", youtube_formats))
    elif content_type == "linkedin_content":
        lines.append(brief_line("LinkedIn selected deliverables", linkedin_formats))
    elif content_type == "blog_post":
        lines.extend(
            [
                brief_line("Blog target length", blog_length),
                brief_line("Blog post format", blog_format),
                brief_line("Blog structure checklist", blog_structure_options),
                brief_line("Blog sections / outline", blog_sections),
                brief_line("Blog SEO keywords", blog_seo_keywords),
                brief_line("Blog image mode", blog_image_mode),
            ]
        )
    elif content_type == "podcast":
        lines.extend(
            [
                brief_line("Podcast format", podcast_format),
                brief_line("Podcast length", podcast_length),
                brief_line("Podcast delivery tone", podcast_tone),
                brief_line("Podcast speaker count", podcast_speakers),
                brief_line("Podcast speaker roles/names", podcast_roles),
                "Instruction: Write the script as labeled speaker turns (e.g. 'Host:', 'Guest:') so it can be voiced per speaker.",
            ]
        )
    elif content_type == "quote_post":
        lines.extend(
            [
                brief_line("Quote book/source", quote_book),
                brief_line("Quote mood/category tags", quote_moods),
                brief_line("Quote character tags", quote_characters),
                brief_line("Quote image destination formats", quote_graphic_formats),
                brief_line("Quote image color style", quote_graphic_theme),
                brief_line("Available character visual assets", character_asset_note(quote_characters, book=quote_book)),
            ]
        )
    elif content_type == "review_pull_quote":
        lines.extend(
            [
                brief_line("Review source rule", "Use only real review text from the knowledge base."),
                brief_line("Pull quote style", style),
                brief_line("Pull quote image destination formats", pull_quote_graphic_formats),
                brief_line("Pull quote image color style", pull_quote_graphic_theme),
            ]
        )
    elif content_type == "press_release":
        lines.extend(
            [
                brief_line("Press release timing", press_timing),
                brief_line("News angle", press_news_angle),
            ]
        )
    elif content_type == "newsletter_blurb":
        lines.append(brief_line("Newsletter structure", newsletter_structure))
    elif content_type == "character_spotlight":
        lines.extend(
            [
                brief_line("Character spotlight subject", character_subject),
                brief_line("Character spotlight destination format", character_format),
                brief_line("Character spotlight image mode", character_image_mode),
                brief_line("Available character visual assets", character_asset_note(character_subject, book=campaign_book)),
            ]
        )

    return "\n".join(lines)


async def generate_campaign_from_fields(
    *,
    campaign_topic,
    campaign_name="",
    campaign_start_date="",
    campaign_cadence="",
    campaign_book,
    campaign_tone,
    campaign_objectives,
    campaign_audience,
    campaign_cta,
    campaign_constraints,
    campaign_formats,
    campaign_widgets,
    status,
    campaign_output,
    campaign_preview,
    campaign_brief_output,
    campaign_path,
    saved_draft_path=None,
    campaign_progress=None,
    campaign_progress_label=None,
    campaign_sections=None,
    podcast_scripts_out=None,
    campaign_form_fields=None,
) -> None:
    selected = [content_type for content_type, _label in CAMPAIGN_FORMATS if content_type in (campaign_formats or [])]
    if not campaign_topic or not str(campaign_topic).strip():
        status.value = "Enter a campaign topic first."
        ui.notify(status.value, type="warning")
        return
    if not selected:
        status.value = "Select at least one campaign content type."
        ui.notify(status.value, type="warning")
        return

    status.value = "Generating campaign..."
    campaign_output.value = ""
    campaign_preview.content = build_campaign_preview_html([])
    if campaign_sections is not None:
        campaign_sections.clear()
    campaign_path.value = ""
    if saved_draft_path is not None:
        saved_draft_path.value = ""
    all_briefs: list[str] = []
    results: list[str] = []
    asset_records: list[dict[str, object]] = []
    asset_total = sum(max(1, int(float(campaign_widgets[content_type]["quantity"].value or 1))) for content_type in selected)
    completed_assets = 0
    started_at = datetime.now()
    # Posting calendar basis. Cadence is optional; when set, each content type's posts are
    # spaced at the cadence interval from the (optional) start date, so all types share the
    # same campaign window and every post lands on a specific calendar date.
    cadence_days = {"daily": 1, "every 2 days": 2, "twice a week": 3, "weekly": 7}.get(
        str(campaign_cadence or "").strip().lower()
    )
    start = None
    start_text = str(campaign_start_date or "").strip()
    if start_text:
        try:
            start = datetime.strptime(start_text, "%Y-%m-%d")
        except ValueError:
            start = None

    def post_schedule(post_index: int):
        """(sort_key, date_label) for the post_index-th post of a content type, or None."""
        if not cadence_days:
            return None
        offset = (max(1, int(post_index)) - 1) * cadence_days
        if start:
            day = start + timedelta(days=offset)
            return (day, day.strftime("%a %b %d, %Y"))
        return (offset, f"Day {1 + offset}")

    try:
        for content_type in selected:
            widgets = campaign_widgets[content_type]
            count = max(1, int(float(widgets["quantity"].value or 1)))
            quantity = str(count)
            # Vary content + visuals across a multi-post set so they don't feel repetitive.
            sibling_angle_names = campaign_set_angle_names(count) if count > 1 else []
            for index in range(1, count + 1):
                label = CAMPAIGN_FORMAT_LABELS.get(content_type, content_type)
                angle = campaign_post_angle(index) if count > 1 else None
                post_theme = ""
                if count > 1:
                    _base_theme = (
                        (widgets.get("visual_style").value if widgets.get("visual_style") else None)
                        or (widgets.get("graphic_theme").value if widgets.get("graphic_theme") else None)
                        or "Gothic"
                    )
                    post_theme = campaign_visual_theme_for_post(_base_theme, index)
                if campaign_progress is not None and campaign_progress_label is not None:
                    await set_generation_progress(
                        progress=campaign_progress,
                        label=campaign_progress_label,
                        status=status,
                        started_at=started_at,
                        completed_steps=completed_assets,
                        total_steps=asset_total + 1,
                        completed_assets=completed_assets,
                        total_assets=asset_total,
                        message=f"Generating {label} asset {index} of {count}...",
                    )
                brief = build_campaign_asset_brief(
                    content_type=content_type,
                    asset_number=index,
                    asset_count=count,
                    campaign_topic=campaign_topic,
                    campaign_book=campaign_book,
                    angle_name=(angle["name"] if angle else ""),
                    angle_guidance=(angle["content"] if angle else ""),
                    sibling_angles=sibling_angle_names,
                    campaign_tone=campaign_tone,
                    campaign_objectives=campaign_objectives,
                    campaign_audience=campaign_audience,
                    campaign_cta=campaign_cta,
                    campaign_constraints=campaign_constraints,
                    style=widgets["style"].value,
                    quantity=quantity,
                    instagram_formats=campaign_widgets["instagram_caption"].get("formats").value,
                    instagram_hook=(campaign_widgets["instagram_caption"].get("hook").value if campaign_widgets["instagram_caption"].get("hook") else ""),
                    youtube_formats=campaign_widgets["youtube_content"].get("formats").value,
                    linkedin_formats=campaign_widgets["linkedin_content"].get("formats").value,
                    blog_length=campaign_widgets["blog_post"].get("length").value,
                    blog_format=campaign_widgets["blog_post"].get("format").value,
                    blog_sections=campaign_widgets["blog_post"].get("sections").value,
                    blog_structure_options=campaign_widgets["blog_post"].get("structure").value,
                    blog_seo_keywords=campaign_widgets["blog_post"].get("seo").value,
                    blog_image_mode=campaign_widgets["blog_post"].get("image_mode").value,
                    podcast_format=campaign_widgets["podcast"].get("format").value,
                    podcast_length=campaign_widgets["podcast"].get("length").value,
                    podcast_tone=campaign_widgets["podcast"].get("tone").value,
                    podcast_speakers=(str(int(campaign_widgets["podcast"]["speakers"].value)) if (campaign_widgets["podcast"].get("speakers") and campaign_widgets["podcast"]["speakers"].value) else ""),
                    podcast_roles=(campaign_widgets["podcast"].get("roles").value if campaign_widgets["podcast"].get("roles") else ""),
                    quote_book=campaign_widgets["quote_post"].get("book").value,
                    quote_moods=campaign_widgets["quote_post"].get("moods").value,
                    quote_characters=campaign_widgets["quote_post"].get("characters").value,
                    quote_graphic_formats=campaign_widgets["quote_post"].get("graphic_formats").value,
                    quote_graphic_theme=campaign_widgets["quote_post"].get("graphic_theme").value,
                    pull_quote_graphic_formats=campaign_widgets["review_pull_quote"].get("graphic_formats").value,
                    pull_quote_graphic_theme=campaign_widgets["review_pull_quote"].get("graphic_theme").value,
                    press_timing=campaign_widgets["press_release"].get("timing").value,
                    press_news_angle=campaign_widgets["press_release"].get("news_angle").value,
                    press_release_goal="",
                    newsletter_structure=campaign_widgets["newsletter_blurb"].get("structure").value,
                    character_subject=campaign_widgets["character_spotlight"].get("character").value,
                    character_format=campaign_widgets["character_spotlight"].get("format").value,
                    character_image_mode=campaign_widgets["character_spotlight"].get("image_mode").value,
                )
                all_briefs.append(brief)
                result = await asyncio.to_thread(run_pipeline, content_type=content_type, topic=brief)
                heading = f"## {label} {index}" if count > 1 else f"## {label}"
                post_slot = post_schedule(index)
                if post_slot is not None:
                    heading = f"{heading} — 🗓 {post_slot[1]}"
                generated_content = result["generated_content"]
                results.append(f"{heading}\n\n{generated_content}")
                if content_type == "podcast" and podcast_scripts_out is not None:
                    podcast_scripts_out.append((f"{label} {index} of {count}".strip(), generated_content))
                preview_fields = {
                    "instagram_formats": campaign_widgets["instagram_caption"].get("formats").value,
                    "instagram_hashtags": "Use selected audience hashtags",
                    "instagram_hook": widgets["style"].value,
                    "youtube_formats": campaign_widgets["youtube_content"].get("formats").value,
                    "youtube_keywords": campaign_widgets["youtube_content"].get("style").value,
                    "youtube_link": campaign_cta,
                    "linkedin_formats": campaign_widgets["linkedin_content"].get("formats").value,
                    "linkedin_angle": campaign_widgets["linkedin_content"].get("style").value,
                    "linkedin_cta": campaign_cta,
                    "quote_moods": campaign_widgets["quote_post"].get("moods").value,
                    "quote_characters": campaign_widgets["quote_post"].get("characters").value,
                    "character_subject": campaign_widgets["character_spotlight"].get("character").value,
                    "character_format": campaign_widgets["character_spotlight"].get("format").value,
                    "character_image_mode": campaign_widgets["character_spotlight"].get("image_mode").value,
                    "blog_format": campaign_widgets["blog_post"].get("format").value,
                    "blog_length": campaign_widgets["blog_post"].get("length").value,
                    "blog_structure": campaign_widgets["blog_post"].get("sections").value,
                    "blog_structure_checklist": campaign_widgets["blog_post"].get("structure").value,
                    "blog_image_mode": campaign_widgets["blog_post"].get("image_mode").value,
                }
                quote_graphic = {}
                if content_type in {"quote_post", "review_pull_quote", "character_spotlight", "blog_post"}:
                    quote_graphic = render_campaign_quote_graphic(
                        content_type=content_type,
                        content=generated_content,
                        widgets=widgets,
                        campaign_topic=str(campaign_topic or "campaign"),
                        asset_number=index,
                        theme_override=post_theme,
                        book=campaign_book,
                    )
                # Campaign image generation for image-capable platforms (BUG-IMG-04 + LinkedIn/YouTube/Newsletter).
                ig_image_paths: list[str] = []
                if content_type in {"instagram_caption", "linkedin_content", "youtube_content", "newsletter_blurb"}:
                    selected_image_formats = [f for f in (widgets.get("generate_images").value or []) if f] if widgets.get("generate_images") else []
                    if selected_image_formats:
                        # Instagram chips are post-format names that map to aspect ratios; the
                        # other platforms select aspect-labeled image formats directly.
                        image_labels = (
                            instagram_visual_formats_for_post(selected_image_formats, None, 5)
                            if content_type == "instagram_caption"
                            else selected_image_formats
                        )
                        user_visual_style = (widgets.get("visual_style").value if widgets.get("visual_style") else "Gothic") or "Gothic"
                        post_visual_style = post_theme or user_visual_style
                        # Per-post visual direction so generated images vary in subject/composition,
                        # not just text. For a multi-post set, don't pin every post to one base image.
                        post_topic = str(campaign_topic or "campaign")
                        if angle:
                            post_topic = f"{post_topic}. Visual treatment for this post ({angle['name']}): {angle['visual']}"
                        post_base_image = "" if count > 1 else ((widgets.get("base_image").value if widgets.get("base_image") else "") or "")
                        image_package = await asyncio.to_thread(
                            render_generated_visual_package,
                            content_type=content_type,
                            content=generated_content,
                            topic=post_topic,
                            format_labels=image_labels,
                            theme_name=post_visual_style,
                            image_content_type=widgets.get("image_content_type").value if widgets.get("image_content_type") else "",
                            base_image_path=post_base_image,
                            book=(widgets.get("book").value if widgets.get("book") else campaign_book),
                        )
                        if isinstance(image_package, dict):
                            ig_image_paths = [str(p) for p in image_package.get("paths", [])]
                            if ig_image_paths:
                                results[-1] += f"\n\n**Generated images:** {len(ig_image_paths)} file(s) — package: {image_package.get('zip_path', '')}"
                asset_records.append(
                    {
                        "content_type": content_type,
                        "label": label,
                        "asset_label": f"Asset {index} of {count}",
                        "asset_index": index,
                        "asset_count": count,
                        "post_date": post_slot[1] if post_slot else "",
                        "post_sort": post_slot[0] if post_slot else None,
                        "style": widgets["style"].value,
                        "content": generated_content,
                        "preview_fields": preview_fields,
                        "quote_graphic": quote_graphic,
                        "image_paths": ig_image_paths,
                        "widgets": widgets,
                    }
                )
                completed_assets += 1
                if campaign_progress is not None and campaign_progress_label is not None:
                    await set_generation_progress(
                        progress=campaign_progress,
                        label=campaign_progress_label,
                        status=status,
                        started_at=started_at,
                        completed_steps=completed_assets,
                        total_steps=asset_total + 1,
                        completed_assets=completed_assets,
                        total_assets=asset_total,
                        message=f"Completed {label} asset {index} of {count}.",
                    )
    except Exception as exc:
        status.value = f"Campaign generation failed: {exc}"
        ui.notify(status.value, type="negative")
        return

    if campaign_progress is not None and campaign_progress_label is not None:
        await set_generation_progress(
            progress=campaign_progress,
            label=campaign_progress_label,
            status=status,
            started_at=started_at,
            completed_steps=asset_total,
            total_steps=asset_total + 1,
            completed_assets=completed_assets,
            total_assets=asset_total,
            message="Saving campaign bundle...",
        )
    # Campaign title + posting schedule (BUG-CM-03/04).
    campaign_title = str(campaign_name or "").strip() or str(campaign_topic).strip()[:60] or "Campaign"
    # Posting calendar grouped by date (only when a cadence was set). Each post already
    # carries its scheduled date, so we group them to show exactly what publishes when.
    scheduled = [
        (
            record.get("post_sort"),
            str(record.get("post_date")),
            content_type_label(str(record.get("content_type") or "content"))
            + (f" {record.get('asset_label')}" if int(record.get("asset_count") or 1) > 1 else ""),
        )
        for record in asset_records
        if record.get("post_date")
    ]
    if scheduled:
        scheduled.sort(key=lambda item: item[0])
        schedule_lines = ["## Posting Calendar", ""]
        current_date = None
        for _sort_key, date_label, slot in scheduled:
            if date_label != current_date:
                schedule_lines.append(f"\n**{date_label}**")
                current_date = date_label
            schedule_lines.append(f"- {slot.strip()}")
        schedule_block = "\n".join(schedule_lines) + "\n\n---\n\n"
    else:
        schedule_block = ""

    combined = f"# {campaign_title}\n\n{schedule_block}" + "\n\n---\n\n".join(results)
    saved_path = save_output(combined, "campaign_mode", slugify_filename(campaign_title, "bundle"))
    draft_record = save_draft(
        title=campaign_title[:90],
        content_type="campaign_mode",
        content=combined,
        source_path=saved_path,
        metadata={
            "campaign_name": campaign_title,
            "formats": selected,
            "asset_count": len(results),
            "topic": str(campaign_topic or "").strip(),
            "start_date": start_text,
            "cadence": str(campaign_cadence or "").strip(),
            "related_book": normalize_selected(campaign_book),
            "brief": "\n\n---\n\n".join(all_briefs),
            "objectives": normalize_selected(campaign_objectives),
            "audience": normalize_selected(campaign_audience),
            "form_fields": campaign_form_fields or {},
        },
    )
    campaign_output.value = combined
    campaign_preview.content = build_campaign_preview_html(asset_records)
    if campaign_sections is not None:
        campaign_sections.clear()
        with campaign_sections:
            if not asset_records:
                ui.label("No assets generated.").classes("mce-muted")
            # Group assets by content type — one accordion section per type (first open),
            # each showing its posts with real text, real generated images, and download.
            grouped_records: dict[str, list] = {}
            for record in asset_records:
                grouped_records.setdefault(str(record.get("content_type") or "content"), []).append(record)
            for group_index, (rec_ct, type_records) in enumerate(grouped_records.items()):
                type_label = content_type_label(rec_ct)
                with ui.expansion(f"{type_label}  ({len(type_records)})", icon="folder_open", value=(group_index == 0)).classes("mce-expansion"):
                    for record in type_records:
                        rec_content = str(record.get("content") or "")
                        rec_label = str(record.get("asset_label") or "").strip()
                        post_date = str(record.get("post_date") or "")
                        rec_widgets = record.get("widgets") if isinstance(record.get("widgets"), dict) else {}
                        rec_title = f"{type_label} {rec_label}".strip()
                        quote_graphic = record.get("quote_graphic") if isinstance(record.get("quote_graphic"), dict) else {}
                        asset_images = [str(p) for p in (record.get("image_paths") or []) if str(p or "").strip()]
                        asset_images += [str(p) for p in (quote_graphic.get("paths") or []) if str(p or "").strip()]
                        with ui.card().classes("mce-subcard w-full"):
                            head_text = rec_title + (f"   ·   🗓 {post_date}" if post_date else "")
                            ui.label(head_text).classes("mce-section-title")
                            ui.markdown(rec_content).classes("w-full mce-rendered-output")
                            asset_image_row = ui.row().classes("w-full mce-gallery")

                            def render_asset_images(row=asset_image_row, paths=tuple(asset_images)) -> None:
                                row.clear()
                                with row:
                                    for image_path in [p for p in paths if p][:8]:
                                        ui.image(image_path).classes("mce-gallery-thumb")

                            render_asset_images()
                            if not asset_images:
                                ui.label("No images generated for this asset.").classes("mce-muted")

                            async def regenerate_asset_images(content=rec_content, ct=rec_ct, wgts=rec_widgets, row=asset_image_row) -> None:
                                gi = wgts.get("generate_images")
                                selected = [f for f in (gi.value or []) if f] if gi else []
                                if not selected:
                                    ui.notify("Pick at least one image format on the content card first.", type="warning")
                                    return
                                labels = instagram_visual_formats_for_post(selected, None, 5) if ct == "instagram_caption" else selected
                                pkg = await asyncio.to_thread(
                                    render_generated_visual_package,
                                    content_type=ct,
                                    content=content,
                                    topic=str(campaign_topic or "campaign"),
                                    format_labels=labels,
                                    theme_name=(wgts.get("visual_style").value if wgts.get("visual_style") else "Gothic") or "Gothic",
                                    image_content_type=wgts.get("image_content_type").value if wgts.get("image_content_type") else "",
                                    base_image_path=(wgts.get("base_image").value if wgts.get("base_image") else "") or "",
                                    book=(wgts.get("book").value if wgts.get("book") else ""),
                                )
                                paths = [str(p) for p in pkg.get("paths", [])] if isinstance(pkg, dict) else []
                                if paths:
                                    render_asset_images(row=row, paths=tuple(paths))
                                    ui.notify(f"Regenerated {len(paths)} image(s).", type="positive")
                                else:
                                    ui.notify((pkg.get("error") if isinstance(pkg, dict) else "") or "No images generated.", type="warning")

                            with ui.row().classes("mce-actions"):
                                make_secondary_button(
                                    "Download",
                                    lambda c=rec_content, t=rec_title: download_docx(c, t or "campaign_asset"),
                                )
                                make_secondary_button(
                                    "Save to Drafts",
                                    lambda c=rec_content, ct=rec_ct, t=rec_title: (
                                        save_draft(title=(t[:90] or ct), content_type=ct, content=c),
                                        ui.notify("Saved to drafts.", type="positive"),
                                    ),
                                )
                                if rec_ct in {"instagram_caption", "linkedin_content", "youtube_content", "newsletter_blurb"}:
                                    make_secondary_button(
                                        "Regenerate images",
                                        lambda fn=regenerate_asset_images: fn(),
                                    )
    campaign_brief_output.value = "\n\n---\n\n".join(all_briefs)
    campaign_path.value = str(saved_path)
    if saved_draft_path is not None:
        saved_draft_path.value = draft_record["path"]
    status.value = f"Campaign ready: {len(results)} assets generated."
    if campaign_progress is not None and campaign_progress_label is not None:
        await finish_generation_progress(
            progress=campaign_progress,
            label=campaign_progress_label,
            status=status,
            total_assets=asset_total,
            message=f"Campaign ready: {len(results)} assets generated.",
        )
    ui.notify(status.value, type="positive")


def slugify_filename(value: str, fallback: str = "podcast") -> str:
    value = (value or "").strip() or fallback
    value = re.sub(r"[^\w\s-]", "", value, flags=re.UNICODE)
    value = re.sub(r"[-\s]+", "_", value).strip("_").lower()
    return value or fallback


def _docx_inline_runs(paragraph, text: str) -> None:
    """Add text to a docx paragraph, rendering **bold** / *italic* markdown runs."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**") or token.startswith("__"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token[1:-1]).italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def build_docx(content: str, stem: str) -> Path:
    """Render markdown-ish text to a .docx and return its path (downloads are .docx)."""
    from docx import Document

    document = Document()
    for raw in str(content or "").splitlines():
        stripped = raw.strip()
        if not stripped or stripped == "---":
            continue
        banner = re.match(r"^={2,}\s*(.+?)\s*={2,}$", stripped)
        if banner:
            document.add_heading(banner.group(1), level=2)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            document.add_heading(re.sub(r"[*_`]", "", heading.group(2)), level=min(len(heading.group(1)), 4))
            continue
        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        if bullet:
            _docx_inline_runs(document.add_paragraph(style="List Bullet"), bullet.group(1))
            continue
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if numbered:
            _docx_inline_runs(document.add_paragraph(style="List Number"), numbered.group(1))
            continue
        _docx_inline_runs(document.add_paragraph(), stripped)

    out_dir = PROJECT_ROOT / "outputs" / "downloads"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{slugify_filename(stem, 'document')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    document.save(str(path))
    return path


def download_docx(content: str, stem: str) -> None:
    """Build a .docx from content and trigger a browser download (with a clear warning if empty)."""
    if not str(content or "").strip():
        ui.notify("Generate content before downloading.", type="warning")
        return
    path = build_docx(content, stem)
    ui.download(str(path), f"{slugify_filename(stem, 'document')}.docx")


def normalize_speaker_label(value: str) -> str:
    value = re.sub(r"[*_`#>]", "", value or "")
    value = value.strip().strip("[]").strip()
    value = re.sub(r"[^A-Za-z0-9_ -]", "", value)
    value = re.sub(r"[-\s]+", "_", value).strip("_")
    return value.upper() or "HOST"


def parse_podcast_segments(script: str) -> list[tuple[str, str]]:
    """Extract speaker-labeled turns, falling back to host narration."""
    segments: list[tuple[str, str]] = []
    current_label = ""
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_label, current_lines
        text = strip_production_tags("\n".join(current_lines).strip())
        if current_label and text:
            segments.append((current_label, text))
        current_lines = []

    for raw_line in (script or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue

        bracket_match = re.match(
            r"^\[(?!SFX:|INTRO MUSIC:|OUTRO MUSIC:|MUSIC BED:|PAUSE:|BEAT|CUT IF|OPTIONAL|USER EDIT:)([A-Za-z0-9_ -]+)\]\s*(.*)$",
            line,
            flags=re.IGNORECASE,
        )
        if bracket_match:
            flush()
            current_label = normalize_speaker_label(bracket_match.group(1))
            remainder = bracket_match.group(2).strip(" :-")
            current_lines = [remainder] if remainder else []
            continue

        colon_match = re.match(r"^([A-Za-z][A-Za-z0-9_ -]{1,40}):\s+(.+)$", line)
        if colon_match and not line.startswith("["):
            flush()
            current_label = normalize_speaker_label(colon_match.group(1))
            current_lines = [colon_match.group(2).strip()]
            continue

        if current_label:
            current_lines.append(line)

    flush()

    if segments:
        return segments

    fallback_lines = []
    for raw_line in (script or "").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or line.startswith("---"):
            continue
        if re.match(r"^\[(SFX:|INTRO MUSIC:|OUTRO MUSIC:|MUSIC BED:|PAUSE:|BEAT)", line, flags=re.IGNORECASE):
            continue
        fallback_lines.append(line)
    fallback_text = strip_production_tags("\n".join(fallback_lines).strip())
    return [("HOST", fallback_text)] if fallback_text else []


def split_tts_text(text: str, max_chars: int = PODCAST_TTS_MAX_CHARS) -> list[str]:
    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= max_chars:
        return [text]

    sentences = re.split(r"(?<=[.!?])\s+", text)
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        proposed = f"{current} {sentence}".strip()
        if len(proposed) <= max_chars:
            current = proposed
            continue
        if current:
            chunks.append(current)
        if len(sentence) <= max_chars:
            current = sentence
            continue
        words = sentence.split()
        current = ""
        for word in words:
            proposed_word = f"{current} {word}".strip()
            if len(proposed_word) > max_chars and current:
                chunks.append(current)
                current = word
            else:
                current = proposed_word
    if current:
        chunks.append(current)
    return chunks


def compact_podcast_segments(segments: list[tuple[str, str]]) -> list[tuple[str, str]]:
    compacted: list[tuple[str, str]] = []
    pending_label = ""
    pending_text = ""

    def flush() -> None:
        nonlocal pending_label, pending_text
        if pending_label and pending_text.strip():
            for chunk in split_tts_text(pending_text):
                compacted.append((pending_label, chunk))
        pending_label = ""
        pending_text = ""

    for label, text in segments:
        text = text.strip()
        proposed = f"{pending_text}\n\n{text}".strip() if pending_text else text
        if label == pending_label and len(proposed) <= PODCAST_TTS_MAX_CHARS:
            pending_text = proposed
            continue
        flush()
        pending_label = label
        pending_text = text
    flush()
    return compacted


def voice_options(voices: list[dict]) -> dict[str, str]:
    return {
        voice["voice_id"]: f"{voice.get('name', 'Unnamed')} · {voice.get('category', 'voice')}"
        for voice in voices
        if voice.get("voice_id")
    }


def voice_preview_url(voice_id: str, voices: list[dict]) -> str:
    for voice in voices:
        if voice.get("voice_id") == voice_id:
            return voice.get("preview_url", "") or ""
    return ""


def build_voice_casting_from_selects(host_voice: str, guest_voice: str, guest_2_voice: str) -> str:
    lines = []
    if host_voice:
        lines.append(f"HOST={host_voice}")
        lines.append(f"NARRATOR={host_voice}")
    if guest_voice:
        lines.extend([f"GUEST={guest_voice}", f"GUEST_1={guest_voice}", f"AUTHOR={guest_voice}"])
    if guest_2_voice:
        lines.extend([f"GUEST_2={guest_2_voice}", f"COHOST={guest_2_voice}"])
    return "\n".join(lines)


def parse_voice_casting(value: str) -> dict[str, str]:
    casting: dict[str, str] = {}
    for raw_line in (value or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if "=" in line:
            label, voice_id = line.split("=", 1)
        elif ":" in line:
            label, voice_id = line.split(":", 1)
        else:
            continue
        casting[normalize_speaker_label(label)] = voice_id.strip()
    return casting


def resolve_voice_for_label(label: str, casting: dict[str, str], fallback_ids: list[str]) -> str:
    normalized = normalize_speaker_label(label)
    if normalized in casting:
        return casting[normalized]
    if normalized.startswith("GUEST_2") or normalized in {"COHOST", "SECOND_GUEST"}:
        return casting.get("GUEST_2") or (fallback_ids[2] if len(fallback_ids) > 2 else fallback_ids[-1])
    if normalized.startswith("GUEST") or normalized in {"AUTHOR", "CRITIC", "INTERVIEWEE"}:
        return casting.get("GUEST") or casting.get("GUEST_1") or (fallback_ids[1] if len(fallback_ids) > 1 else fallback_ids[0])
    return casting.get("HOST") or casting.get("NARRATOR") or fallback_ids[0]


async def render_podcast_audio_native(
    *,
    script: str,
    document_name: str,
    voices: list[dict],
    host_voice: str,
    guest_voice: str,
    guest_2_voice: str,
    model_id: str,
    stability: float,
    similarity_boost: float,
    style: float,
    speed: float,
    speaker_boost: bool,
    preview_only: bool,
    status,
    audio_player,
    full_audio_path,
    package_path,
    saved_draft_path=None,
    progress=None,
    progress_label=None,
) -> None:
    if not script or not script.strip():
        status.value = "Generate or paste a podcast script first."
        ui.notify(status.value, type="warning")
        return
    if not os.getenv("ELEVENLABS_API_KEY"):
        status.value = "ELEVENLABS_API_KEY is missing. Add it to .env and restart the server."
        ui.notify(status.value, type="negative")
        return
    if not voices:
        status.value = "No ElevenLabs voices are loaded yet."
        ui.notify(status.value, type="warning")
        return

    selected_voice_ids = [voice_id for voice_id in [host_voice, guest_voice, guest_2_voice] if voice_id]
    if not selected_voice_ids:
        selected_voice_ids = [voice["voice_id"] for voice in voices if voice.get("voice_id")][:1]
    if not selected_voice_ids:
        status.value = "No usable ElevenLabs voices found."
        ui.notify(status.value, type="negative")
        return

    segments = compact_podcast_segments(parse_podcast_segments(script))
    if preview_only:
        segments = segments[:PODCAST_PREVIEW_MAX_SEGMENTS]
    if not segments:
        status.value = "No renderable spoken text found in the podcast script."
        ui.notify(status.value, type="warning")
        return

    status.value = "Rendering ElevenLabs audio..."
    casting = parse_voice_casting(build_voice_casting_from_selects(host_voice, guest_voice, guest_2_voice))
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_stem = slugify_filename(document_name or "podcast_audio")
    render_mode = "preview" if preview_only else "full"
    output_dir = PROJECT_ROOT / "outputs" / f"{file_stem}_{render_mode}_audio_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)

    rendered_paths: list[Path] = []
    started_at = datetime.now()
    try:
        if progress is not None and progress_label is not None:
            await set_generation_progress(
                progress=progress,
                label=progress_label,
                status=status,
                started_at=started_at,
                completed_steps=0,
                total_steps=len(segments) + 2,
                completed_assets=0,
                total_assets=len(segments),
                message=f"Preparing {render_mode} audio render...",
            )
        for index, (label, text) in enumerate(segments, start=1):
            voice_id = resolve_voice_for_label(label, casting, selected_voice_ids)
            segment_path = output_dir / f"{index:03d}_{label.lower()}.mp3"
            if progress is not None and progress_label is not None:
                await set_generation_progress(
                    progress=progress,
                    label=progress_label,
                    status=status,
                    started_at=started_at,
                    completed_steps=index - 1,
                    total_steps=len(segments) + 2,
                    completed_assets=index - 1,
                    total_assets=len(segments),
                    message=f"Rendering voice segment {index} of {len(segments)}...",
                )
            await asyncio.to_thread(
                synthesize_speech,
                text=text,
                voice_id=voice_id,
                output_path=segment_path,
                model_id=model_id or ELEVENLABS_MODEL_OPTIONS[0],
                output_format=PODCAST_OUTPUT_FORMAT,
                strip_tags=True,
                stability=stability,
                similarity_boost=similarity_boost,
                style=style,
                speed=speed,
                use_speaker_boost=speaker_boost,
            )
            rendered_paths.append(segment_path)
            if progress is not None and progress_label is not None:
                await set_generation_progress(
                    progress=progress,
                    label=progress_label,
                    status=status,
                    started_at=started_at,
                    completed_steps=index,
                    total_steps=len(segments) + 2,
                    completed_assets=index,
                    total_assets=len(segments),
                    message=f"Completed voice segment {index} of {len(segments)}.",
                )

        if progress is not None and progress_label is not None:
            await set_generation_progress(
                progress=progress,
                label=progress_label,
                status=status,
                started_at=started_at,
                completed_steps=len(segments) + 1,
                total_steps=len(segments) + 2,
                completed_assets=len(rendered_paths),
                total_assets=len(segments),
                message="Combining segments into MP3...",
            )
        combined = AudioSegment.empty()
        pause = AudioSegment.silent(duration=650)
        for index, segment_path in enumerate(rendered_paths):
            if index:
                combined += pause
            combined += AudioSegment.from_file(segment_path)

        output_path = output_dir / f"{file_stem}_{render_mode}.mp3"
        combined.export(output_path, format="mp3", bitrate="128k")

        manifest_path = output_dir / "podcast_audio_manifest.md"
        manifest_path.write_text(
            "\n".join(
                [
                    "# Podcast Audio Manifest",
                    "",
                    f"Mode: {render_mode}",
                    f"Document: {document_name or file_stem}",
                    f"Model: {model_id or ELEVENLABS_MODEL_OPTIONS[0]}",
                    f"Rendered segments: {len(rendered_paths)}",
                    "",
                    "## Voice Casting",
                    build_voice_casting_from_selects(host_voice, guest_voice, guest_2_voice) or "Auto-assigned first available voice.",
                    "",
                    "## Segment Files",
                    *[f"- {path.name}" for path in rendered_paths],
                    "",
                    f"Combined MP3: {output_path.name}",
                ]
            ),
            encoding="utf-8",
        )

        zip_path = output_dir.with_suffix(".zip")
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.write(output_path, arcname=output_path.name)
            archive.write(manifest_path, arcname=manifest_path.name)
            for segment_path in rendered_paths:
                archive.write(segment_path, arcname=f"segments/{segment_path.name}")
    except Exception as exc:
        status.value = f"Audio render failed: {exc}"
        ui.notify(status.value, type="negative")
        return

    full_audio_path.value = str(output_path)
    package_path.value = str(zip_path)
    audio_draft_record = save_draft(
        title=f"{document_name or file_stem} {render_mode} audio",
        content_type="podcast_audio",
        content=manifest_path.read_text(encoding="utf-8"),
        source_path=output_path,
        metadata={
            "mode": render_mode,
            "package_path": str(zip_path),
            "segment_count": len(rendered_paths),
            "model": model_id or ELEVENLABS_MODEL_OPTIONS[0],
        },
    )
    if saved_draft_path is not None:
        saved_draft_path.value = audio_draft_record["path"]
    audio_player.set_source(output_path)
    status.value = f"Rendered {render_mode} MP3 from {len(rendered_paths)} voice segment(s)."
    if progress is not None and progress_label is not None:
        await finish_generation_progress(
            progress=progress,
            label=progress_label,
            status=status,
            total_assets=len(rendered_paths),
            message=f"Rendered {render_mode} MP3 from {len(rendered_paths)} voice segment(s).",
        )
    ui.notify(status.value, type="positive")


def _split_markdown_list(value: str) -> list[str]:
    if not value:
        return []
    cleaned = re.sub(r"^[\s*-]+", "", value.strip(), flags=re.MULTILINE)
    parts = re.split(r",|\n", cleaned)
    return [part.strip(" -*`") for part in parts if part.strip(" -*`")]


def _markdown_field(block: str, field_name: str) -> str:
    pattern = rf"###\s+{re.escape(field_name)}\s*\n+(.*?)(?=\n###|\n---|\Z)"
    match = re.search(pattern, block, flags=re.DOTALL | re.IGNORECASE)
    return match.group(1).strip() if match else ""


def read_quote_bank_entries() -> list[dict[str, object]]:
    path = PROJECT_ROOT / "knowledge_base" / "primary" / "quote_bank.md"
    if not path.exists():
        return []

    text = path.read_text(encoding="utf-8")
    blocks = re.split(r"\n(?=##\s+Quote\b)", text)
    entries: list[dict[str, object]] = []
    for block in blocks:
        if "### Quote" not in block:
            continue
        heading = re.search(r"##\s+(Quote\s+[^\n]+)", block)
        entry_id = heading.group(1).strip() if heading else "Quote"
        quote_text = _markdown_field(block, "Quote")
        if not quote_text:
            continue
        entries.append(
            {
                "id": entry_id,
                "book": _markdown_field(block, "Book") or "Unknown book",
                "speaker": _markdown_field(block, "Speaker / Source") or "Unknown source",
                "characters": _split_markdown_list(_markdown_field(block, "Character Tags")),
                "moods": _split_markdown_list(_markdown_field(block, "Mood Tags")),
                "spoiler": _markdown_field(block, "Spoiler Level") or "Unspecified",
                "quote": quote_text.strip().strip('"'),
            }
        )
    return entries


def quote_book_options() -> list[str]:
    books = [entry["book"] for entry in read_quote_bank_entries() if entry.get("book")]
    ordered = list(dict.fromkeys([str(book) for book in books]))
    return ["All books", *ordered] if ordered else ["All books"]


def quote_entries_for_book(book_source: str) -> list[dict[str, object]]:
    entries = read_quote_bank_entries()
    if not book_source or book_source == "All books":
        return entries
    return [entry for entry in entries if str(entry.get("book", "")).lower() == str(book_source).lower()]


def quote_options_for_book(book_source: str) -> list[str]:
    return [entry["quote"] for entry in quote_entries_for_book(book_source)]


def quote_characters_for_book(book_source: str) -> list[str]:
    available: list[str] = []
    for entry in quote_entries_for_book(book_source):
        available.extend(entry.get("characters") or [])
    available_set = {name for name in available if name}
    ordered = [name for name in CHARACTER_TAGS if name in available_set]
    extras = sorted(available_set.difference(ordered))
    return ordered + extras or list(CHARACTER_TAGS)


def read_real_review_entries() -> list[dict[str, object]]:
    path = PROJECT_ROOT / "knowledge_base" / "secondary" / "real_reviews.md"
    if not path.exists():
        return []

    text = path.read_text(encoding="utf-8")
    blocks = re.split(r"\n(?=##\s+(?:Review Source|Reader Review Source)\b)", text)
    entries: list[dict[str, object]] = []
    for block in blocks:
        heading = re.search(r"##\s+(.+)", block)
        if not heading:
            continue
        review_id = heading.group(1).strip()
        source = _markdown_field(block, "Source") or "Unknown source"
        reviewer = _markdown_field(block, "Reviewer") or "Unknown reviewer"
        quotes_block = _markdown_field(block, "Real Quotes")
        quotes = [line.strip(" -*\"") for line in quotes_block.splitlines() if line.strip().startswith("-")]
        if not quotes:
            continue
        entries.append(
            {
                "id": review_id,
                "book": _markdown_field(block, "Book") or "Unknown book",
                "source": source,
                "reviewer": reviewer,
                "review_type": _markdown_field(block, "Review Type") or "Unspecified",
                "rating": _markdown_field(block, "Rating / Award") or _markdown_field(block, "Rating") or "",
                "date": _markdown_field(block, "Date") or "",
                "summary": _markdown_field(block, "Summary of Review") or "",
                "link": _markdown_field(block, "Source Link") or "",
                "quotes": quotes,
            }
        )
    return entries


def review_book_options() -> list[str]:
    books = [entry["book"] for entry in read_real_review_entries() if entry.get("book")]
    ordered = list(dict.fromkeys([str(book) for book in books]))
    # Always offer every catalog book so each one is selectable here, even before
    # it has verified reviews. The review source / quote sub-dropdowns stay data-
    # driven, so they simply stay empty for a book until real reviews are added.
    return list(dict.fromkeys([*QUOTE_BOOK_OPTIONS, *ordered]))


def review_entries_for_book(book_source: str) -> list[dict[str, object]]:
    entries = read_real_review_entries()
    if not book_source or book_source == "All books":
        return entries
    return [entry for entry in entries if str(entry.get("book", "")).lower() == str(book_source).lower()]


def review_source_options(book_source: str) -> list[str]:
    return [f"{entry['source']} — {entry['reviewer']}" for entry in review_entries_for_book(book_source)]


def review_quotes_for_source(book_source: str, review_source: str) -> list[str]:
    entries = review_entries_for_book(book_source)
    if review_source:
        entries = [entry for entry in entries if f"{entry['source']} — {entry['reviewer']}" == review_source]
    quotes: list[str] = []
    for entry in entries:
        quotes.extend(entry.get("quotes") or [])
    return list(dict.fromkeys(quotes))


def build_quote_preview_html(content_type: str, fields: dict[str, object]) -> str:
    if content_type == "quote_post":
        quote_text = normalize_selected(fields.get("quote_text"))
        source = normalize_selected(fields.get("quote_source"))
        moods = normalize_selected(fields.get("quote_moods"))
        characters = normalize_selected(fields.get("quote_characters"))
        spoiler = normalize_selected(fields.get("quote_spoiler"))
        request = normalize_selected(fields.get("quote_request"))
        return f"""
        <div style="border-radius:20px;padding:18px;background:linear-gradient(180deg,#fff,#fbf6f7);border:1px solid rgba(142,31,47,.14);box-shadow:0 18px 36px rgba(61,31,41,.08);">
            <div style="text-transform:uppercase;letter-spacing:.18em;font-size:11px;color:#8E1F2F;margin-bottom:12px;">Quote Post Preview</div>
            <div style="font-size:18px;line-height:1.55;font-weight:700;margin-bottom:12px;white-space:pre-wrap;">“{quote_text}”</div>
            <div style="font-size:13px;color:#6b6b6b;margin-bottom:6px;"><strong>Source:</strong> {source}</div>
            <div style="font-size:13px;color:#6b6b6b;margin-bottom:6px;"><strong>Moods:</strong> {moods}</div>
            <div style="font-size:13px;color:#6b6b6b;margin-bottom:6px;"><strong>Characters:</strong> {characters}</div>
            <div style="font-size:13px;color:#6b6b6b;margin-bottom:6px;"><strong>Spoiler:</strong> {spoiler}</div>
            <div style="font-size:13px;color:#6b6b6b;"><strong>Request:</strong> {request}</div>
        </div>
        """

    if content_type == "review_pull_quote":
        mode = normalize_selected(fields.get("review_mode"))
        quote_text = normalize_selected(fields.get("review_quote_text"))
        source = normalize_selected(fields.get("review_source"))
        attribution = normalize_selected(fields.get("review_attribution"))
        promo_line = normalize_selected(fields.get("review_promo_line"))
        return f"""
        <div style="border-radius:20px;padding:18px;background:linear-gradient(180deg,#fff,#f5f7fb);border:1px solid rgba(17,17,20,.12);box-shadow:0 18px 36px rgba(61,31,41,.08);">
            <div style="text-transform:uppercase;letter-spacing:.18em;font-size:11px;color:#8E1F2F;margin-bottom:12px;">Review Pull Quote Preview</div>
            <div style="font-size:18px;line-height:1.55;font-weight:700;margin-bottom:12px;white-space:pre-wrap;">{quote_text}</div>
            <div style="font-size:13px;color:#6b6b6b;margin-bottom:6px;"><strong>Mode:</strong> {mode}</div>
            <div style="font-size:13px;color:#6b6b6b;margin-bottom:6px;"><strong>Source:</strong> {source}</div>
            <div style="font-size:13px;color:#6b6b6b;margin-bottom:6px;"><strong>Attribution:</strong> {attribution}</div>
            <div style="font-size:13px;color:#6b6b6b;"><strong>Brand promo line:</strong> {promo_line}</div>
        </div>
        """

    return """
    <div style="border-radius:20px;padding:20px;background:#ffffff;border:1px dashed rgba(17,17,20,.16);color:#6b6b6b;">
        Select a quote workflow to preview grounded quote selections.
    </div>
    """


# Per-format aspect ratio + badge for the Instagram Output Studio mock cards.
IG_FORMAT_RENDER = {
    "feed post": ("4 / 5", "Post · 4:5"),
    "reel": ("9 / 16", "Reel · 9:16"),
    "story": ("9 / 16", "Story · 9:16"),
    "carousel": ("4 / 5", "Carousel · 4:5"),
}


def build_instagram_format_card(fmt: str, block: dict, image_path: str = "", format_images=None) -> str:
    """Render one Instagram mock card for a single post format at its true aspect ratio.

    ``format_images`` is the list of generated image paths for this format (used by the
    per-format Regenerate flow); when empty the uploaded ``image_path`` is shown.
    """
    aspect, badge = IG_FORMAT_RENDER.get(fmt.lower(), IG_FORMAT_RENDER["feed post"])
    images = [str(p) for p in (format_images or []) if str(p or "").strip()]
    primary_url = image_data_url(images[0]) if images else image_data_url(str(image_path or ""))

    def fill(url_value: str, label: str = "IMAGE") -> str:
        if url_value:
            return f'<img src="{escape(url_value, quote=True)}" alt="" style="width:100%;height:100%;object-fit:cover;display:block;">'
        return (
            f'<div style="width:100%;height:100%;display:flex;align-items:center;justify-content:center;'
            f'background:linear-gradient(135deg,#1c1418,#6f182c);color:#f8efe8;font-weight:800;letter-spacing:.06em;font-size:12px;">{escape(label)}</div>'
        )

    def image_fill(label: str = "IMAGE") -> str:
        return fill(primary_url, label)

    header = (
        '<div style="padding:9px 12px;display:flex;align-items:center;gap:8px;">'
        '<div style="width:28px;height:28px;border-radius:50%;background:linear-gradient(135deg,#2a0712,#8E1F2F);color:#fff;display:flex;align-items:center;justify-content:center;font-family:Georgia,serif;font-size:15px;">M</div>'
        '<div style="flex:1;font-weight:800;font-size:12px;">mortalvengeance</div>'
        f'<div style="font-size:9px;font-weight:800;color:#8E1F2F;border:1px solid rgba(142,31,47,.3);border-radius:999px;padding:2px 7px;white-space:nowrap;">{escape(badge)}</div>'
        '</div>'
    )
    key = fmt.lower()

    if key == "story":
        overlay = normalize_selected(block.get("overlay text"))
        cta = normalize_selected(block.get("cta"))
        overlay_html = (
            '<div style="position:absolute;inset:0;display:flex;align-items:center;justify-content:center;text-align:center;padding:18px;">'
            f'<div style="font-family:Georgia,serif;font-weight:900;font-size:24px;line-height:1.1;color:#fff;text-shadow:0 2px 14px rgba(0,0,0,.6);">{escape(overlay)}</div>'
            '</div>'
            if overlay != "Not specified" else ""
        )
        cta_html = (
            f'<div style="position:absolute;left:0;right:0;bottom:14px;text-align:center;color:#fff;font-size:12px;font-weight:700;text-shadow:0 1px 8px rgba(0,0,0,.6);">{escape(cta)}</div>'
            if cta != "Not specified" else ""
        )
        body_area = f'<div style="position:relative;aspect-ratio:{aspect};background:#160f13;">{image_fill("STORY")}{overlay_html}{cta_html}</div>'
        return f'<div style="width:240px;border-radius:18px;overflow:hidden;background:#fff;border:1px solid rgba(0,0,0,.1);box-shadow:0 14px 30px rgba(61,31,41,.1);">{header}{body_area}</div>'

    if key == "carousel":
        slides = block.get("slides") or []
        raw_caption = str(block.get("caption") or (slides[0] if slides else "")).strip()
        caption = markdown_to_html(raw_caption) if raw_caption else ""
        hashtags = escape(normalize_selected(block.get("hashtags"))) if block.get("hashtags") else ""
        total = len(slides)
        if slides:
            slide_cards = "".join(
                (
                    f'<div style="flex:0 0 auto;width:150px;aspect-ratio:{aspect};border-radius:12px;overflow:hidden;position:relative;background:#160f13;border:1px solid rgba(0,0,0,.1);">'
                    f'{fill(image_data_url(images[i - 1]) if i - 1 < len(images) else primary_url)}'
                    f'<div style="position:absolute;top:6px;right:6px;background:rgba(0,0,0,.7);color:#fff;border-radius:999px;font-size:10px;padding:1px 7px;font-weight:800;">{i}/{total}</div>'
                    f'<div style="position:absolute;inset:auto 0 0 0;padding:8px;background:linear-gradient(transparent,rgba(0,0,0,.8));color:#fff;font-size:11px;line-height:1.3;">{escape(str(slide_text)[:140])}</div>'
                    '</div>'
                )
                for i, slide_text in enumerate(slides, 1)
            )
        else:
            slide_cards = f'<div style="flex:0 0 auto;width:150px;aspect-ratio:{aspect};border-radius:12px;overflow:hidden;background:#160f13;">{image_fill("SLIDES")}</div>'
        deck = f'<div style="display:flex;gap:10px;overflow-x:auto;padding:12px;">{slide_cards}</div>'
        caption_area = (
            '<div style="padding:6px 14px 14px;">'
            + (
                f'<div style="font-size:13px;line-height:1.45;white-space:pre-wrap;"><strong>mortalvengeance</strong> {caption}</div>'
                if caption
                else '<div style="font-size:12px;color:#9a8f86;font-style:italic;">Caption preview unavailable for this asset.</div>'
            )
            + (f'<div style="font-size:12px;color:#2454a6;margin-top:6px;">{hashtags}</div>' if hashtags else '')
            + '</div>'
        )
        return f'<div style="width:340px;border-radius:18px;overflow:hidden;background:#fff;border:1px solid rgba(0,0,0,.1);box-shadow:0 14px 30px rgba(61,31,41,.1);">{header}{deck}{caption_area}</div>'

    # Feed post / Reel / default.
    caption = markdown_to_html(block.get("caption")) if block.get("caption") else "Generate to see the caption."
    hashtags = escape(normalize_selected(block.get("hashtags"))) if block.get("hashtags") else ""
    label = "REEL COVER" if key == "reel" else "IMAGE"
    body_area = f'<div style="aspect-ratio:{aspect};background:#160f13;">{image_fill(label)}</div>'
    caption_area = (
        '<div style="padding:10px 14px 14px;">'
        '<div style="display:flex;justify-content:space-between;font-size:11px;margin-bottom:8px;color:#333;font-weight:800;"><span>Like · Comment · Share</span><span>Save</span></div>'
        f'<div style="font-size:13px;line-height:1.45;white-space:pre-wrap;"><strong>mortalvengeance</strong> {caption}</div>'
        f'<div style="font-size:12px;color:#2454a6;margin-top:6px;">{hashtags}</div>'
        '</div>'
    )
    width = "240px" if key == "reel" else "300px"
    return f'<div style="width:{width};border-radius:18px;overflow:hidden;background:#fff;border:1px solid rgba(0,0,0,.1);box-shadow:0 14px 30px rgba(61,31,41,.1);">{header}{body_area}{caption_area}</div>'


# Card palettes for review_pull_quote (BUG-RPQ-03).
REVIEW_CARD_STYLES = {
    "dark gothic": {"bg": "linear-gradient(160deg,#141019,#2a1622)", "fg": "#f6efe9", "muted": "#d7a8b3", "accent": "#e7c9cf"},
    "light minimal": {"bg": "linear-gradient(180deg,#fffdfb,#f4f1ec)", "fg": "#1a1417", "muted": "#8E1F2F", "accent": "#6b6b6b"},
    "brand default": {"bg": "linear-gradient(160deg,#2a0712,#8E1F2F)", "fg": "#fff6f3", "muted": "#f0c9cf", "accent": "#f3d9dd"},
}


def build_review_quote_card(fmt, style, image_type, quote, attribution, promo, image_path="") -> str:
    """Render one styled pull-quote card for a given Instagram format (BUG-RPQ-03)."""
    aspect, badge = IG_FORMAT_RENDER.get(str(fmt).lower(), IG_FORMAT_RENDER["feed post"])
    pal = REVIEW_CARD_STYLES.get(str(style or "").strip().lower(), REVIEW_CARD_STYLES["dark gothic"])
    image_mode = str(image_type or "").lower()
    use_image = "overlay" in image_mode  # character or book-cover overlay
    url = image_data_url(str(image_path or ""))

    if use_image and url:
        bg_layer = (
            f'<img src="{escape(url, quote=True)}" alt="" style="position:absolute;inset:0;width:100%;height:100%;object-fit:cover;">'
            '<div style="position:absolute;inset:0;background:linear-gradient(180deg,rgba(8,5,8,.35),rgba(8,5,8,.8));"></div>'
        )
        text_color, muted, accent = "#ffffff", "#e7c9cf", "#f0d4d9"
    elif use_image:
        bg_layer = (
            '<div style="position:absolute;inset:0;background:linear-gradient(135deg,#1c1418,#6f182c);"></div>'
            '<div style="position:absolute;inset:0;background:linear-gradient(180deg,rgba(8,5,8,.18),rgba(8,5,8,.72));"></div>'
            '<div style="position:absolute;top:10px;left:12px;font-size:9px;letter-spacing:.14em;text-transform:uppercase;color:rgba(255,255,255,.6);">'
            f'{escape("book cover" if "cover" in image_mode else "character")} image</div>'
        )
        text_color, muted, accent = "#ffffff", "#e7c9cf", "#f0d4d9"
    else:
        bg_layer = f'<div style="position:absolute;inset:0;background:{pal["bg"]};"></div>'
        text_color, muted, accent = pal["fg"], pal["muted"], pal["accent"]

    quote_html = markdown_to_html(str(quote or "")[:400])
    attribution_html = escape(attribution) if attribution and attribution != "Not specified" else "Verified reader review"
    promo_html = (
        f'<div style="position:relative;margin-top:14px;font-size:11px;letter-spacing:.16em;text-transform:uppercase;color:{muted};font-weight:800;">{escape(promo)}</div>'
        if promo and promo != "Not specified" else ""
    )
    badge_html = (
        f'<div style="position:absolute;top:12px;right:12px;font-size:9px;font-weight:800;color:{text_color};'
        f'border:1px solid rgba(255,255,255,.35);border-radius:999px;padding:2px 8px;background:rgba(0,0,0,.18);">{escape(badge)}</div>'
    )
    width = "260px" if aspect == "9 / 16" else "300px"
    return (
        f'<div style="width:{width};border-radius:20px;overflow:hidden;box-shadow:0 18px 40px rgba(20,10,16,.28);">'
        f'<div style="position:relative;aspect-ratio:{aspect};display:flex;flex-direction:column;justify-content:center;padding:26px 24px;">'
        f'{bg_layer}{badge_html}'
        f'<div style="position:absolute;top:-8px;left:14px;font-family:Georgia,serif;font-size:90px;line-height:1;color:rgba(255,255,255,.10);">&ldquo;</div>'
        f'<blockquote style="position:relative;margin:0;font-family:Georgia,serif;font-weight:600;font-size:clamp(17px,3.4vw,23px);line-height:1.35;color:{text_color};white-space:pre-wrap;">{quote_html}</blockquote>'
        f'<div style="position:relative;margin-top:16px;font-size:13px;color:{accent};">&mdash; {attribution_html}</div>'
        f'{promo_html}'
        '</div></div>'
    )


def build_preview_html(content_type: str, fields: dict[str, object]) -> str:
    def text(value) -> str:
        return normalize_selected(value)

    def safe(value) -> str:
        return escape(text(value))

    def body(value) -> str:
        return markdown_to_html(value)

    def image_markup(image_path, fallback_label: str) -> str:
        url = image_data_url(str(image_path or ""))
        if url:
            return f'<img src="{escape(url, quote=True)}" alt="Uploaded creative" style="width:100%;height:100%;object-fit:cover;display:block;">'
        return f"""
        <div style="width:100%;height:100%;display:flex;align-items:center;justify-content:center;background:linear-gradient(135deg,#1c1418,#6f182c);color:#f8efe8;font-weight:800;letter-spacing:.08em;text-align:center;padding:20px;">
            {escape(fallback_label)}
        </div>
        """

    def first_line(value, fallback: str) -> str:
        raw = text(value)
        if raw == "Not specified":
            return fallback
        for line in raw.splitlines():
            cleaned = line.strip().strip("#*- ")
            if cleaned:
                return cleaned[:110]
        return fallback

    if content_type == "quote_post":
        image_path = str(fields.get("image_path") or "")
        if image_path and Path(image_path).exists():
            return quote_platform_mockup_html([image_path], str(fields.get("generated_visual_package_path") or ""))
        candidates = extract_quote_candidates(str(fields.get("draft") or ""))
        quote_text = candidates[0] if candidates else first_line(fields.get("draft"), "Select quote cards to generate images.")
        return build_quote_preview_html(
            "quote_post",
            {
                "quote_text": quote_text,
                "quote_source": fields.get("quote_source") or "Generated quote draft",
                "quote_moods": fields.get("quote_moods"),
                "quote_characters": fields.get("quote_characters"),
                "quote_spoiler": "Grounded selection",
                "quote_request": "Use the Quote Image Builder below the draft to select cards and render platform mockups.",
            },
        )

    if content_type == "instagram_caption":
        selected = [str(f).strip() for f in (fields.get("instagram_formats") or []) if str(f).strip()] or ["Feed post"]
        parsed = parse_instagram_multiformat(text(fields.get("draft")))
        image_path = str(fields.get("image_path") or "")
        format_images = fields.get("instagram_format_images") or {}
        cards = "".join(
            build_instagram_format_card(fmt, parsed.get(fmt.lower(), {}), image_path, format_images.get(fmt.lower()))
            for fmt in selected
        )
        return f'<div style="display:flex;flex-wrap:wrap;gap:18px;justify-content:center;align-items:flex-start;">{cards}</div>'

    if content_type == "linkedin_content":
        formats = safe(fields.get("linkedin_formats"))
        angle = safe(fields.get("linkedin_angle"))
        cta = safe(fields.get("linkedin_cta"))
        scheduled = (escape(str(fields.get("post_date"))) + " · Public") if str(fields.get("post_date") or "").strip() else "Preview · Public"
        # Split the draft into per-format sections so each LinkedIn deliverable is a
        # labeled block (badge) — and render document/carousel formats as a slide deck.
        segments = parse_headed_sections(str(fields.get("draft") or ""), _LINKEDIN_FORMAT_NAMES)
        blocks_html = ""
        for seg_label, seg_body in segments:
            if seg_label:
                blocks_html += (
                    f'<div style="display:inline-block;font-size:11px;font-weight:800;color:#0a66c2;'
                    f'border:1px solid rgba(10,102,194,.3);border-radius:999px;padding:2px 10px;margin:14px 0 6px;">{escape(seg_label)}</div>'
                )
            is_deck = ("carousel" in seg_label.lower() or "document" in seg_label.lower()
                       or bool(re.search(r"(?im)^\s*slide\s*\d", seg_body)))
            seg_slides = parse_linkedin_slides(seg_body) if is_deck else []
            if seg_slides:
                slide_cards = "".join(
                    '<div style="flex:0 0 auto;width:210px;aspect-ratio:1.91/1;border-radius:10px;'
                    'background:linear-gradient(160deg,#0a2540,#0a66c2);color:#fff;padding:13px;'
                    'border:1px solid rgba(0,0,0,.1);overflow:hidden;">'
                    f'<div style="font-size:10px;font-weight:800;opacity:.75;letter-spacing:.06em;">SLIDE {i} / {len(seg_slides)}</div>'
                    f'<div style="font-size:12px;line-height:1.4;margin-top:7px;overflow:hidden;">{markdown_to_html(str(slide_text)[:240])}</div>'
                    '</div>'
                    for i, slide_text in enumerate(seg_slides, 1)
                )
                blocks_html += f'<div style="display:flex;gap:10px;overflow-x:auto;padding:4px 0 8px;">{slide_cards}</div>'
            elif seg_body.strip():
                blocks_html += f'<div style="font-size:14px;line-height:1.5;white-space:pre-wrap;margin-bottom:4px;">{markdown_to_html(seg_body)}</div>'
        return f"""
        <div style="max-width:520px;margin:0 auto;border-radius:16px;background:#fff;border:1px solid rgba(0,0,0,.12);box-shadow:0 18px 40px rgba(61,31,41,.08);overflow:hidden;">
            <div style="padding:16px;display:flex;gap:12px;align-items:center;">
                <div style="width:48px;height:48px;border-radius:4px;background:#0a66c2;color:white;display:flex;align-items:center;justify-content:center;font-weight:900;font-size:22px;">in</div>
                <div style="flex:1;">
                    <div style="font-weight:800;font-size:14px;">Alejandro Torres De La Rocha</div>
                    <div style="font-size:12px;color:#666;">Author · Mortal Vengeance</div>
                    <div style="font-size:11px;color:#777;">{scheduled}</div>
                </div>
            </div>
            <div style="padding:0 16px 14px;">{blocks_html}</div>
            <div style="aspect-ratio:1.91/1;background:#111;">{image_markup(fields.get("image_path"), "LINKEDIN VISUAL")}</div>
            <div style="padding:12px 16px;border-top:1px solid #eee;color:#666;font-size:12px;">
                <strong>Angle:</strong> {angle} · <strong>CTA:</strong> {cta} · <strong>Formats:</strong> {formats}
            </div>
            <div style="display:flex;justify-content:space-around;border-top:1px solid #eee;padding:10px 0;color:#555;font-size:13px;font-weight:700;">
                <span>Like</span><span>Comment</span><span>Repost</span><span>Send</span>
            </div>
        </div>
        """

    if content_type == "youtube_content":
        formats = safe(fields.get("youtube_formats"))
        keywords = safe(fields.get("youtube_keywords"))
        link = safe(fields.get("youtube_link"))
        scheduled = escape(str(fields.get("post_date"))) if str(fields.get("post_date") or "").strip() else "Preview"
        # Split the draft into per-deliverable sections (Title / Description / Pinned
        # comment / Long-form video / ...) so each is a labeled block, not one blob.
        segments = parse_headed_sections(str(fields.get("draft") or ""), _YOUTUBE_DELIVERABLE_NAMES)
        title_seg = next((seg_body for (seg_label, seg_body) in segments if seg_label.lower() == "title"), "")
        title = escape(first_line(title_seg or fields.get("draft"), "Mortal Vengeance: A Grim Tale"))
        blocks_html = ""
        for seg_label, seg_body in segments:
            if seg_label:
                blocks_html += (
                    '<div style="display:inline-block;font-size:11px;font-weight:800;color:#8E1F2F;'
                    'border:1px solid rgba(142,31,47,.3);border-radius:999px;padding:2px 10px;margin:12px 0 6px;">'
                    f'{escape(seg_label)}</div>'
                )
            if seg_body.strip():
                blocks_html += f'<div style="font-size:13px;line-height:1.45;white-space:pre-wrap;color:#333;margin-bottom:4px;">{markdown_to_html(seg_body)}</div>'
        return f"""
        <div style="max-width:560px;margin:0 auto;border-radius:18px;background:#fff;border:1px solid rgba(0,0,0,.12);box-shadow:0 18px 40px rgba(61,31,41,.08);overflow:hidden;">
            <div style="aspect-ratio:16/9;background:#111;position:relative;">{image_markup(fields.get("image_path"), "YOUTUBE THUMBNAIL")}</div>
            <div style="padding:14px 16px;display:flex;gap:12px;">
                <div style="width:42px;height:42px;border-radius:50%;background:#8E1F2F;color:#fff;display:flex;align-items:center;justify-content:center;font-family:Georgia,serif;font-size:23px;">M</div>
                <div style="flex:1;">
                    <div style="font-size:16px;font-weight:800;line-height:1.3;margin-bottom:5px;">{title}</div>
                    <div style="font-size:12px;color:#666;margin-bottom:10px;">Mortal Vengeance · {scheduled}</div>
                    {blocks_html}
                    <div style="font-size:12px;color:#666;margin-top:10px;padding-top:8px;border-top:1px solid #eee;"><strong>Formats:</strong> {formats} · <strong>Keywords:</strong> {keywords} · <strong>Link:</strong> {link}</div>
                </div>
            </div>
        </div>
        """

    if content_type == "press_release":
        draft_raw = text(fields.get("draft") or fields.get("topic"))
        lines = [line.strip() for line in draft_raw.splitlines() if line.strip()]
        headline = escape(first_line(draft_raw, "Press Release Headline"))
        body_source = "\n\n".join(lines[1:8]) if len(lines) > 1 else draft_raw
        body_html = body(body_source[:1800])
        dateline_parts = [
            text(fields.get("pr_city")) if text(fields.get("pr_city")) != "Not specified" else "",
            text(fields.get("pr_state")) if text(fields.get("pr_state")) != "Not specified" else "",
        ]
        dateline = safe(", ".join(part for part in dateline_parts if part) or "Dateline")
        release_date = safe(fields.get("pr_release_date"))
        contact = safe(fields.get("pr_contact_name"))
        website = safe(fields.get("pr_website"))
        news_angle = safe(fields.get("pr_news_angle"))
        assets = safe(fields.get("pr_required_assets"))
        return f"""
        <div style="max-width:720px;margin:0 auto;background:#f7f3ea;border:1px solid rgba(44,24,31,.18);box-shadow:0 20px 42px rgba(61,31,41,.14);padding:18px;">
            <div style="border-top:5px solid #1e1b19;border-bottom:2px solid #1e1b19;text-align:center;padding:10px 0 8px;margin-bottom:14px;">
                <div style="font-family:Georgia,serif;font-size:42px;line-height:1;font-weight:900;letter-spacing:.02em;color:#1d1a18;">THE TELL TALES INK PRESS</div>
                <div style="display:flex;justify-content:space-between;font-size:11px;text-transform:uppercase;letter-spacing:.12em;color:#5c524b;margin-top:8px;">
                    <span>FOR IMMEDIATE RELEASE</span><span>{release_date}</span><span>{PRESS_RELEASE_DESTINATION}</span>
                </div>
            </div>
            <div style="display:grid;grid-template-columns:2fr 1fr;gap:18px;">
                <div>
                    <div style="font-size:11px;text-transform:uppercase;letter-spacing:.16em;color:#8E1F2F;font-weight:900;margin-bottom:8px;">{dateline}</div>
                    <h1 style="font-family:Georgia,serif;font-size:38px;line-height:1.02;margin:0 0 10px;color:#17120f;">{headline}</h1>
                    <div style="font-size:15px;line-height:1.58;white-space:pre-wrap;color:#26211e;column-count:2;column-gap:20px;">{body_html}</div>
                </div>
                <aside style="border-left:1px solid rgba(44,24,31,.22);padding-left:14px;">
                    <div style="font-size:12px;text-transform:uppercase;letter-spacing:.14em;font-weight:900;border-bottom:1px solid rgba(44,24,31,.2);padding-bottom:6px;margin-bottom:10px;">Newswire Notes</div>
                    <div style="font-size:13px;line-height:1.45;color:#322b27;"><strong>News angle</strong><br>{news_angle}</div>
                    <div style="font-size:13px;line-height:1.45;color:#322b27;margin-top:12px;"><strong>Media contact</strong><br>{contact}<br>{website}</div>
                    <div style="font-size:13px;line-height:1.45;color:#322b27;margin-top:12px;"><strong>Assets</strong><br>{assets}</div>
                </aside>
            </div>
        </div>
        """

    if content_type == "blog_post":
        blog_format = safe(fields.get("blog_format"))
        blog_length = safe(fields.get("blog_length"))
        structure = safe(fields.get("blog_structure_options") or fields.get("blog_structure_checklist") or fields.get("blog_sections") or fields.get("blog_structure"))
        seo = safe(fields.get("blog_seo_keywords"))
        image_mode = safe(fields.get("blog_image_mode"))
        blog_character = safe(fields.get("blog_character"))
        visual_style = safe(fields.get("blog_visual_style"))
        visual_formats = safe(fields.get("blog_visual_formats"))
        draft = body(text(fields.get("draft") or fields.get("topic"))[:700])
        title = escape(first_line(fields.get("draft") or fields.get("topic"), "Mortal Vengeance Feature"))
        return f"""
        <article style="max-width:680px;margin:0 auto;border-radius:22px;background:#fffdfb;border:1px solid rgba(44,24,31,.13);box-shadow:0 18px 40px rgba(61,31,41,.09);overflow:hidden;">
            <div style="aspect-ratio:16/9;background:#120d10;">{image_markup(fields.get("image_path"), "BLOG VISUAL")}</div>
            <div style="padding:24px;">
                <div style="text-transform:uppercase;letter-spacing:.16em;font-size:11px;color:#8E1F2F;font-weight:900;margin-bottom:10px;">{blog_format}</div>
                <h2 style="font-family:Georgia,serif;font-size:34px;line-height:1.08;margin:0 0 12px;color:#171217;">{title}</h2>
                <div style="font-size:13px;color:#71686a;margin-bottom:14px;"><strong>Length:</strong> {blog_length} · <strong>Image:</strong> {image_mode} · <strong>Character:</strong> {blog_character}</div>
                <div style="font-size:14px;line-height:1.58;white-space:pre-wrap;color:#2f292b;margin-bottom:14px;">{draft}</div>
                <div style="font-size:12px;color:#71686a;"><strong>Structure:</strong> {structure}</div>
                <div style="font-size:12px;color:#71686a;margin-top:6px;"><strong>SEO:</strong> {seo}</div>
                <div style="font-size:12px;color:#71686a;margin-top:6px;"><strong>Visuals:</strong> {visual_style} · {visual_formats}</div>
            </div>
        </article>
        """

    if content_type == "character_spotlight":
        subject = text(fields.get("character_subject") or fields.get("cs_character"))
        if subject == "Not specified":
            return """
            <div style="border-radius:20px;padding:30px 24px;background:#faf7f3;border:1px dashed rgba(44,24,31,.22);color:#8a8079;text-align:center;font-size:14px;">
                Select a character above to preview the spotlight.
            </div>
            """
        focus = safe(fields.get("character_focus") or fields.get("cs_focus"))
        format_label = text(fields.get("character_format") or fields.get("cs_platform_format"))
        image_mode = text(fields.get("character_image_mode") or fields.get("cs_image_mode"))
        draft_text = text(fields.get("draft"))
        portrait = resolve_character_portrait_asset(subject, book=fields.get("related_book")) if image_mode != "use uploaded image" else None
        portrait_url = image_data_url(portrait) if portrait else ""
        if image_mode == "text-only highlight image":
            portrait_markup = image_markup("", escape(subject).upper())
        elif portrait_url:
            portrait_markup = f'<img src="{escape(portrait_url, quote=True)}" alt="{escape(subject, quote=True)} portrait" style="width:100%;height:100%;object-fit:cover;display:block;">'
        else:
            portrait_markup = image_markup(fields.get("image_path"), "CHARACTER PORTRAIT")

        lookup = format_label.lower()
        if "story" in lookup or "reel" in lookup:
            container_style = "max-width:360px;aspect-ratio:9/16"
            layout = "display:flex;flex-direction:column"
            image_style = "height:58%;"
            text_style = "padding:18px;"
        elif "youtube" in lookup or "horizontal" in lookup or "blog cover" in lookup:
            container_style = "max-width:620px;aspect-ratio:16/9"
            layout = "display:grid;grid-template-columns:46% 1fr"
            image_style = "height:100%;"
            text_style = "padding:20px;"
        else:
            container_style = "max-width:500px;aspect-ratio:1/1"
            layout = "display:grid;grid-template-columns:46% 1fr"
            image_style = "height:100%;"
            text_style = "padding:20px;"

        max_copy = 520 if "blog" in lookup else 240
        display_draft = body(draft_text[:max_copy]) + ("..." if len(draft_text) > max_copy else "")
        return f"""
        <div style="{container_style};margin:0 auto;border-radius:22px;background:#fffdfb;border:1px solid rgba(44,24,31,.13);box-shadow:0 18px 40px rgba(61,31,41,.09);overflow:hidden;">
            <div style="{layout};width:100%;height:100%;">
                <div style="background:#090908;{image_style}">{portrait_markup}</div>
                <div style="{text_style}display:flex;flex-direction:column;gap:10px;overflow:hidden;">
                    <div style="text-transform:uppercase;letter-spacing:.18em;font-size:11px;color:#8E1F2F;font-weight:800;">Character Spotlight</div>
                    <div style="font-family:Georgia,serif;font-size:clamp(24px,4vw,34px);line-height:1.05;font-weight:900;color:#171217;">{escape(subject)}</div>
                    <div style="font-size:12px;color:#71686a;"><strong>{escape(format_label)}</strong> · {escape(image_mode)}</div>
                    <div style="font-size:12px;color:#71686a;"><strong>Focus:</strong> {focus}</div>
                    <div style="font-size:14px;line-height:1.45;white-space:pre-wrap;color:#2f292b;">{display_draft}</div>
                </div>
            </div>
        </div>
        """

    if content_type == "newsletter_blurb":
        subject = text(fields.get("newsletter_subject"))
        subject_display = escape(subject) if subject != "Not specified" else "Your subject line appears here"
        preview_text = text(fields.get("newsletter_preview"))
        preview_display = escape(preview_text) if preview_text != "Not specified" else "Preview text shown beside the subject in the inbox"
        structure = safe(fields.get("newsletter_structure"))
        newsletter_body = body(text(fields.get("draft") or fields.get("topic"))[:1200])
        sender = "Mortal Vengeance"
        return f"""
        <div style="max-width:560px;margin:0 auto;border-radius:16px;background:#fff;border:1px solid rgba(0,0,0,.12);box-shadow:0 18px 40px rgba(61,31,41,.08);overflow:hidden;font-family:Helvetica,Arial,sans-serif;">
            <div style="background:#f1f3f4;padding:10px 16px;font-size:12px;color:#5f6368;border-bottom:1px solid #e0e0e0;">Inbox</div>
            <div style="padding:14px 16px;border-bottom:1px solid #eee;display:flex;gap:12px;align-items:flex-start;">
                <div style="width:40px;height:40px;border-radius:50%;background:linear-gradient(135deg,#2a0712,#8E1F2F);color:#fff;display:flex;align-items:center;justify-content:center;font-family:Georgia,serif;font-size:20px;flex:none;">M</div>
                <div style="flex:1;min-width:0;">
                    <div style="display:flex;justify-content:space-between;font-size:13px;"><strong>{escape(sender)}</strong><span style="color:#5f6368;font-size:12px;">now</span></div>
                    <div style="font-size:14px;font-weight:700;color:#202124;margin-top:2px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;">{subject_display}</div>
                    <div style="font-size:13px;color:#5f6368;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;">{preview_display}</div>
                </div>
            </div>
            <div style="padding:20px 22px;">
                <div style="text-transform:uppercase;letter-spacing:.14em;font-size:11px;color:#8E1F2F;font-weight:800;margin-bottom:10px;">{structure}</div>
                <h2 style="font-family:Georgia,serif;font-size:24px;line-height:1.2;margin:0 0 14px;color:#1a1417;">{subject_display}</h2>
                <div style="font-size:14px;line-height:1.6;white-space:pre-wrap;color:#2f292b;">{newsletter_body}</div>
            </div>
        </div>
        """

    if content_type == "review_pull_quote":
        quote_raw = text(fields.get("review_quote_text") or fields.get("draft"))
        if quote_raw == "Not specified":
            quote_raw = "Select or paste a quote to preview the pull-quote card."
        # If the resolved text is a full generated draft, preview the first quote in it.
        elif len(quote_raw) > 220 or quote_raw.strip().count("\n") >= 2:
            candidates = extract_quote_candidates(quote_raw)
            if candidates:
                quote_raw = candidates[0]
        selected = [str(f).strip() for f in (fields.get("review_instagram_formats") or []) if str(f).strip()] or ["Feed post"]
        style = text(fields.get("review_card_style"))
        image_type = text(fields.get("review_image_type"))
        source = text(fields.get("review_attribution") or fields.get("review_source"))
        attribution = source if source != "Not specified" else ""
        promo = text(fields.get("review_promo_line"))
        image_path = str(fields.get("image_path") or "")
        cards = "".join(
            build_review_quote_card(fmt, style, image_type, quote_raw, attribution, promo, image_path)
            for fmt in selected
        )
        return f'<div style="display:flex;flex-wrap:wrap;gap:18px;justify-content:center;align-items:flex-start;">{cards}</div>'

    return f"""
    <div style="border-radius:20px;padding:20px;background:#ffffff;border:1px dashed rgba(17,17,20,.16);color:#6b6b6b;">
        Select a content type to see a platform-style preview.
    </div>
    """


_REFUSAL_RE = re.compile(
    r"(?i)\b(i'?ll wait|i will wait|until you tell me|once you (?:confirm|tell me|say)|"
    r"received[,. ].*(?:wait|generate)|as an ai|i (?:cannot|can'?t|am unable to) (?:help|generate|create|produce)|"
    r"i'?m (?:sorry|unable|happy to help|ready)|let me know when|just say the word|ready to generate|"
    r"waiting for your|tell me to (?:proceed|generate|start)|provide (?:the|more) (?:details|caption))\b"
)


def looks_like_refusal(value: str) -> bool:
    """True when a string looks like an LLM acknowledgement/refusal rather than real
    content — so it never gets selected as quote text and painted onto an image."""
    return bool(_REFUSAL_RE.search(str(value or "")))


def extract_quote_for_graphic(content: str) -> str:
    text = str(content or "").strip()
    if not text or looks_like_refusal(text):
        return ""
    # Prefer an explicit on-image text line from the structured output.
    label_match = re.search(
        r"(?im)^\s*(?:overlay text|image text suggestion|quote|pull quote|review quote|caption text|graphic text)\s*:\s*(.+)$",
        text,
    )
    if label_match:
        candidate = label_match.group(1).strip().strip('"“”')
        if candidate and not candidate.startswith("<"):
            return candidate[:260]
    curly_match = re.search(r"[“\"]([^”\"]{18,260})[”\"]", text)
    if curly_match:
        return curly_match.group(1).strip()
    # Line scan: skip "=== HEADER ===" scaffolding and strip field-label prefixes
    # so the overlay never shows "=== FEED POST ===" or "Caption:" (BUG-IMG-03).
    label_prefix = re.compile(
        r"(?i)^(?:caption|hashtags|image direction|overlay text|cta|attribution|hook|slide\s*\d*[^:]*)\s*:\s*"
    )
    for line in text.splitlines():
        cleaned = line.strip()
        if re.match(r"^={2,}.*={2,}$", cleaned):
            continue
        cleaned = label_prefix.sub("", cleaned).strip("#*- ").strip('"“”')
        if 18 <= len(cleaned) <= 260 and not cleaned.startswith("<") and not looks_like_refusal(cleaned):
            return cleaned
    fallback = label_prefix.sub("", text.strip().splitlines()[0] if text.strip().splitlines() else text)[:260]
    return "" if looks_like_refusal(fallback) else fallback


def clean_quote_candidate(value: str) -> str:
    cleaned = str(value or "").strip().strip('"“”')
    cleaned = re.sub(r"^\s*(?:quote|pull quote|graphic text|caption text)\s*\d*\s*[:\-]\s*", "", cleaned, flags=re.I)
    cleaned = re.sub(r"^\s*(?:\d+[\).\]]|[-*•])\s*", "", cleaned)
    return cleaned.strip().strip('"“”')


def extract_quote_candidates(content: str, limit: int = 12) -> list[str]:
    """Extract multiple plausible quote-card lines from generated content."""
    text = str(content or "").strip()
    if not text:
        return []
    candidates: list[str] = []

    for match in re.finditer(
        r"(?im)^\s*(?:quote|quote option|pull quote|graphic text|caption text)\s*\d*\s*[:\-]\s*(.+)$",
        text,
    ):
        candidates.append(clean_quote_candidate(match.group(1)))

    for match in re.finditer(r"[“\"]([^”\"]{18,260})[”\"]", text):
        candidates.append(clean_quote_candidate(match.group(1)))

    for line in text.splitlines():
        cleaned = clean_quote_candidate(line)
        lower = cleaned.lower()
        if (
            18 <= len(cleaned) <= 260
            and not lower.startswith(("caption", "hashtags", "cta", "source", "mood", "character", "image", "alt text"))
            and ":" not in cleaned[:22]
            and not looks_like_refusal(cleaned)
        ):
            candidates.append(cleaned)

    unique: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = re.sub(r"\s+", " ", candidate.lower()).strip()
        if candidate and key not in seen and not looks_like_refusal(candidate):
            seen.add(key)
            unique.append(candidate)
        if len(unique) >= limit:
            break
    return unique


def parse_instagram_multiformat(draft: str) -> dict[str, dict]:
    """Parse the labeled '=== FORMAT ===' Instagram output into per-format field dicts.

    Returns {format_key: {caption, hashtags, overlay text, cta, image direction, slides[]}}.
    """
    text = str(draft or "")
    parts = re.split(r"(?im)^\s*={2,}\s*(.+?)\s*={2,}\s*$", text)
    if len(parts) < 3:
        return {}

    def parse_block(body: str) -> dict:
        fields: dict[str, object] = {}
        slides: list[str] = []
        current = None
        for line in body.splitlines():
            m_field = re.match(r"(?i)^\s*(caption|hashtags|overlay text|cta|image direction)\s*:\s*(.*)$", line)
            m_slide = re.match(r"(?i)^\s*slide\s*(\d+)\s*(?:\([^)]*\))?\s*:\s*(.*)$", line)
            if m_field:
                current = m_field.group(1).lower()
                fields[current] = m_field.group(2).strip()
            elif m_slide:
                current = ("slide", len(slides))
                slides.append(m_slide.group(2).strip())
            elif current is not None and line.strip():
                if isinstance(current, tuple):
                    slides[current[1]] = (slides[current[1]] + "\n" + line.strip()).strip()
                else:
                    fields[current] = (str(fields[current]) + "\n" + line.strip()).strip()
        if slides:
            fields["slides"] = slides
        return fields

    blocks: dict[str, dict] = {}
    iterator = iter(parts[1:])
    for header, segment in zip(iterator, iterator):
        blocks[header.strip().lower()] = parse_block(segment)
    return blocks


_LINKEDIN_FORMAT_NAMES = ["Standard post", "Article", "Document / carousel", "Poll", "Newsletter"]
_YOUTUBE_DELIVERABLE_NAMES = [
    "Long-form video", "Short", "Community post", "Premiere",
    "Title", "Description", "Pinned comment", "Talking points", "SEO title set",
]


def _section_heading(line: str, known: set):
    """Return the heading label if a line is a section heading, else None.
    Matches '## Heading', '**Heading**', or a bare known section name."""
    match = re.match(r"^\s{0,3}(?:#{1,6}\s+(.+?)|\*\*(.+?)\*\*)\s*:?\s*$", line)
    if match:
        return (match.group(1) or match.group(2) or "").strip().strip(":").strip()
    bare = line.strip().rstrip(":").strip()
    return bare if bare.lower() in known else None


def parse_headed_sections(draft: str, known_names=()) -> list:
    """Split a multi-format draft into ``(heading, body)`` segments on markdown/bold
    headings (or bare known section names). One ``("", draft)`` entry if no headings."""
    text = str(draft or "").strip()
    if not text:
        return []
    known = {name.lower() for name in known_names}
    segments: list = []
    label, buf, found = "", [], False
    for line in text.splitlines():
        head = _section_heading(line, known)
        if head is not None:
            if label or "\n".join(buf).strip():
                segments.append((label, "\n".join(buf).strip()))
            label, buf, found = head, [], True
        else:
            buf.append(line)
    if label or "\n".join(buf).strip():
        segments.append((label, "\n".join(buf).strip()))
    if not found:
        return [("", text)]
    return [(lbl, body) for (lbl, body) in segments if (lbl or body)]


def parse_linkedin_slides(body: str) -> list:
    """Extract ``Slide N: ...`` blocks from a LinkedIn document/carousel segment."""
    slides: list = []
    current = None
    for line in str(body or "").splitlines():
        match = re.match(r"(?i)^\s*slide\s*(\d+)\s*(?:\([^)]*\))?\s*[:\-.]?\s*(.*)$", line)
        if match:
            current = len(slides)
            slides.append(match.group(2).strip())
        elif current is not None and line.strip():
            slides[current] = (slides[current] + "\n" + line.strip()).strip()
    return [slide for slide in slides if slide]


def extract_speaker_from_content(content: str) -> str:
    """Best-effort extraction of a quote's speaker/attribution from generated text."""
    text = str(content or "")
    label_match = re.search(
        r"(?im)^\s*(?:speaker|character|attribution|attributed to|said by|spoken by|—\s*by)\s*[:\-]\s*(.+)$",
        text,
    )
    if label_match:
        name = label_match.group(1).strip().strip('"“”')
        if name and name.lower() != "not specified":
            return name[:60]
    # Dash attribution at the end of a line, e.g. "— Alex Herrera" or "- María García".
    for line in text.splitlines():
        dash_match = re.search(r"[—–-]\s*([A-ZÁÉÍÓÚÑ][\w'’.\- ]{1,48})\s*$", line.strip())
        if dash_match:
            name = dash_match.group(1).strip().strip(".")
            if 1 <= len(name.split()) <= 5 and name.lower() not in {"mortal vengeance"}:
                return name[:60]
    return ""


# Maps each selected Instagram post format to the image format(s) it should render.
# Carousel expands to N distinct slides; story/reel/post each render one image.
INSTAGRAM_POST_FORMAT_VISUALS = {
    "feed post": ["Instagram Post (4:5)"],
    "post": ["Instagram Post (4:5)"],
    "reel": ["Instagram Reel (9:16)"],
    "story": ["Instagram Story (9:16)"],
}


def _carousel_slide_labels(slide_count: int) -> list[str]:
    count = max(2, min(10, int(slide_count or 5)))
    return [f"Carousel Slide {i} (4:5)" for i in range(1, count + 1)]


# Per-format caption/hashtag rules injected into the generation prompt.
INSTAGRAM_FORMAT_SPECS = {
    "feed post": [
        "Caption: full caption up to 2200 characters — open with the hook line, then a body built from the book/source content, end with a clean CTA.",
        "Hashtags: 10-15 mixed, audience-native hashtags.",
        "Image direction: one line of visual direction.",
    ],
    "reel": [
        "Caption: 3-5 punchy lines plus a CTA. Keep it minimal.",
        "Hashtags: 3-5 hashtags maximum.",
        "Image direction: one line of visual direction for the cover frame.",
    ],
    "story": [
        "Overlay text: 1-6 words meant to sit large ON the image (not a paragraph caption).",
        "CTA: a swipe / link / poll prompt.",
        "Hashtags: 1 hashtag maximum, or 'none'.",
        "Image direction: one line of visual direction.",
    ],
}


def instagram_generation_instructions(formats, carousel_slides: int = 5, image_content_type: str = "") -> list[str]:
    """Build the structured per-format output instructions for Instagram generation."""
    selected = [str(f).strip() for f in (formats or []) if str(f).strip()] or ["Feed post"]
    carousel_n = max(2, min(10, int(carousel_slides or 5)))
    lines = [
        "Instruction: The opening hook is ONLY the caption's first line / tonal direction. Build the full caption body from the related book/source content and the knowledge base — never expand the caption from the hook text alone.",
        f"Instruction: Image content type for every generated image: {image_content_type or 'text / quote graphic'}.",
        "Instruction: Produce a SEPARATE format-specific deliverable for EACH selected format below. Use EXACTLY these labeled blocks (header line '=== FORMAT ===' followed by the listed fields), one block per format, and output nothing outside these blocks:",
    ]
    for fmt in selected:
        key = fmt.lower()
        lines.append(f"=== {fmt.upper()} ===")
        if key == "carousel":
            lines.append("Caption: intro hook + a 'swipe →' teaser.")
            lines.append("Hashtags: 5-10 hashtags.")
            for i in range(1, carousel_n + 1):
                role = "cover hook" if i == 1 else ("CTA" if i == carousel_n else "content beat from book")
                lines.append(f"Slide {i} ({role}): <text drawn from book content>")
            lines.append("Image direction: one line of visual direction for the slides.")
        else:
            lines.extend(INSTAGRAM_FORMAT_SPECS.get(key, INSTAGRAM_FORMAT_SPECS["feed post"]))
    return lines


REVIEW_INSTAGRAM_FORMAT_VISUALS = {
    "feed post": "Instagram Post (4:5)",
    "feed": "Instagram Post (4:5)",
    "story": "Instagram Story (9:16)",
    "reel": "Instagram Reel (9:16)",
}

# Caption length rules per Instagram format for pull-quote cards.
REVIEW_FORMAT_CAPTION_RULES = {
    "feed post": "Caption: 1-2 short paragraphs that frame the quote, plus 5-10 hashtags.",
    "story": "Overlay note: the quote sits large ON the image; Caption: a 1-line CTA only (swipe/link), no hashtags.",
    "reel": "Caption: 2-3 punchy lines plus a CTA; 3-5 hashtags max.",
}


def review_visual_formats_for_instagram(instagram_formats, explicit_visual_formats=None) -> list[str]:
    """Map review_pull_quote Instagram format choices to image format labels."""
    labels: list[str] = []
    for fmt in (instagram_formats or []):
        label = REVIEW_INSTAGRAM_FORMAT_VISUALS.get(str(fmt).strip().lower())
        if label:
            labels.append(label)
    for fmt in (explicit_visual_formats or []):
        labels.append(str(fmt))
    seen: set[str] = set()
    result: list[str] = []
    for label in labels:
        if label and label not in seen:
            seen.add(label)
            result.append(label)
    return result or ["Instagram Post (4:5)"]


def review_generation_instructions(instagram_formats, image_type: str = "", card_style: str = "", quote_mode: str = "", manual_quote: str = "") -> list[str]:
    """Structured per-format pull-quote card instructions (BUG-RPQ-04)."""
    selected = [str(f).strip() for f in (instagram_formats or []) if str(f).strip()] or ["Feed post"]
    mode = str(quote_mode or "").lower()
    lines = []
    if "paste" in mode and str(manual_quote or "").strip():
        lines.append(f"Instruction: Use this EXACT pasted pull quote verbatim, do not rewrite it: \"{str(manual_quote).strip()}\".")
    elif "ai-generated" in mode or "ai generated" in mode:
        lines.append("Instruction: Write 3 fresh pull quotes in the STYLE of real reader/critic reviews (do not fabricate a specific outlet or attribute a real reviewer). Mark them clearly as brand-created, style-of-review quotes.")
    else:
        lines.append("Instruction: Use ONLY real review text from the knowledge base; preserve exact wording and keep the real source/attribution.")
    lines.append(f"Instruction: Image type for the cards: {image_type or 'Text-only card'}. Card style: {card_style or 'Dark gothic'}.")
    lines.append("Instruction: Output discrete, ready-to-use pull-quote CARDS. For EACH quote use EXACTLY this labeled block and nothing else between blocks:")
    lines.append("=== PULL QUOTE ===")
    lines.append("Quote: <the pull quote text>")
    lines.append("Attribution: <reviewer name / outlet, or 'brand-created style-of-review'>")
    for fmt in selected:
        rule = REVIEW_FORMAT_CAPTION_RULES.get(fmt.lower(), REVIEW_FORMAT_CAPTION_RULES["feed post"])
        lines.append(f"{fmt} {rule}")
    lines.append("Image direction: <one line of visual direction for the card>")
    lines.append("Instruction: Produce 3-5 such pull-quote card blocks.")
    return lines


def instagram_visual_formats_for_post(post_formats, explicit_visual_formats=None, carousel_slides: int = 5) -> list[str]:
    """Derive the image formats to render from the selected Instagram post formats.

    Carousel expands to ``carousel_slides`` slides; story/reel/post each map to one
    image. Any explicitly chosen image types are appended. Order-preserving de-dupe.
    """
    labels: list[str] = []
    for fmt in (post_formats or []):
        key = str(fmt).strip().lower()
        if key == "carousel":
            labels.extend(_carousel_slide_labels(carousel_slides))
        else:
            labels.extend(INSTAGRAM_POST_FORMAT_VISUALS.get(key, []))
    for fmt in (explicit_visual_formats or []):
        labels.append(str(fmt))
    seen: set[str] = set()
    result: list[str] = []
    for label in labels:
        if label and label not in seen:
            seen.add(label)
            result.append(label)
    return result or ["Instagram Post (4:5)"]


def formats_for_character_spotlight(selection) -> list[str]:
    selected = normalize_selected(selection)
    lookup = selected.lower()
    if "instagram reel" in lookup:
        return ["Instagram Reel (9:16)"]
    if "instagram story" in lookup:
        return ["Instagram Story (9:16)"]
    if "instagram carousel" in lookup:
        return ["Instagram Post (4:5)", "Square Post (1:1)"]
    if "instagram post" in lookup:
        return ["Instagram Post (4:5)"]
    if "youtube cover" in lookup:
        return ["YouTube Image Cover (16:9)"]
    if "youtube post" in lookup:
        return ["YouTube Post (1:1)"]
    if "linkedin post square" in lookup:
        return ["LinkedIn Post Square (1:1)"]
    if "linkedin post horizontal" in lookup:
        return ["LinkedIn Post Horizontal (1.91:1)"]
    if "blog cover" in lookup:
        return ["Blog Cover (16:9)"]
    if "blog square" in lookup:
        return ["Blog Square (1:1)"]
    if "blog horizontal" in lookup:
        return ["Blog Horizontal (1.91:1)"]
    return ["Instagram Post (4:5)"]


def render_generated_visual_package(
    *,
    content_type: str,
    content: str,
    topic: str,
    character_name: str = "",
    format_labels=None,
    theme_name: str = "",
    image_content_type: str = "",
    base_image_path: str = "",
    book: str = "",
) -> dict[str, object]:
    """Render a basic downloadable image package for visual content workflows."""
    if content_type not in {"quote_post", "review_pull_quote", "character_spotlight", "blog_post", "instagram_caption", "linkedin_content", "youtube_content", "newsletter_blurb"}:
        return {}
    # An EXPLICIT empty format list means "skip images" (e.g. all Instagram image
    # chips cleared); None still falls through to the per-type defaults below.
    if format_labels is not None and len(list(format_labels)) == 0:
        return {}

    quote = extract_quote_for_graphic(content)
    if not quote:
        quote = extract_quote_for_graphic(topic)
    if not quote and content_type in {"instagram_caption", "blog_post", "character_spotlight", "linkedin_content", "youtube_content", "newsletter_blurb"}:
        quote = first_nonempty_line(content) or first_nonempty_line(topic)
    if not quote:
        return {}

    output_dir = PROJECT_ROOT / "outputs" / "generated_visuals" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    file_stem = slugify_filename(f"{content_type}_{quote}", "generated_visual")
    default_format_labels = {
        "blog_post": ["Blog Cover (16:9)", "Blog Square (1:1)", "Blog Horizontal (1.91:1)"],
        "instagram_caption": ["Instagram Post (4:5)"],
        "linkedin_content": ["LinkedIn Post Horizontal (1.91:1)"],
        "youtube_content": ["YouTube Image Cover (16:9)"],
        "newsletter_blurb": ["Link Preview (1.91:1)"],
    }
    if not format_labels:
        format_labels = default_format_labels.get(content_type, ["Instagram Post (4:5)"])

    use_case_lookup = {
        "blog_post": "Blog hero images",
        "instagram_caption": "Instagram posts",
        "character_spotlight": "Character illustrations",
        "review_pull_quote": "Social quote cards",
        "linkedin_content": "LinkedIn post images",
        "youtube_content": "YouTube thumbnails",
        "newsletter_blurb": "Email header images",
    }
    content_kind_lookup = {
        "blog_post": "character" if character_name else "hero",
        "instagram_caption": "social post",
        "character_spotlight": "character",
        "review_pull_quote": "quote typography",
        "linkedin_content": "social post",
        "youtube_content": "youtube image",
        "newsletter_blurb": "email header",
    }
    if content_type == "instagram_caption":
        attribution = character_name or extract_speaker_from_content(content) or "Mortal Vengeance"
    elif content_type == "review_pull_quote":
        attribution = character_name or extract_speaker_from_content(content) or "Real reader review"
    else:
        attribution = character_name or extract_speaker_from_content(content) or "Mortal Vengeance"
    # Image content type (BUG-IG-03) chooses provider intent and whether quote text is drawn on-image.
    image_kind = (image_content_type or "").lower()
    content_kind = content_kind_lookup.get(content_type, "")
    render_text_override = None
    if content_type == "instagram_caption":
        if "portrait" in image_kind:
            content_kind = "character portrait illustration"
            render_text_override = False
        elif "abstract" in image_kind or "mood" in image_kind:
            content_kind = "abstract mood illustration"
            render_text_override = False
        elif "mixed" in image_kind:
            content_kind = "social post illustration"
            render_text_override = None
        else:  # text / quote graphic
            content_kind = "social post quote typography"
            render_text_override = True
    elif content_type == "review_pull_quote":
        if "character" in image_kind:
            content_kind = "character portrait with quote typography overlay"
            render_text_override = True
        elif "cover" in image_kind:
            content_kind = "book cover with quote typography overlay"
            render_text_override = True
        else:  # text-only card
            content_kind = "social post quote typography"
            render_text_override = True
    # Carousel slides get distinct text per slide so they don't render identically.
    # Keyed by the specific carousel label (not by position) so non-carousel sizes in
    # the SAME set (feed/story/reel) keep the shared quote — no positional leakage.
    carousel_labels = [lbl for lbl in (format_labels or []) if "Carousel Slide" in str(lbl)]
    carousel_slide_texts = None
    if carousel_labels:
        _slide_candidates = extract_quote_candidates(content) or []
        carousel_slide_texts = {
            lbl: _slide_candidates[i]
            for i, lbl in enumerate(carousel_labels)
            if i < len(_slide_candidates) and _slide_candidates[i]
        }
    # Base image chosen by the user: composite the quote onto it locally and skip
    # fresh generation (the chosen image IS the visual).
    if base_image_path and Path(str(base_image_path)).exists():
        usable_labels = [lbl for lbl in format_labels if lbl in QUOTE_GRAPHIC_FORMATS]
        if usable_labels:
            try:
                return render_quote_cards(
                    quote=quote,
                    attribution=attribution,
                    format_labels=usable_labels,
                    brand_title="Mortal Vengeance",
                    output_dir=output_dir,
                    file_stem=file_stem,
                    theme_name=theme_name or "Gothic",
                    character_name=character_name,
                    background_image_path=str(base_image_path),
                    slide_texts=carousel_slide_texts,
                    book=book,
                )
            except Exception as exc:
                return {"error": str(exc), "quote": quote}

    external_paths: list[str] = []
    external_errors: list[str] = []
    external_prompts: list[str] = []
    for label in format_labels:
        # Only carousel slides get distinct per-slide text; other multi-size sets
        # reuse the same quote so a single post stays consistent across its sizes.
        slide_quote = (carousel_slide_texts or {}).get(label) or quote
        external = generate_external_visual(
            use_case=use_case_lookup.get(content_type, "Social posts"),
            content=slide_quote,
            topic=topic,
            style=theme_name,
            attribution=attribution,
            format_label=label,
            output_dir=output_dir,
            file_stem=file_stem,
            content_kind=content_kind,
            render_text=render_text_override,
        )
        if external.get("path"):
            external_paths.append(str(external["path"]))
        else:
            external_errors.append(f"{label}: {external.get('provider', 'external')} - {external.get('error', 'No image returned')}")
        if external.get("prompt"):
            external_prompts.append(str(external["prompt"]))

    if external_paths:
        zip_path = output_dir / f"{file_stem}_external_images.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in external_paths:
                file_path = Path(path)
                if file_path.exists():
                    archive.write(file_path, arcname=file_path.name)
        result = {
            "paths": external_paths,
            "zip_path": str(zip_path),
            "quote": quote,
            "provider": "external",
            "prompts": external_prompts,
        }
        if external_errors:
            result["error"] = "; ".join(external_errors)
        return result

    try:
        return render_quote_cards(
            quote=quote,
            attribution=attribution,
            format_labels=format_labels,
            brand_title="Mortal Vengeance",
            output_dir=output_dir,
            file_stem=file_stem,
            theme_name=theme_name or "Gothic",
            character_name=character_name or extract_speaker_from_content(content),
            slide_texts=carousel_slide_texts,
            book=book,
        )
    except Exception as exc:
        return {"error": str(exc), "quote": quote}


def render_selected_quote_visual_package(
    *,
    quotes: list[str],
    attribution: str,
    format_labels=None,
    theme_name: str = "",
    book: str = "",
) -> dict[str, object]:
    """Render a downloadable image package for selected quote candidates."""
    selected_quotes = [clean_quote_candidate(quote) for quote in quotes if clean_quote_candidate(quote)]
    if not selected_quotes:
        return {"error": "Select at least one quote card to generate images."}
    output_dir = PROJECT_ROOT / "outputs" / "generated_visuals" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    output_dir.mkdir(parents=True, exist_ok=True)
    all_paths: list[str] = []
    errors: list[str] = []
    formats = format_labels or ["Instagram Post (4:5)"]
    speaker = "" if normalize_selected(attribution) == "Not specified" else normalize_selected(attribution)

    for index, quote in enumerate(selected_quotes, start=1):
        external_rendered = False
        quote_external_errors: list[str] = []
        for label in formats:
            external = generate_external_visual(
                use_case="Instagram quote cards",
                content=quote,
                topic=quote,
                style=theme_name,
                attribution=speaker or "Mortal Vengeance",
                format_label=label,
                output_dir=output_dir,
                file_stem=slugify_filename(f"quote_card_{index}_{quote}", f"quote_card_{index}"),
                content_kind="quote typography",
            )
            if external.get("path"):
                all_paths.append(str(external["path"]))
                external_rendered = True
            else:
                quote_external_errors.append(
                    f"{quote[:60]} / {label}: {external.get('provider', 'external')} - {external.get('error', 'No image returned')}"
                )
        if external_rendered:
            errors.extend(quote_external_errors)
            continue
        try:
            package = render_quote_cards(
                quote=quote,
                attribution=speaker or "Mortal Vengeance",
                format_labels=formats,
                brand_title="Mortal Vengeance",
                output_dir=output_dir,
                file_stem=slugify_filename(f"quote_card_{index}_{quote}", f"quote_card_{index}"),
                theme_name=theme_name or "Gothic",
                character_name=speaker,
                book=book,
            )
            all_paths.extend(str(path) for path in package.get("paths", []))
        except Exception as exc:
            errors.extend(quote_external_errors)
            errors.append(f"{quote[:60]}: {exc}")

    zip_path = output_dir / "selected_quote_cards.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in all_paths:
            file_path = Path(path)
            if file_path.exists():
                archive.write(file_path, arcname=file_path.name)

    result: dict[str, object] = {
        "paths": all_paths,
        "zip_path": str(zip_path),
        "quotes": selected_quotes,
    }
    if errors:
        result["error"] = "; ".join(errors)
    return result


def quote_platform_mockup_html(paths: list[str], package_path: str = "", error: str = "") -> str:
    if error and not paths:
        return f"""
        <div style="border-radius:18px;background:#fff7f7;border:1px solid rgba(142,31,47,.22);padding:16px;color:#8E1F2F;">
            {escape(error)}
        </div>
        """
    if not paths:
        return """
        <div style="border-radius:18px;background:#fff;border:1px dashed rgba(44,24,31,.18);padding:16px;color:#71686a;">
            Select quote cards and generate images to preview them here.
        </div>
        """

    cards = []
    for path in paths[:12]:
        file_path = Path(path)
        label = file_path.stem.replace("_", " ").title()
        src = image_data_url(file_path)
        slug = file_path.stem.lower()
        if "youtube" in slug or "cover" in slug:
            frame = "YouTube thumbnail"
            aspect = "16/9"
            chrome = '<div style="height:26px;background:#0f0f0f;color:white;font-size:12px;padding:6px 10px;">Mortal Vengeance · Thumbnail preview</div>'
        elif "linkedin" in slug or "horizontal" in slug or "link" in slug:
            frame = "LinkedIn / link preview"
            aspect = "1.91/1"
            chrome = '<div style="height:34px;background:#fff;border-bottom:1px solid #e6e2df;color:#0a66c2;font-weight:800;font-size:13px;padding:8px 10px;">LinkedIn post preview</div>'
        elif "story" in slug or "reel" in slug or "9x16" in slug:
            frame = "Instagram story / reel"
            aspect = "9/16"
            chrome = '<div style="height:34px;background:#111;color:white;font-size:12px;padding:8px 10px;">@mortalvengeance</div>'
        else:
            frame = "Instagram post"
            aspect = "4/5" if "4x5" in slug or "post" in slug else "1/1"
            chrome = '<div style="height:38px;background:#fff;border-bottom:1px solid #eee;font-weight:800;font-size:13px;padding:10px;">mortalvengeance</div>'
        cards.append(
            f"""
            <article style="border:1px solid rgba(44,24,31,.14);border-radius:18px;background:#fff;overflow:hidden;box-shadow:0 14px 28px rgba(61,31,41,.07);">
                {chrome}
                <div style="aspect-ratio:{aspect};background:#120d10;display:flex;align-items:center;justify-content:center;">
                    <img src="{escape(src, quote=True)}" alt="{escape(label, quote=True)}" style="width:100%;height:100%;object-fit:contain;display:block;">
                </div>
                <div style="padding:10px 12px;font-size:12px;color:#71686a;">
                    <strong>{escape(frame)}</strong><br>{escape(file_path.name)}
                </div>
            </article>
            """
        )
    package_note = (
        f'<div style="font-size:12px;color:#71686a;margin-top:10px;"><strong>Download package:</strong> {escape(package_path)}</div>'
        if package_path
        else ""
    )
    error_note = (
        f'<div style="font-size:12px;color:#8E1F2F;margin-top:8px;"><strong>Render note:</strong> {escape(error)}</div>'
        if error
        else ""
    )
    return f"""
    <div>
        <div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:14px;">
            {''.join(cards)}
        </div>
        {package_note}
        {error_note}
    </div>
    """


def attribution_for_campaign_asset(content_type: str, widgets: dict[str, object]) -> str:
    if content_type == "review_pull_quote":
        return "Real reader review"
    if content_type == "quote_post":
        characters = normalize_selected(widgets.get("characters").value if widgets.get("characters") else "")
        if characters != "Not specified":
            return characters
    if content_type == "character_spotlight":
        character = normalize_selected(widgets.get("character").value if widgets.get("character") else "")
        if character != "Not specified":
            return character
    return "Mortal Vengeance"


def campaign_graphic_formats(content_type: str, widgets: dict[str, object]) -> list[str]:
    if content_type == "quote_post" and widgets.get("graphic_formats"):
        formats = widgets["graphic_formats"].value or []
        return list(formats) if isinstance(formats, list) else [str(formats)]
    if content_type == "review_pull_quote" and widgets.get("graphic_formats"):
        formats = widgets["graphic_formats"].value or []
        return list(formats) if isinstance(formats, list) else [str(formats)]
    if content_type == "character_spotlight" and widgets.get("format"):
        return formats_for_character_spotlight(widgets["format"].value)
    if content_type == "blog_post":
        return ["Blog Cover (16:9)", "Blog Square (1:1)", "Blog Horizontal (1.91:1)"]
    return ["Instagram Post (4:5)"]


def campaign_graphic_theme(content_type: str, widgets: dict[str, object]) -> str:
    if content_type in {"quote_post", "review_pull_quote"} and widgets.get("graphic_theme"):
        return str(widgets["graphic_theme"].value or "Gothic")
    if content_type == "blog_post":
        return "Press"
    return "Gothic"


def render_campaign_quote_graphic(
    *,
    content_type: str,
    content: str,
    widgets: dict[str, object],
    campaign_topic: str,
    asset_number: int,
    theme_override: str = "",
    book: str = "",
) -> dict[str, object]:
    quote = extract_quote_for_graphic(content)
    if not quote and content_type in {"character_spotlight", "blog_post"}:
        quote = first_nonempty_line(content) or campaign_topic
    if not quote:
        return {}
    output_dir = PROJECT_ROOT / "outputs" / "campaign_quote_graphics" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    file_stem = slugify_filename(f"{content_type}_{asset_number}_{campaign_topic}", "campaign_quote")
    # Prefer the per-type book widget (quote_post has its own), else the campaign book.
    book_widget = widgets.get("book")
    selected_book = normalize_selected(book_widget.value) if book_widget else (book or "")
    try:
        return render_quote_cards(
            quote=quote,
            attribution=attribution_for_campaign_asset(content_type, widgets),
            format_labels=campaign_graphic_formats(content_type, widgets),
            brand_title="Mortal Vengeance",
            output_dir=output_dir,
            file_stem=file_stem,
            theme_name=(theme_override or campaign_graphic_theme(content_type, widgets)),
            book=selected_book,
            character_name=normalize_selected(
                widgets.get("characters").value
                if widgets.get("characters")
                else widgets.get("character").value
                if widgets.get("character")
                else ""
            ),
        )
    except Exception as exc:
        return {"error": str(exc), "quote": quote}


def campaign_asset_preview_html(asset: dict[str, object]) -> str:
    content_type = str(asset.get("content_type") or "")
    label = escape(str(asset.get("label") or content_type))
    content = str(asset.get("content") or "")
    style = escape(str(asset.get("style") or "default"))
    quantity_label = escape(str(asset.get("asset_label") or "Asset"))
    scheduled_date = escape(str(asset.get("post_date") or ""))
    escaped_content = escape(content)
    fields = asset.get("preview_fields") if isinstance(asset.get("preview_fields"), dict) else {}
    quote_graphic = asset.get("quote_graphic") if isinstance(asset.get("quote_graphic"), dict) else {}
    asset_image_paths = asset.get("image_paths") if isinstance(asset.get("image_paths"), list) else []
    first_image = str(asset_image_paths[0]) if asset_image_paths else ""

    if content_type in {"instagram_caption", "youtube_content", "linkedin_content", "character_spotlight"}:
        preview = build_preview_html(content_type, {"draft": content, "image_path": first_image, "post_date": str(asset.get("post_date") or ""), **fields})
    elif content_type in {"quote_post", "review_pull_quote"}:
        quote_text = extract_quote_for_graphic(content)
        preview = build_quote_preview_html(
            "quote_post" if content_type == "quote_post" else "review_pull_quote",
            {
                "quote_text": quote_text,
                "quote_source": quote_graphic.get("zip_path") or "Campaign generated quote",
                "quote_moods": fields.get("quote_moods", ""),
                "quote_characters": fields.get("quote_characters", ""),
                "quote_spoiler": "Campaign asset",
                "quote_request": style,
                "review_quote_text": quote_text,
                "review_source": "Campaign generated from review grounding",
                "review_attribution": attribution_for_campaign_asset(content_type, asset.get("widgets", {}) if isinstance(asset.get("widgets"), dict) else {}),
                "review_mode": style,
                "review_promo_line": "Mortal Vengeance",
            },
        )
        paths = quote_graphic.get("paths") or []
        first_path = str(paths[0]) if paths else ""
        if first_path and Path(first_path).exists():
            preview = f"""
            {preview}
            <div style="margin-top:14px;border-radius:16px;overflow:hidden;border:1px solid rgba(142,31,47,.18);">
                <img src="{escape(image_data_url(first_path), quote=True)}" alt="Generated quote graphic" style="width:100%;display:block;">
            </div>
            <div style="font-size:12px;color:#6b6b6b;margin-top:8px;"><strong>Graphic:</strong> {escape(first_path)}</div>
            <div style="font-size:12px;color:#6b6b6b;"><strong>Package:</strong> {escape(str(quote_graphic.get("zip_path") or ""))}</div>
            """
        elif quote_graphic.get("error"):
            preview = f'{preview}<div style="font-size:12px;color:#8E1F2F;margin-top:8px;">Graphic render note: {escape(str(quote_graphic.get("error")))}</div>'
    else:
        preview = f"""
        <div style="border-radius:18px;background:#fff;border:1px solid rgba(44,24,31,.12);box-shadow:0 14px 30px rgba(61,31,41,.07);padding:18px;">
            <div style="font-size:13px;line-height:1.55;white-space:pre-wrap;color:#2f292b;">{markdown_to_html(content)}</div>
        </div>
        """

    return f"""
    <article style="border-radius:22px;background:#fffdfb;border:1px solid rgba(44,24,31,.13);box-shadow:0 18px 36px rgba(61,31,41,.08);padding:18px;margin-bottom:18px;">
        <div style="display:flex;gap:12px;align-items:flex-start;justify-content:space-between;margin-bottom:14px;">
            <div>
                <div style="font-family:Georgia,serif;font-size:22px;font-weight:800;color:#171217;">{label}</div>
                <div style="font-size:12px;color:#71686a;">{quantity_label} · {style}{(' · 🗓 ' + scheduled_date) if scheduled_date else ''}</div>
            </div>
            <div style="border-radius:999px;background:#f8eceb;color:#8E1F2F;padding:6px 10px;font-size:11px;font-weight:800;text-transform:uppercase;letter-spacing:.08em;">{escape(content_type)}</div>
        </div>
        {preview}
        <details style="margin-top:14px;">
            <summary style="cursor:pointer;color:#5d3a40;font-weight:700;font-size:13px;list-style:none;display:flex;align-items:center;gap:6px;"><span style="font-size:11px;">▸</span> View generated text</summary>
            <pre style="white-space:pre-wrap;font-family:ui-monospace,Menlo,monospace;font-size:12px;line-height:1.5;background:#fbf7f4;border:1px solid rgba(44,24,31,.10);border-radius:12px;padding:12px;overflow:auto;">{escaped_content}</pre>
        </details>
    </article>
    """


def build_campaign_preview_html(assets: list[dict[str, object]]) -> str:
    if not assets:
        return """
        <div style="border-radius:20px;padding:20px;background:#ffffff;border:1px dashed rgba(17,17,20,.16);color:#6b6b6b;">
            Campaign previews will appear here separated by content type.
        </div>
        """
    grouped: dict[str, list[dict[str, object]]] = {}
    for asset in assets:
        grouped.setdefault(str(asset.get("content_type") or "content"), []).append(asset)
    sections = []
    for content_type, content_assets in grouped.items():
        label = CAMPAIGN_FORMAT_LABELS.get(content_type, content_type)
        cards = "\n".join(campaign_asset_preview_html(asset) for asset in content_assets)
        sections.append(
            f"""
            <section style="margin-bottom:22px;">
                <div style="font-family:Georgia,serif;font-size:24px;font-weight:800;margin:0 0 12px;color:#171217;">{escape(label)}</div>
                {cards}
            </section>
            """
        )
    return "\n".join(sections)


async def generate_draft_from_fields(
    *,
    content_type: str,
    topic,
    related_book,
    platform,
    social_objectives,
    audience,
    cta,
    constraints,
    quote_book,
    quote_moods,
    character_tags,
    podcast_format,
    podcast_speakers,
    podcast_roles,
    podcast_tone,
    podcast_length,
    elevenlabs_model,
    elevenlabs_voice_ids,
    podcast_show_title,
    podcast_episode_title,
    podcast_episode_number,
    blog_length,
    blog_format,
    blog_sections,
    blog_structure_options,
    blog_seo_keywords,
    blog_image_mode,
    blog_character,
    instagram_formats,
    instagram_hashtags,
    instagram_hook,
    instagram_image_content_type="",
    instagram_carousel_slides=5,
    cs_character,
    cs_focus,
    cs_platform_format,
    cs_image_mode,
    nl_subject,
    nl_preview,
    nl_structure,
    pr_timing,
    pr_embargo_date,
    pr_city,
    pr_state,
    pr_release_date,
    pr_contact_name,
    pr_contact_title,
    pr_organization,
    pr_contact_email,
    pr_contact_phone,
    pr_website,
    pr_news_angle,
    pr_release_goal,
    pr_primary_announcement,
    pr_supporting_proof,
    pr_quote_source,
    pr_target_media,
    pr_required_assets,
    review_book="",
    review_source="",
    review_quote="",
    review_mode="",
    review_attribution="",
    review_promo_line="",
    review_graphic_formats=None,
    review_visual_style="",
    review_quote_mode="",
    review_manual_quote="",
    review_image_type="",
    review_instagram_formats=None,
    review_card_style="",
    status,
    filtered_context_path,
    prompt_path,
    draft_path,
    generated_output,
    visual_format_labels=None,
    visual_theme_name: str = "",
    saved_draft_path=None,
    structured_brief_output=None,
    uploaded_image_path: str = "",
    generated_visual_path=None,
    generated_visual_package_path=None,
    form_fields_snapshot=None,
) -> None:
    status.value = "Generating..."
    generated_output.value = ""
    filtered_context_path.value = ""
    prompt_path.value = ""
    draft_path.value = ""
    if saved_draft_path is not None:
        saved_draft_path.value = ""
    if generated_visual_path is not None:
        generated_visual_path.value = ""
    if generated_visual_package_path is not None:
        generated_visual_package_path.value = ""

    brief = build_brief(
        content_type=content_type,
        topic=topic,
        related_book=related_book,
        platform=platform,
        social_objectives=social_objectives,
        audience=audience,
        cta=cta,
        constraints=constraints,
        quote_book=quote_book,
        quote_moods=quote_moods,
        character_tags=character_tags,
        podcast_format=podcast_format,
        podcast_speakers=podcast_speakers,
        podcast_roles=podcast_roles,
        podcast_tone=podcast_tone,
        podcast_length=podcast_length,
        elevenlabs_model=elevenlabs_model,
        elevenlabs_voice_ids=elevenlabs_voice_ids,
        podcast_show_title=podcast_show_title,
        podcast_episode_title=podcast_episode_title,
        podcast_episode_number=podcast_episode_number,
        blog_length=blog_length,
        blog_format=blog_format,
        blog_sections=blog_sections,
        blog_structure_options=blog_structure_options,
        blog_seo_keywords=blog_seo_keywords,
        blog_image_mode=blog_image_mode,
        blog_character=blog_character,
        instagram_formats=instagram_formats,
        instagram_hashtags=instagram_hashtags,
        instagram_hook=instagram_hook,
        instagram_image_content_type=instagram_image_content_type,
        instagram_carousel_slides=instagram_carousel_slides,
        cs_character=cs_character,
        cs_focus=cs_focus,
        cs_platform_format=cs_platform_format,
        cs_image_mode=cs_image_mode,
        nl_subject=nl_subject,
        nl_preview=nl_preview,
        nl_structure=nl_structure,
        pr_timing=pr_timing,
        pr_embargo_date=pr_embargo_date,
        pr_city=pr_city,
        pr_state=pr_state,
        pr_release_date=pr_release_date,
        pr_contact_name=pr_contact_name,
        pr_contact_title=pr_contact_title,
        pr_organization=pr_organization,
        pr_contact_email=pr_contact_email,
        pr_contact_phone=pr_contact_phone,
        pr_website=pr_website,
        pr_news_angle=pr_news_angle,
        pr_release_goal=pr_release_goal,
        pr_primary_announcement=pr_primary_announcement,
        pr_supporting_proof=pr_supporting_proof,
        pr_quote_source=pr_quote_source,
        pr_target_media=pr_target_media,
        pr_required_assets=pr_required_assets,
        review_book=review_book,
        review_source=review_source,
        review_quote=review_quote,
        review_mode=review_mode,
        review_attribution=review_attribution,
        review_promo_line=review_promo_line,
        review_graphic_formats=review_graphic_formats,
        review_visual_style=review_visual_style,
        review_quote_mode=review_quote_mode,
        review_manual_quote=review_manual_quote,
        review_image_type=review_image_type,
        review_instagram_formats=review_instagram_formats,
        review_card_style=review_card_style,
        uploaded_image_path=uploaded_image_path,
    )
    if structured_brief_output is not None:
        structured_brief_output.value = brief

    try:
        if content_type in IMAGE_AWARE_CONTENT_TYPES and uploaded_image_path:
            result = await asyncio.to_thread(
                run_pipeline_with_image,
                content_type=content_type,
                topic=brief,
                image_path=uploaded_image_path,
            )
        else:
            result = await asyncio.to_thread(
                run_pipeline,
                content_type=content_type,
                topic=brief,
            )
    except Exception as exc:
        status.value = f"Generation failed: {exc}"
        ui.notify(status.value, type="negative")
        return

    status.value = "Draft ready"
    filtered_context_path.value = str(result["filtered_context_path"])
    prompt_path.value = str(result["prompt_path"])
    draft_path.value = str(result["draft_path"])
    generated_output.value = result["generated_content"]
    character_for_visual = normalize_selected(
        cs_character
        if content_type == "character_spotlight"
        else blog_character
        if content_type == "blog_post" and "character portrait" in str(blog_image_mode or "").lower()
        else character_tags
        if content_type == "quote_post"
        else ""
    )
    visual_package = (
        {}
        if content_type == "quote_post"
        else render_generated_visual_package(
            content_type=content_type,
            content=result["generated_content"],
            topic=brief,
            character_name="" if character_for_visual == "Not specified" else character_for_visual,
            format_labels=visual_format_labels,
            theme_name=visual_theme_name,
            image_content_type=(
                instagram_image_content_type if content_type == "instagram_caption"
                else review_image_type if content_type == "review_pull_quote"
                else ""
            ),
            book=(review_book if content_type == "review_pull_quote" else related_book),
        )
    )
    if generated_visual_path is not None:
        paths = visual_package.get("paths") if isinstance(visual_package, dict) else []
        generated_visual_path.value = str(paths[0]) if paths else str(visual_package.get("error", "") if isinstance(visual_package, dict) else "")
    if generated_visual_package_path is not None and isinstance(visual_package, dict):
        generated_visual_package_path.value = str(visual_package.get("zip_path") or "")
    topic_lines = str(topic or "").strip().splitlines()
    draft_title = (
        str(podcast_episode_title or "").strip()
        or str(podcast_show_title or "").strip()
        or (topic_lines[0][:90] if topic_lines else "")
        or f"{content_type} draft"
    )
    draft_record = save_draft(
        title=draft_title,
        content_type=content_type,
        content=result["generated_content"],
        source_path=result["draft_path"],
        metadata={
            "topic": str(topic or "").strip(),
            "related_book": normalize_selected(related_book),
            "brief": brief,
            "platform": normalize_selected(platform),
            "audience": normalize_selected(audience),
            "constraints": normalize_selected(constraints),
            "uploaded_image_path": uploaded_image_path,
            "form_fields": form_fields_snapshot or {},
        },
    )
    if saved_draft_path is not None:
        saved_draft_path.value = draft_record["path"]
    ui.notify("Content generated successfully.", type="positive")




NICEGUI_APP_CSS = r"""
:root {
    --font-display: "Playfair Display", Georgia, "Times New Roman", serif;
    --font-ui: "Inter", Arial, ui-sans-serif, system-ui, sans-serif;
    --mce-bg: #fbf7f1;
    --mce-bg-2: #fffaf6;
    --mce-card: #ffffff;
    --mce-card-solid: #ffffff;
    --mce-surface-soft: #fffaf6;
    --mce-ink: #161014;
    --mce-muted: #6f6668;
    --mce-border: #eadfda;
    --mce-border-strong: rgba(143, 15, 45, 0.25);
    --mce-accent: #8f0f2d;
    --mce-accent-dark: #260d1a;
    --mce-wine-dark: #140911;
    --mce-red: #c91f3a;
    --mce-red-bright: #df2f4b;
    --mce-accent-soft: #f8e5e4;
    --mce-gold: #d9a85c;
    --mce-gold-soft: #f4d49a;
    --mce-success: #2e7d4f;
    --mce-shadow: 0 24px 70px rgba(70, 20, 25, 0.18);
    --mce-shadow-soft: 0 14px 36px rgba(40, 20, 20, 0.07);
}

/* NTH-02: dark mode palette (overrides the CSS variables the theme is built on). */
body.mce-dark {
    --mce-bg: #14100f;
    --mce-bg-2: #1c1715;
    --mce-card: rgba(34, 28, 30, 0.96);
    --mce-card-solid: #221c1e;
    --mce-ink: #f4ece6;
    --mce-muted: #b8a8ab;
    --mce-border: rgba(255, 255, 255, 0.12);
    --mce-border-strong: rgba(214, 130, 146, 0.42);
    --mce-accent-soft: #3a1822;
    --mce-shadow: 0 22px 70px rgba(0, 0, 0, 0.55);
    --mce-shadow-soft: 0 12px 36px rgba(0, 0, 0, 0.45);
}
body.mce-dark .mce-tabs,
body.mce-dark .mce-dashboard-pill {
    background: rgba(255, 255, 255, 0.06);
}
body.mce-dark .mce-card,
body.mce-dark .mce-subcard,
body.mce-dark .mce-expansion {
    background: var(--mce-card-solid);
}
.mce-dark-toggle {
    color: var(--mce-accent) !important;
}

html, body, #app, .q-layout, .q-page, .nicegui-content {
    min-height: 100%;
    background:
        radial-gradient(900px 420px at 7% 0%, rgba(142, 31, 47, 0.13), transparent 56%),
        radial-gradient(780px 360px at 96% 4%, rgba(234, 213, 173, 0.22), transparent 52%),
        linear-gradient(180deg, var(--mce-bg-2) 0%, var(--mce-bg) 100%) !important;
    color: var(--mce-ink);
    font-family: var(--font-ui);
}

.nicegui-content {
    padding: 0 !important;
    max-width: none !important;
}

.mce-shell {
    width: 100%;
    max-width: 1320px;
    margin: 0 auto;
    padding: 34px 28px 56px;
    gap: 22px;
}

.mce-hero {
    position: relative;
    width: 100%;
    overflow: hidden;
    border-radius: 24px;
    cursor: pointer;
    line-height: 0;
    box-shadow: 0 24px 70px rgba(70, 20, 25, 0.18);
    transition: transform 0.25s ease, box-shadow 0.25s ease;
}

.mce-hero:hover {
    transform: translateY(-2px);
    box-shadow: 0 30px 84px rgba(70, 20, 25, 0.26);
}

/* Pre-composed banner artwork (text, logo and CTA are baked into the image). */
.mce-banner-img {
    display: block;
    width: 100%;
    height: auto;
    border-radius: 24px;
}

.mce-kicker {
    font-size: 12px;
    line-height: 1;
    letter-spacing: 0.18em;
    text-transform: uppercase;
    color: var(--mce-gold);
    font-weight: 700;
}

.mce-hero-title {
    margin-top: 18px;
    max-width: 720px;
    font-family: var(--font-display);
    font-size: clamp(34px, 4.2vw, 56px);
    line-height: 0.98;
    font-weight: 800;
    letter-spacing: -0.02em;
}

.mce-hero-title span {
    color: var(--mce-red-bright);
}

.mce-hero-subtitle {
    margin-top: 16px;
    max-width: 620px;
    color: #f3e9df;
    font-size: 18px;
    line-height: 1.5;
}

.mce-hero-chips {
    display: flex;
    gap: 14px;
    margin-top: 24px;
    flex-wrap: wrap;
    justify-content: center;
}

.mce-hero-chip {
    border: 1px solid rgba(244, 212, 154, 0.35);
    background: rgba(0, 0, 0, 0.32);
    color: #f8efe4;
    border-radius: 999px;
    padding: 10px 16px;
    font-size: 13px;
}

.mce-bottom-banner {
    position: relative;
    margin-top: 26px;
    width: 100%;
    overflow: hidden;
    border-radius: 24px;
    cursor: pointer;
    line-height: 0;
    box-shadow: 0 24px 70px rgba(70, 20, 25, 0.18);
    transition: transform 0.25s ease, box-shadow 0.25s ease;
}

.mce-bottom-banner:hover {
    transform: translateY(-2px);
    box-shadow: 0 30px 84px rgba(70, 20, 25, 0.26);
}

.mce-bottom-banner-title {
    font-family: var(--font-display);
    font-size: 28px;
    line-height: 1.05;
    margin: 0;
    font-weight: 800;
}

.mce-bottom-banner-title span {
    color: var(--mce-red-bright);
}

.mce-bottom-banner-sub {
    margin-top: 8px;
    font-size: 13px;
    color: rgba(255, 255, 255, 0.7);
}

.mce-cta {
    background: linear-gradient(135deg, #c61e3c, #97122a) !important;
    color: #ffffff !important;
    border: 1px solid rgba(255, 255, 255, 0.2) !important;
    border-radius: 16px !important;
    padding: 16px 28px !important;
    font-weight: 800 !important;
    white-space: nowrap;
}

.mce-tabs-wrap {
    width: 100%;
    display: flex;
    justify-content: center;
    margin-top: 2px;
}

.mce-tabs {
    width: auto;
    background: rgba(255, 255, 255, 0.42);
    border: 1px solid var(--mce-border);
    border-radius: 999px;
    padding: 4px;
    box-shadow: 0 8px 26px rgba(41, 24, 31, 0.05);
}

.mce-tabs .q-tab {
    min-height: 40px;
    border-radius: 999px;
    padding: 0 18px;
    color: var(--mce-muted);
    font-size: 12px;
    font-weight: 850;
    letter-spacing: 0.04em;
}

.mce-tabs .q-tab--active {
    color: #fff !important;
    background: linear-gradient(135deg, var(--mce-accent), #b61734);
    font-weight: 900;
    box-shadow: 0 8px 18px rgba(143, 15, 45, 0.32);
}

.mce-tabs .q-tab__indicator {
    display: none;
}

.mce-tab-panels,
.mce-panel,
.q-tab-panel {
    background: transparent !important;
    box-shadow: none !important;
}

.mce-panel {
    padding: 0 !important;
}

.mce-grid {
    width: 100%;
    display: grid;
    grid-template-columns: minmax(0, 1fr) minmax(360px, 430px);
    gap: 24px;
    align-items: start;
}

/* UI-06: single-column until the right panel has content. */
.mce-grid-single {
    grid-template-columns: minmax(0, 1fr);
}

.mce-card {
    width: 100%;
    min-width: 0;
    border: 1px solid var(--mce-border);
    border-radius: 26px;
    background: var(--mce-card);
    box-shadow: var(--mce-shadow-soft);
    padding: 24px;
    backdrop-filter: blur(10px);
}

.mce-sticky {
    position: sticky;
    top: 22px;
    align-self: start;
}

.mce-card-title {
    margin: 0;
    font-family: var(--font-display);
    font-size: 25px;
    line-height: 1.08;
    font-weight: 850;
    letter-spacing: -0.04em;
    color: var(--mce-ink);
}

.mce-card-subtitle {
    margin-top: 8px;
    color: var(--mce-muted);
    font-size: 13px;
    line-height: 1.45;
}

.mce-section-title {
    color: var(--mce-accent-dark);
    font-size: 12px;
    font-weight: 900;
    letter-spacing: 0.09em;
    text-transform: uppercase;
}

.mce-stack {
    gap: 14px;
}

.mce-stack .q-field,
.mce-stack .q-select,
.mce-stack .q-textarea,
.mce-stack .q-input {
    width: 100%;
}

.mce-stack .q-field__control {
    border-radius: 15px !important;
    background: rgba(255, 255, 255, 0.72) !important;
}

.mce-stack .q-field__native,
.mce-stack .q-field__input,
.mce-stack textarea,
.mce-stack input {
    color: var(--mce-ink) !important;
}

.mce-stack textarea {
    resize: vertical;
}

.mce-subcard {
    width: 100%;
    border: 1px solid var(--mce-border);
    border-radius: 20px;
    background: rgba(255, 255, 255, 0.68);
    box-shadow: none;
    padding: 16px;
}

.mce-subcard .q-card__section {
    padding: 0;
}

.mce-two-col {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 12px;
    width: 100%;
}

.mce-actions {
    display: flex;
    flex-wrap: wrap;
    gap: 10px;
    padding-top: 6px;
}

.mce-actions .q-btn,
.q-btn.mce-primary,
.q-btn.mce-secondary {
    min-height: 42px;
    border-radius: 13px;
    font-weight: 850;
    letter-spacing: 0.01em;
}

.q-btn.mce-primary,
.mce-primary {
    background: linear-gradient(180deg, var(--mce-accent), var(--mce-accent-dark)) !important;
    color: white !important;
    box-shadow: 0 12px 24px rgba(142, 31, 47, 0.26);
}

.q-btn.mce-secondary,
.mce-secondary {
    background: rgba(142, 31, 47, 0.08) !important;
    color: var(--mce-accent-dark) !important;
    border: 1px solid rgba(142, 31, 47, 0.16) !important;
}

.mce-primary:hover,
.mce-secondary:hover {
    transform: translateY(-1px);
}

.mce-preview {
    width: 100%;
    min-height: 150px;
    max-height: 470px;
    overflow: auto;
    border: 1px dashed var(--mce-border-strong);
    border-radius: 18px;
    background: rgba(255, 255, 255, 0.64);
    padding: 14px;
}

.mce-output-textarea textarea {
    min-height: 230px !important;
    font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, "Liberation Mono", monospace;
    font-size: 12.5px !important;
    line-height: 1.58 !important;
}

.mce-script-textarea textarea {
    min-height: 320px !important;
    font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, "Liberation Mono", monospace;
    font-size: 12.8px !important;
    line-height: 1.62 !important;
}

.mce-path-field input,
.mce-path-field textarea {
    font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, "Liberation Mono", monospace;
    font-size: 12px !important;
}

/* UX-09: read-only fields look distinct from editable inputs. */
.mce-readonly .q-field__control {
    background: rgba(61, 31, 41, 0.05) !important;
}
.mce-readonly .q-field__control:before {
    border-style: dashed !important;
    border-color: rgba(61, 31, 41, 0.18) !important;
}
.mce-readonly .q-field__control:after {
    display: none !important;
}
.mce-readonly input,
.mce-readonly textarea {
    cursor: default !important;
    color: #5c524b !important;
}

/* UI-04/05: suggest-button status states. */
.mce-status-loading {
    color: #8E1F2F !important;
    font-weight: 700 !important;
}
.mce-status-success {
    color: #1f8a4c !important;
    font-weight: 800 !important;
}
.mce-status-error {
    color: #c0392b !important;
    font-weight: 700 !important;
}

/* UI-03: image upload success state + thumbnail. */
.mce-upload-success {
    color: #1f8a4c !important;
    font-weight: 800 !important;
}
.mce-upload-thumb {
    width: 96px;
    height: 96px;
    border-radius: 12px;
    object-fit: cover;
    border: 2px solid #1f8a4c;
    margin-top: 6px;
}

/* BUG-IMG-02: visible gallery of generated image assets. */
.mce-gallery {
    display: flex;
    flex-wrap: wrap;
    gap: 10px;
}
.mce-gallery-thumb {
    width: 140px;
    height: 140px;
    border-radius: 12px;
    object-fit: cover;
    border: 1px solid var(--mce-border);
}

/* UX-08: saved drafts list view. */
.mce-saved-list {
    border: 1px solid var(--mce-border);
    border-radius: 14px;
    overflow: hidden;
    gap: 0 !important;
}
.mce-saved-row {
    display: flex;
    align-items: center;
    gap: 12px;
    width: 100%;
    padding: 10px 14px;
    border-bottom: 1px solid var(--mce-border);
    cursor: pointer;
    flex-wrap: nowrap;
}
.mce-saved-row:last-child {
    border-bottom: none;
}
.mce-saved-row:hover:not(.mce-saved-head) {
    background: rgba(142, 31, 47, 0.06);
}
.mce-saved-head {
    cursor: default;
    background: rgba(61, 31, 41, 0.05);
    font-size: 11px;
    letter-spacing: 0.08em;
    text-transform: uppercase;
    font-weight: 800;
    color: #6b6b6b;
}
.mce-saved-col-title {
    flex: 2;
    min-width: 0;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
    font-weight: 700;
}
.mce-saved-col-type {
    flex: 1;
    min-width: 0;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
}
.mce-saved-col-date {
    flex: 1;
    min-width: 0;
    text-align: right;
    color: #6b6b6b;
    font-size: 12px;
}

/* UX-06: dashboard live-stats block. */
.mce-dashboard-stats {
    display: flex;
    flex-direction: column;
    gap: 12px;
    padding: 20px 22px;
    justify-content: center;
}
.mce-stat-row {
    display: flex;
    align-items: baseline;
    justify-content: space-between;
    gap: 12px;
}
.mce-stat-label {
    font-size: 12px;
    color: #6b6b6b;
}
.mce-stat-value {
    font-family: var(--font-display);
    font-weight: 800;
    font-size: 15px;
    color: #17120f;
    text-align: right;
    max-width: 62%;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
}

/* ---- Three-card dashboard summary row (This session / Quick tip / Glance) ---- */
.mce-dash-row {
    display: grid;
    grid-template-columns: 1.3fr 1fr 1fr;
    gap: 28px;
    width: 100%;
    align-items: stretch;
}
.mce-dash-card {
    border-radius: 28px;
    padding: 30px 34px;
    min-height: 268px;
    box-shadow: 0 18px 45px rgba(20, 20, 30, 0.08);
    display: flex;
    flex-direction: column;
}
.mce-dash-card-title {
    font-family: var(--font-display);
    font-weight: 700;
    font-size: 30px;
    line-height: 1.1;
    margin: 0 0 22px;
}
/* Card 1 — This session (dark). */
.mce-session-card {
    background: #11121a;
    color: #ffffff;
    position: relative;
}
.mce-session-card::before {
    content: "";
    position: absolute;
    left: 24px;
    top: 32px;
    bottom: 32px;
    width: 3px;
    border-radius: 99px;
    background: #e03b52;
}
.mce-session-inner {
    padding-left: 26px;
    display: flex;
    flex-direction: column;
    height: 100%;
}
.mce-session-inner .mce-dash-card-title {
    color: #ffffff;
}
.mce-session-row {
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 14px;
    padding: 18px 0;
    border-top: 1px solid rgba(255, 255, 255, 0.14);
    font-family: var(--font-ui);
    font-size: 16px;
    color: #cfd0d8;
}
.mce-session-row:first-of-type {
    border-top: none;
}
.mce-session-value {
    font-weight: 700;
    color: #f04c61;
    text-align: right;
    max-width: 55%;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
}
/* Card 2 — Quick tip (cream). */
.mce-tip-card {
    background: #f8f0e5;
    color: #1f2430;
}
.mce-tip-header {
    display: flex;
    align-items: center;
    gap: 16px;
    font-family: var(--font-ui);
    font-weight: 700;
    font-size: 20px;
    color: #1f2430;
}
.mce-tip-icon {
    width: 54px;
    height: 54px;
    min-width: 54px;
    border-radius: 50%;
    background: #c61f3d;
    color: #ffffff;
    display: grid;
    place-items: center;
    box-shadow: 0 8px 20px rgba(198, 31, 61, 0.25);
}
.mce-tip-icon .q-icon {
    font-size: 26px;
}
.mce-tip-copy {
    margin: 26px 0 auto;
    font-family: var(--font-ui);
    font-size: 19px;
    line-height: 1.55;
    color: #1f2430;
}
.mce-tip-button {
    margin-top: 26px;
    align-self: flex-start;
    border: 1px solid rgba(255, 255, 255, 0.2) !important;
    background: linear-gradient(135deg, #c61e3c, #97122a) !important;
    border-radius: 16px;
    padding: 12px 24px;
    font-family: var(--font-ui);
    font-weight: 700;
    font-size: 16px;
    text-transform: none;
    box-shadow: 0 8px 20px rgba(198, 31, 61, 0.28);
}
.mce-tip-button,
.mce-tip-button .q-btn__content {
    color: #ffffff !important;
}
/* Card 3 — Content at a glance (white). */
.mce-glance-card {
    background: #ffffff;
    color: #2a2a2a;
}
.mce-glance-title {
    color: #15151a;
    margin-bottom: 10px;
}
.mce-glance-row {
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 14px;
    padding: 16px 0;
    border-top: 1px solid #e9e2dc;
    font-family: var(--font-ui);
    font-size: 16px;
    color: #2a2a2a;
}
.mce-glance-left {
    display: flex;
    align-items: center;
    gap: 14px;
}
.mce-glance-icon {
    color: #8a8077;
    font-size: 20px;
}
.mce-glance-value {
    font-family: var(--font-ui);
    font-weight: 800;
    font-size: 18px;
    color: #15151a;
}

.mce-expansion {
    width: 100%;
    border: 1px solid var(--mce-border);
    border-radius: 18px;
    background: rgba(255, 255, 255, 0.58);
    overflow: hidden;
}

.mce-expansion .q-item {
    min-height: 48px;
    color: var(--mce-accent-dark);
    font-weight: 850;
}

.mce-note {
    padding: 12px 14px;
    border: 1px solid rgba(142, 31, 47, 0.14);
    border-radius: 16px;
    background: #fff8f8;
    color: #5d3a40;
    font-size: 13px;
    line-height: 1.45;
}

.mce-status-row {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 10px;
}

.mce-status-chip {
    padding: 12px 14px;
    border: 1px solid var(--mce-border);
    border-radius: 16px;
    background: #fff;
}

.mce-status-chip small {
    display: block;
    color: var(--mce-muted);
    font-size: 10px;
    font-weight: 900;
    letter-spacing: 0.09em;
    text-transform: uppercase;
    margin-bottom: 4px;
}

.mce-status-chip strong {
    color: var(--mce-success);
}

.mce-audio {
    width: 100%;
    border-radius: 999px;
}

.mce-divider {
    height: 1px;
    width: 100%;
    margin: 4px 0;
    background: var(--mce-border);
}

.mce-muted {
    color: var(--mce-muted);
}

.mce-dashboard-top {
    display: grid;
    grid-template-columns: minmax(0, 1fr) minmax(280px, 430px);
    gap: 18px;
    align-items: stretch;
}

.mce-dashboard-logo-card {
    min-height: 220px;
    overflow: hidden;
    border-radius: 28px;
    border: 1px solid var(--mce-border);
    background: #160d14;
    box-shadow: var(--mce-shadow-soft);
}

.mce-dashboard-logo-card img {
    width: 100%;
    height: 100%;
    object-fit: cover;
}

.mce-dashboard-feature {
    position: relative;
    min-height: 220px;
    overflow: hidden;
    border-radius: 28px;
    border: 1px solid rgba(255, 255, 255, 0.14);
    background: #220a15;
    box-shadow: var(--mce-shadow);
    cursor: pointer;
}

.mce-dashboard-feature img {
    width: 100%;
    height: 100%;
    min-height: 220px;
    object-fit: cover;
    display: block;
}

.mce-dashboard-feature:hover,
.mce-dashboard-card:hover,
.mce-quick-card:hover {
    transform: translateY(-2px);
}

.mce-dashboard-filter-row {
    display: flex;
    flex-wrap: wrap;
    gap: 10px;
}

.mce-dashboard-pill {
    padding: 9px 14px;
    border-radius: 999px;
    border: 1px solid var(--mce-border);
    background: rgba(255, 255, 255, 0.72);
    color: var(--mce-muted);
    font-size: 12px;
    font-weight: 850;
    cursor: pointer;
    transition: background 0.15s ease, color 0.15s ease;
}
.mce-dashboard-pill:hover {
    border-color: var(--mce-accent);
    color: var(--mce-accent);
}

.mce-dashboard-pill-active {
    background: linear-gradient(180deg, var(--mce-accent), var(--mce-accent-dark));
    color: white;
    border-color: transparent;
}

.mce-dashboard-grid {
    display: grid;
    grid-template-columns: repeat(5, minmax(0, 1fr));
    gap: 14px;
}

.mce-dashboard-card,
.mce-quick-card,
.mce-draft-card {
    min-width: 0;
    border: 1px solid var(--mce-border);
    border-radius: 18px;
    background: rgba(255, 255, 255, 0.78);
    box-shadow: var(--mce-shadow-soft);
    transition: transform 160ms ease, box-shadow 160ms ease;
    cursor: pointer;
}

.mce-dashboard-card {
    min-height: 154px;
    padding: 18px;
}

.mce-dashboard-icon {
    width: 48px;
    height: 48px;
    border-radius: 999px;
    object-fit: cover;
    margin-bottom: 14px;
}

.mce-dashboard-card-title,
.mce-dash-section-title {
    font-family: var(--font-display);
    color: var(--mce-ink);
    font-weight: 850;
    letter-spacing: -0.03em;
}

.mce-dashboard-card-title {
    font-size: 16px;
    line-height: 1.1;
}

.mce-dashboard-card-copy {
    margin-top: 7px;
    color: #4e4649;
    font-size: 12px;
    line-height: 1.35;
}

.mce-dashboard-card-arrow {
    margin-top: auto;
    color: var(--mce-accent-dark);
    font-size: 22px;
    line-height: 1;
}

.mce-dashboard-lower {
    display: grid;
    grid-template-columns: minmax(0, 0.88fr) minmax(0, 1.12fr);
    gap: 18px;
}

.mce-dashboard-panel {
    border: 1px solid var(--mce-border);
    border-radius: 22px;
    background: rgba(255, 255, 255, 0.7);
    box-shadow: var(--mce-shadow-soft);
    padding: 18px;
}

.mce-dash-section-title {
    font-size: 20px;
}

.mce-draft-card {
    padding: 12px 14px;
}

.mce-quick-grid {
    display: grid;
    grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 12px;
}

.mce-quick-card {
    min-height: 120px;
    padding: 16px;
    text-align: center;
}

.mce-quick-card img {
    width: 42px;
    height: 42px;
    object-fit: cover;
    margin: 0 auto 10px;
    border-radius: 14px;
}

.mce-quick-card strong,
.mce-draft-card strong {
    color: var(--mce-ink);
    font-size: 13px;
}

.mce-quick-card span,
.mce-draft-card span {
    display: block;
    margin-top: 5px;
    color: var(--mce-muted);
    font-size: 11.5px;
    line-height: 1.35;
}

@media (max-width: 1100px) {
    .mce-grid {
        grid-template-columns: 1fr;
    }

    .mce-sticky {
        position: static;
        max-height: none;
    }

    .mce-dashboard-top,
    .mce-dashboard-lower,
    .mce-dash-row {
        grid-template-columns: 1fr;
    }

    .mce-dashboard-grid {
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }

    .mce-quick-grid {
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }
}

@media (max-width: 720px) {
    .mce-shell {
        padding: 20px 13px 42px;
    }

    .mce-hero {
        border-radius: 22px;
        padding: 24px;
    }

    .mce-tabs-wrap {
        justify-content: stretch;
    }

    .mce-tabs {
        width: 100%;
        border-radius: 18px;
    }

    .mce-tabs .q-tabs__content {
        flex-wrap: wrap;
    }

    .mce-card {
        border-radius: 22px;
        padding: 18px;
    }

    .mce-two-col,
    .mce-status-row {
        grid-template-columns: 1fr;
    }

    .mce-dashboard-grid,
    .mce-quick-grid {
        grid-template-columns: 1fr;
    }
}

/* UI-11: turn NiceGUI's faint "Connection lost / reconnecting" popup into a clearly
   visible branded toast (the framework toggles aria-hidden; we only restyle it). */
.nicegui-error-popup {
    bottom: 0 !important;
    left: 50% !important;
    right: auto !important;
    transform: translateX(-50%) !important;
    margin: 0 0 24px !important;
    padding: 14px 24px 14px 48px !important;
    border: none !important;
    border-radius: 14px !important;
    background: linear-gradient(135deg, #8E1F2F, #4d1020) !important;
    color: #fff6f3 !important;
    box-shadow: 0 18px 44px rgba(20, 10, 16, 0.5) !important;
    gap: 2px !important;
    max-width: min(92vw, 460px) !important;
    font-size: 13px !important;
    line-height: 1.4 !important;
}
.body--light .nicegui-error-popup,
.body--dark .nicegui-error-popup {
    background: linear-gradient(135deg, #8E1F2F, #4d1020) !important;
}
.nicegui-error-popup > span:first-child {
    font-weight: 800 !important;
    font-size: 14px !important;
}
.nicegui-error-popup:dir(ltr) > span:first-child::before,
.nicegui-error-popup:dir(rtl) > span:first-child::before {
    left: 18px !important;
    right: auto !important;
    top: 14px !important;
}
"""


_ASSETS_STATIC_MOUNTED = False


def install_nicegui_theme() -> None:
    global _ASSETS_STATIC_MOUNTED
    ui.colors(primary="#8f0f2d", secondary="#140911", accent="#d9a85c", positive="#2E7D4F")
    if not _ASSETS_STATIC_MOUNTED:
        try:
            app.add_static_files("/assets", str(PROJECT_ROOT / "assets"))
        except Exception:
            pass
        _ASSETS_STATIC_MOUNTED = True
    ui.add_head_html(
        '<link rel="preconnect" href="https://fonts.googleapis.com">'
        '<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>'
        '<link href="https://fonts.googleapis.com/css2?family=Playfair+Display:wght@600;700;800;900&family=Inter:wght@400;500;600;700;800&display=swap" rel="stylesheet">'
    )
    ui.add_head_html(f"<style>{NICEGUI_APP_CSS}</style>")


def make_primary_button(label: str, on_click):
    return ui.button(label, on_click=on_click).props("unelevated no-caps").classes("mce-primary")


def make_secondary_button(label: str, on_click):
    return ui.button(label, on_click=on_click).props("unelevated no-caps").classes("mce-secondary")


def section_heading(title: str, subtitle: str = "") -> None:
    ui.label(title).classes("mce-card-title")
    if subtitle:
        ui.label(subtitle).classes("mce-card-subtitle")
    ui.element("div").classes("mce-divider")


def subcard_heading(title: str) -> None:
    ui.label(title).classes("mce-section-title")


def apply_field_props(component, props: str = "outlined dense"):
    return component.props(props).classes("w-full")


def base_image_picker():
    """Campaign 'base image' picker that re-scans disk on focus so a freshly
    generated image shows up in the list without a manual page refresh."""
    picker = apply_field_props(
        ui.select(list_existing_images(), label="Base image (optional)", with_input=True).props("clearable")
    )
    picker.on("focus", lambda _event: picker.set_options(list_existing_images()))
    return picker


def attach_podcast_script_toolbox(script_textarea, marker: str) -> None:
    """Add a 'Sound & performance toolbox' under a podcast script editor: pick a
    performance cue / emotional tag or a stock transition / sound effect and insert it
    into the script at the cursor. Performance cues shape ElevenLabs delivery (v3);
    transitions / SFX / music are producer notes (stripped before TTS)."""
    script_textarea.classes(add=marker)
    with ui.expansion("Sound & performance toolbox", icon="graphic_eq", caption="Insert performance cues, transitions & sound effects into the script").classes("mce-expansion"):
        with ui.column().classes("mce-stack w-full"):
            ui.label(
                "Pick a tag and insert it into the script at your cursor. Performance cues shape "
                "ElevenLabs delivery; transitions and sound effects are producer notes (not spoken)."
            ).classes("mce-muted")
            cue_select = apply_field_props(
                ui.select(PODCAST_PERFORMANCE_CUES, label="Performance cue / emotion", with_input=True).props("clearable")
            )
            sfx_select = apply_field_props(
                ui.select(PODCAST_SOUND_EFFECTS, label="Transition / sound effect", with_input=True).props("clearable")
            )

            async def insert_tag(tag, own_line: bool) -> None:
                tag = str(tag or "").strip()
                if not tag:
                    ui.notify("Pick a cue or effect first.", type="warning")
                    return
                payload = f"\n{tag}\n" if own_line else f"{tag} "
                new_value = await ui.run_javascript(
                    "(() => {"
                    f"  const ta = document.querySelector('.{marker} textarea');"
                    "  if (!ta) return null;"
                    "  const s = (ta.selectionStart != null) ? ta.selectionStart : ta.value.length;"
                    "  const e = (ta.selectionEnd != null) ? ta.selectionEnd : ta.value.length;"
                    f"  const ins = {json.dumps(payload)};"
                    "  ta.value = ta.value.slice(0, s) + ins + ta.value.slice(e);"
                    "  ta.dispatchEvent(new Event('input', {bubbles: true}));"
                    "  ta.focus();"
                    "  const p = s + ins.length;"
                    "  ta.setSelectionRange(p, p);"
                    "  return ta.value;"
                    "})()"
                )
                if new_value is not None:
                    script_textarea.value = new_value

            with ui.row().classes("mce-actions"):
                make_secondary_button("Insert cue", lambda: insert_tag(cue_select.value, False))
                make_secondary_button("Insert effect", lambda: insert_tag(sfx_select.value, True))


def readonly_input(label: str, value: str = "", mono: bool = True):
    # mono=True keeps the monospace 'path field' look (for file paths); prose status
    # fields pass mono=False so human-readable text isn't shown in a code font.
    classes = "w-full mce-readonly" + (" mce-path-field" if mono else "")
    return ui.input(label=label, value=value).props("readonly outlined dense").classes(classes)


def readonly_textarea(label: str, value: str = ""):
    return ui.textarea(label=label, value=value).props("readonly outlined autogrow").classes("w-full mce-output-textarea mce-readonly")


def dashboard_asset(name: str) -> str:
    path = DASHBOARD_ASSET_DIR / name
    return str(path) if path.exists() else ""


@ui.page("/")
def index() -> None:
    install_nicegui_theme()

    content_types = list_supported_content_types()
    generator_content_types = [content_type for content_type in content_types if content_type != "podcast"]
    generator_content_type_labels = {ct: content_type_label(ct) for ct in generator_content_types}
    default_content_type = "blog_post" if "blog_post" in generator_content_types else generator_content_types[0]

    with ui.column().classes("mce-shell"):
        with ui.element("div").classes("mce-hero").on(
            "click", lambda: open_generator(default_content_type)
        ):
            ui.image(dashboard_asset("telltales-ink-hero.webp")).classes("mce-banner-img")

        with ui.row().classes("mce-tabs-wrap items-center"):
            with ui.tabs().classes("mce-tabs") as tabs:
                dashboard_tab = ui.tab("Dashboard")
                generator_tab = ui.tab("Generator")
                saved_drafts_tab = ui.tab("Saved Drafts")
                campaign_tab = ui.tab("Campaign Mode")
                podcast_tab = ui.tab("Podcast Studio")
                chapter_promos_tab = ui.tab("Chapter Promos")

            dark_mode = ui.dark_mode(value=False)
            dark_state = {"on": False}

            def toggle_dark() -> None:
                dark_state["on"] = not dark_state["on"]
                dark_mode.value = dark_state["on"]
                action = "add" if dark_state["on"] else "remove"
                ui.run_javascript(f"document.body.classList.{action}('mce-dark')")

            ui.button(icon="dark_mode", on_click=toggle_dark).props("flat round dense").classes("mce-dark-toggle").tooltip("Toggle dark mode")

        def open_generator(content_type_name: str | None = None) -> None:
            tabs.value = generator_tab
            if content_type_name and content_type_name in generator_content_type_labels:
                content_type.value = content_type_name
                apply_visibility(content_type_name)

        def open_campaign() -> None:
            tabs.value = campaign_tab

        def open_podcast() -> None:
            tabs.value = podcast_tab

        def open_chapter_promos() -> None:
            tabs.value = chapter_promos_tab

        def open_drafts() -> None:
            tabs.value = saved_drafts_tab

        session_stats = {"drafts": 0, "last_type": "—", "last_topic": "—"}
        glance_values: dict[str, object] = {}
        dashboard_refresh: dict[str, object] = {"stats": None, "recent": None, "glance": None}

        def update_drafts_badge() -> None:
            count = len(list_saved_drafts(250))
            saved_drafts_tab.props(f'label="Saved Drafts ({count})"')

        def refresh_dashboard() -> None:
            for key in ("stats", "recent", "glance"):
                fn = dashboard_refresh.get(key)
                if callable(fn):
                    fn()
            update_drafts_badge()

        def _on_tab_change(_event) -> None:
            refresh_dashboard()
            ui.run_javascript(
                'window.scrollTo({top:0,behavior:"smooth"});'
                'const c=document.querySelector(".nicegui-content")||document.querySelector(".q-page-container");'
                'if(c)c.scrollTo({top:0,behavior:"smooth"});'
            )

        tabs.on_value_change(_on_tab_change)
        update_drafts_badge()

        def record_session_draft(content_type_value: str, topic_text: str = "", count: int = 1) -> None:
            """Record generation activity for the 'This session' card — used by every
            generation path (generator, campaign, podcast, chapter promos)."""
            session_stats["drafts"] += max(1, int(count or 1))
            if content_type_value:
                session_stats["last_type"] = content_type_label(content_type_value)
            topic_clean = str(topic_text or "").strip()
            if topic_clean:
                session_stats["last_topic"] = topic_clean[:60]
            refresh_dashboard()

        with ui.tab_panels(tabs, value=dashboard_tab).classes("w-full mce-tab-panels"):
            with ui.tab_panel(dashboard_tab).classes("mce-panel"):
                with ui.column().classes("mce-stack w-full"):
                    with ui.element("section").classes("mce-dash-row"):
                        # Card 1 — live session stats (dark).
                        with ui.element("div").classes("mce-dash-card mce-session-card"):
                            with ui.element("div").classes("mce-session-inner"):
                                ui.label("This session").classes("mce-dash-card-title")
                                with ui.element("div").classes("mce-session-row"):
                                    ui.label("Drafts generated").classes("mce-session-label")
                                    stat_drafts = ui.label("0").classes("mce-session-value")
                                with ui.element("div").classes("mce-session-row"):
                                    ui.label("Last content type").classes("mce-session-label")
                                    stat_last_type = ui.label("—").classes("mce-session-value")
                                with ui.element("div").classes("mce-session-row"):
                                    ui.label("Last topic").classes("mce-session-label")
                                    stat_last_topic = ui.label("—").classes("mce-session-value")

                        # Card 2 — quick tip pointing at Campaign Mode (cream).
                        with ui.element("div").classes("mce-dash-card mce-tip-card"):
                            with ui.element("div").classes("mce-tip-header"):
                                with ui.element("div").classes("mce-tip-icon"):
                                    ui.icon("bolt")
                                ui.label("Quick tip").classes("mce-tip-title")
                            ui.label(
                                "Use Campaign Mode to map out your content calendar and stay ahead."
                            ).classes("mce-tip-copy")
                            ui.button(
                                "Go to Campaign Mode",
                                on_click=lambda: open_campaign(),
                            ).props("no-caps unelevated icon-right=arrow_forward").classes("mce-tip-button")

                        # Card 3 — aggregate content stats (white).
                        with ui.element("div").classes("mce-dash-card mce-glance-card"):
                            ui.label("Content at a glance").classes("mce-dash-card-title mce-glance-title")
                            for _g_icon, _g_label, _g_key in (
                                ("description", "Saved Drafts", "saved"),
                                ("event", "Content Generated", "generated"),
                                ("layers", "Formats Available", "formats"),
                                ("mic", "Podcast Episodes", "podcast"),
                            ):
                                with ui.element("div").classes("mce-glance-row"):
                                    with ui.element("div").classes("mce-glance-left"):
                                        ui.icon(_g_icon).classes("mce-glance-icon")
                                        ui.label(_g_label).classes("mce-glance-label")
                                    glance_values[_g_key] = ui.label("—").classes("mce-glance-value")

                    def refresh_session_stats() -> None:
                        stat_drafts.set_text(str(session_stats["drafts"]))
                        stat_last_type.set_text(session_stats["last_type"])
                        stat_last_topic.set_text(session_stats["last_topic"])

                    dashboard_refresh["stats"] = refresh_session_stats

                    def refresh_glance_stats() -> None:
                        outputs_dir = PROJECT_ROOT / "outputs"
                        generated = 0
                        podcasts = 0
                        if outputs_dir.exists():
                            for entry in outputs_dir.iterdir():
                                if entry.is_file():
                                    generated += 1
                                    if entry.name.lower().startswith("podcast"):
                                        podcasts += 1
                        glance_values["saved"].set_text(str(len(list_saved_drafts(250))))
                        glance_values["generated"].set_text(str(generated))
                        glance_values["formats"].set_text(f"{len(content_types)}+")
                        glance_values["podcast"].set_text(str(podcasts))

                    dashboard_refresh["glance"] = refresh_glance_stats
                    refresh_glance_stats()

                    dashboard_filters = ["All Formats", "Social Media", "Long-Form", "PR & Media", "Quotes", "Email"]
                    dashboard_cards = [
                        ("Podcast", "Full script + show notes with timestamps.", "podcast_icon.webp", open_podcast, "Long-Form"),
                        ("Instagram Caption", "Captions + hashtags for the feed.", "instagram_caption_icon.webp", lambda: open_generator("instagram_caption"), "Social Media"),
                        ("YouTube", "Titles, descriptions, and talking points.", "youtube_content_icon.webp", lambda: open_generator("youtube_content"), "Social Media"),
                        ("LinkedIn", "Posts and articles that build authority.", "linkedin_content_icon.webp", lambda: open_generator("linkedin_content"), "Social Media"),
                        ("Press Release", "News-ready copy for media and PR.", "press_release_icon.webp", lambda: open_generator("press_release"), "PR & Media"),
                        ("Quote Post", "Book quotes with branded graphics.", "quote_post_icon.webp", lambda: open_generator("quote_post"), "Quotes"),
                        ("Pull Quote", "Shareable quotes for promotions.", "review_pull_quote_icon.webp", lambda: open_generator("review_pull_quote"), "Quotes"),
                        ("Blog Post", "Long-form articles that inform and inspire.", "blog_post_icon.webp", lambda: open_generator("blog_post"), "Long-Form"),
                        ("Newsletter Blurb", "Email-ready content that drives opens.", "newsletter_blurb_icon.webp", lambda: open_generator("newsletter_blurb"), "Email"),
                        ("Character Spotlight", "Deep dives into characters and arcs.", "character_spotlight_icon.webp", lambda: open_generator("character_spotlight"), "Social Media"),
                    ]

                    pill_elements: dict[str, object] = {}
                    card_elements: list[tuple[object, str]] = []

                    def apply_dashboard_filter(selected: str) -> None:
                        for name, pill in pill_elements.items():
                            if name == selected:
                                pill.classes(add="mce-dashboard-pill-active")
                            else:
                                pill.classes(remove="mce-dashboard-pill-active")
                        for card_el, category in card_elements:
                            card_el.set_visibility(selected == "All Formats" or category == selected)

                    with ui.element("div").classes("mce-dashboard-filter-row"):
                        for filter_name in dashboard_filters:
                            pill = ui.label(filter_name).classes("mce-dashboard-pill")
                            pill.on("click", lambda _event, name=filter_name: apply_dashboard_filter(name))
                            pill_elements[filter_name] = pill

                    with ui.element("section").classes("mce-dashboard-grid"):
                        for title, copy, icon, action, category in dashboard_cards:
                            card = ui.card().classes("mce-dashboard-card")
                            card.on("click", lambda _event, action=action: action())
                            card_elements.append((card, category))
                            with card:
                                ui.image(dashboard_asset(icon)).classes("mce-dashboard-icon")
                                ui.label(title).classes("mce-dashboard-card-title")
                                ui.label(copy).classes("mce-dashboard-card-copy")
                                ui.icon("arrow_forward").classes("mce-dashboard-card-arrow")

                    apply_dashboard_filter("All Formats")

                    with ui.element("section").classes("mce-dashboard-lower"):
                        with ui.element("div").classes("mce-dashboard-panel"):
                            with ui.row().classes("items-center justify-between w-full"):
                                ui.label("Recent Drafts").classes("mce-dash-section-title")
                                ui.button("View all drafts", on_click=open_drafts).props("flat no-caps dense").classes("mce-secondary")
                            recent_drafts_container = ui.column().classes("w-full mce-stack")

                            def refresh_recent_drafts() -> None:
                                recent_drafts_container.clear()
                                with recent_drafts_container:
                                    records = list_saved_drafts(4)
                                    if not records:
                                        ui.label("No saved drafts yet. Generate content to see it here.").classes("mce-dashboard-card-copy")
                                        return
                                    for record in records:
                                        rec_title = record.get("title") or "Untitled draft"
                                        rec_type = content_type_label(record.get("content_type") or "content")
                                        rec_when = str(record.get("updated_at") or record.get("created_at") or "").replace("T", " ")
                                        draft_card = ui.card().classes("mce-draft-card")
                                        draft_card.on("click", lambda _event: open_drafts())
                                        with draft_card:
                                            ui.label(rec_title).classes("mce-dashboard-card-title")
                                            ui.label(f"{rec_type} · {rec_when}").classes("mce-dashboard-card-copy")

                            dashboard_refresh["recent"] = refresh_recent_drafts
                            refresh_recent_drafts()

                        with ui.element("div").classes("mce-dashboard-panel"):
                            with ui.row().classes("items-center justify-between w-full"):
                                ui.label("Quick Start").classes("mce-dash-section-title")
                                ui.button("Explore templates", on_click=lambda: open_generator("blog_post")).props("flat no-caps dense").classes("mce-secondary")
                            quick_cards = [
                                ("Start with a Brief", "Open the default content builder.", "quick_brief.webp", lambda: open_generator(default_content_type)),
                                ("Campaign Generator", "Create multiple assets from one brief.", "quick_campaign.webp", open_campaign),
                                ("Continue Latest Draft", "Pick up where you left off most recently.", "quick_continue.webp", open_drafts),
                                ("Browse Templates", "Explore proven formats and examples.", "quick_templates.webp", lambda: open_generator("blog_post")),
                            ]
                            with ui.element("div").classes("mce-quick-grid"):
                                for title, copy, icon, action in quick_cards:
                                    quick = ui.element("div").classes("mce-quick-card")
                                    quick.on("click", lambda _event, action=action: action())
                                    with quick:
                                        ui.image(dashboard_asset(icon))
                                        ui.html(f"<strong>{escape(title)}</strong><span>{escape(copy)}</span>", sanitize=False)

                    with ui.element("div").classes("mce-bottom-banner").on(
                        "click", lambda: open_generator(default_content_type)
                    ):
                        ui.image(dashboard_asset("telltales-ink-bottom-banner.webp")).classes("mce-banner-img")

            with ui.tab_panel(chapter_promos_tab).classes("mce-panel"):
                with ui.element("section").classes("mce-grid"):
                    with ui.card().classes("mce-card"):
                        section_heading(
                            "Chapter Promos",
                            "Read a novel chapter, then generate a chapter summary plus promos for Instagram, Blog, LinkedIn, and YouTube.",
                        )
                        with ui.column().classes("mce-stack w-full"):
                            chapter_books = chapter_reader.list_chapter_books()
                            chapter_book = apply_field_props(
                                ui.select(
                                    chapter_books,
                                    value=(chapter_books[0] if chapter_books else None),
                                    label="Novel",
                                ),
                            )
                            _initial_chapters = chapter_reader.chapter_options(chapter_books[0]) if chapter_books else {}
                            chapter_select = apply_field_props(
                                ui.select(
                                    _initial_chapters,
                                    value=next(iter(_initial_chapters), None),
                                    label="Chapter",
                                    with_input=True,
                                ),
                            )
                            chapter_platforms = apply_field_props(
                                ui.select(
                                    list(CHAPTER_PROMO_PLATFORMS.keys()),
                                    value=list(CHAPTER_PROMO_PLATFORMS.keys()),
                                    multiple=True,
                                    label="Generate promos for",
                                ),
                                "outlined dense use-chips clearable",
                            )

                            def refresh_chapter_options(_event=None) -> None:
                                options = chapter_reader.chapter_options(chapter_book.value) if chapter_book.value else {}
                                chapter_select.set_options(options, value=next(iter(options), None))

                            chapter_book.on_value_change(refresh_chapter_options)

                            chapter_status = readonly_input("Status", "Ready", mono=False)
                            chapter_progress = ui.linear_progress(value=0, show_value=False).classes("w-full")
                            chapter_progress.visible = False
                            chapter_progress_label = ui.label("Ready").classes("mce-muted")
                            chapter_output = ui.textarea(
                                label="Chapter summary + promos (editable)", value="",
                            ).props("outlined autogrow").classes("w-full mce-output-textarea")
                            chapter_saved_path = readonly_input("Saved draft path", "")

                            async def generate_chapter_promos() -> None:
                                book = chapter_book.value
                                chapter_id = chapter_select.value
                                platform_labels = [p for p in (chapter_platforms.value or []) if p]
                                if not book or not chapter_id:
                                    chapter_status.value = "Pick a novel and a chapter first."
                                    ui.notify(chapter_status.value, type="warning")
                                    return
                                if not platform_labels:
                                    chapter_status.value = "Select at least one platform."
                                    ui.notify(chapter_status.value, type="warning")
                                    return
                                chapter = await asyncio.to_thread(chapter_reader.get_chapter, book, chapter_id)
                                if not chapter:
                                    chapter_status.value = "That chapter could not be read."
                                    ui.notify(chapter_status.value, type="warning")
                                    return
                                started_at = datetime.now()
                                chapter_progress.visible = True
                                total_assets = len(platform_labels) + 1
                                total_steps = total_assets + 1
                                done = 0
                                chapter_status.value = "Reading chapter and summarizing..."
                                await set_generation_progress(
                                    progress=chapter_progress, label=chapter_progress_label, status=chapter_status,
                                    started_at=started_at, completed_steps=done, total_steps=total_steps,
                                    completed_assets=done, total_assets=total_assets, message="Summarizing chapter...",
                                )
                                chapter_text = chapter_reader.chapter_text_for_prompt(chapter)
                                summary_prompt = (
                                    f'You are a book-marketing strategist for the novel "{book}".\n'
                                    f"Read this chapter ({chapter['label']}) and produce a tight, SPOILER-AWARE marketing brief.\n\n"
                                    f'CHAPTER TEXT:\n"""\n{chapter_text}\n"""\n\n'
                                    "Return in markdown with these sections:\n"
                                    "- **Summary:** 3-5 sentences on what happens (no major twist or ending spoilers).\n"
                                    "- **Key characters:** the characters featured in this chapter.\n"
                                    "- **Themes & mood:** the core emotional beats and tone.\n"
                                    "- **Promo hooks:** 3 spoiler-free teaser lines usable on social media.\n"
                                )
                                try:
                                    summary = (await asyncio.to_thread(generate_text, summary_prompt)).strip()
                                except Exception as exc:
                                    chapter_status.value = f"Chapter summary failed: {exc}"
                                    ui.notify(chapter_status.value, type="negative")
                                    chapter_progress.visible = False
                                    return
                                done += 1
                                combined = f"# {book} — {chapter['label']}\n\n## Chapter Summary\n\n{summary}\n"
                                empty_platforms: list[str] = []
                                for platform_label in platform_labels:
                                    content_type_value = CHAPTER_PROMO_PLATFORMS[platform_label]
                                    await set_generation_progress(
                                        progress=chapter_progress, label=chapter_progress_label, status=chapter_status,
                                        started_at=started_at, completed_steps=done, total_steps=total_steps,
                                        completed_assets=done, total_assets=total_assets,
                                        message=f"Generating {platform_label} promo...",
                                    )
                                    brief = (
                                        f"Topic: Promotional {platform_label} content teasing {chapter['label']} of the novel.\n"
                                        f"Related book/source: {book}\n"
                                        f"Chapter: {chapter['label']}\n\n"
                                        f"Chapter marketing brief (internal grounding — summary, characters, themes, hooks):\n{summary}\n\n"
                                        "Instruction: Create promotional content that teases THIS chapter to build interest in the book. "
                                        "Stay spoiler-free — do not reveal major twists or the ending. Match the book's dark, literary tone. "
                                        "Ground every reference in the chapter brief above; do not invent plot details."
                                    )
                                    try:
                                        result = await asyncio.to_thread(run_pipeline, content_type_value, brief)
                                        generated = str(result.get("generated_content") or "").strip()
                                        if not generated:
                                            generated = f"_(No content returned for {platform_label}. Try regenerating.)_"
                                            empty_platforms.append(platform_label)
                                    except Exception as exc:
                                        generated = f"_{platform_label} promo failed: {exc}_"
                                        empty_platforms.append(platform_label)
                                    combined += f"\n\n---\n\n## {platform_label} Promo\n\n{generated}"
                                    done += 1
                                chapter_output.value = combined
                                try:
                                    record = save_draft(
                                        title=f"{book} — {chapter['label']} promos"[:90],
                                        content_type="chapter_promo",
                                        content=combined,
                                        metadata={"book": book, "chapter": chapter["label"], "platforms": platform_labels},
                                    )
                                    chapter_saved_path.value = record["path"]
                                except Exception:
                                    pass
                                await finish_generation_progress(
                                    progress=chapter_progress, label=chapter_progress_label, status=chapter_status,
                                    total_assets=total_assets, message="Chapter promos ready.",
                                )
                                if empty_platforms:
                                    ui.notify(
                                        f"Generated — but no content came back for: {', '.join(empty_platforms)}. Try regenerating.",
                                        type="warning",
                                    )
                                else:
                                    ui.notify("Chapter promos generated and saved to drafts.", type="positive")
                                record_session_draft("chapter_promo", f"{book} — {chapter['label']}")
                                update_drafts_badge()

                            with ui.row().classes("mce-actions"):
                                make_primary_button("Generate Chapter Promos", generate_chapter_promos)
                                make_secondary_button(
                                    "Download (.docx)",
                                    lambda: download_docx(chapter_output.value, f"{chapter_book.value or 'chapter'} promos"),
                                )

            with ui.tab_panel(generator_tab).classes("mce-panel"):
                with ui.element("section").classes("mce-grid"):
                    with ui.card().classes("mce-card"):
                        section_heading("Content Brief", "Shape the prompt, then generate a draft with the existing pipeline.")
                        with ui.column().classes("mce-stack w-full"):
                            content_type = apply_field_props(
                                ui.select(generator_content_type_labels, value=default_content_type, label="Content type"),
                            )
                            topic = apply_field_props(
                                ui.textarea(label="Topic", placeholder="What should this content be about?"),
                                "outlined autogrow",
                            )
                            related_book = apply_field_props(
                                ui.select(
                                    QUOTE_BOOK_OPTIONS,
                                    value=QUOTE_BOOK_OPTIONS[0] if QUOTE_BOOK_OPTIONS else None,
                                    label="Related book/source (optional)",
                                ),
                            )
                            platform = apply_field_props(
                                ui.select(PLATFORM_OPTIONS, value=PLATFORM_OPTIONS[0], label="Platform / destination"),
                            )
                            social_objectives = apply_field_props(
                                ui.select(SOCIAL_OBJECTIVES, multiple=True, label="Social objectives"),
                                "outlined dense use-chips clearable",
                            )
                            audience = apply_field_props(
                                ui.select(AUDIENCE_OPTIONS, multiple=True, label="Audience"),
                                "outlined dense use-chips clearable",
                            )
                            cta = apply_field_props(
                                ui.input(label="CTA / next step", placeholder="What should the reader do next?"),
                            )
                            constraints = apply_field_props(
                                ui.select(CONSTRAINT_OPTIONS, multiple=True, label="Constraints"),
                                "outlined dense use-chips clearable",
                            )

                            preview_refresh: dict[str, object] = {"fn": None}

                            async def save_social_image(event, path_input, status_label, platform_name: str, preview=None) -> None:
                                suffix = Path(event.file.name or "").suffix.lower()
                                if suffix not in ALLOWED_IMAGE_EXTENSIONS:
                                    status_label.set_text("Use PNG, JPG, JPEG, or WEBP.")
                                    status_label.classes(remove="mce-upload-success")
                                    ui.notify("Please upload a PNG, JPG, JPEG, or WEBP image.", type="warning")
                                    return
                                IMAGE_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
                                destination = IMAGE_UPLOAD_DIR / safe_upload_filename(event.file.name)
                                await event.file.save(destination)
                                path_input.value = str(destination)
                                status_label.set_text(f"✓ Uploaded: {event.file.name}")
                                status_label.classes(add="mce-upload-success")
                                if preview is not None:
                                    preview.set_source(image_data_url(str(destination)))
                                    preview.set_visibility(True)
                                refresh = preview_refresh.get("fn")
                                if callable(refresh):
                                    refresh()
                                ui.notify(f"{platform_name} image uploaded.", type="positive")

                            def attach_asset_picker(path_input, status_label, thumb=None) -> None:
                                """Add a 'choose an existing image' picker beside an upload field (BUG-IMG-01)."""
                                picker = apply_field_props(
                                    ui.select({}, label="Or choose an existing image", with_input=True).props("clearable"),
                                    "outlined dense",
                                )

                                def apply_choice(value) -> None:
                                    if not value:
                                        return
                                    path_input.value = str(value)
                                    status_label.set_text(f"✓ Selected: {Path(str(value)).name}")
                                    status_label.classes(add="mce-upload-success")
                                    if thumb is not None:
                                        thumb.set_source(image_data_url(str(value)))
                                        thumb.set_visibility(True)
                                    refresh = preview_refresh.get("fn")
                                    if callable(refresh):
                                        refresh()

                                picker.on("focus", lambda _event: picker.set_options(list_existing_images()))
                                picker.on_value_change(lambda event: apply_choice(event.value))
                                picker.set_options(list_existing_images())

                            with ui.card().classes("mce-subcard") as instagram_card:
                                subcard_heading("Instagram Settings")
                                instagram_formats = apply_field_props(
                                    ui.select(INSTAGRAM_FORMAT_OPTIONS, multiple=True, label="Instagram formats"),
                                    "outlined dense use-chips clearable",
                                )
                                ui.label("Each selected format gets its own caption, image, and aspect ratio (Post 4:5, Reel/Story 9:16, Carousel slides).").classes("mce-muted")
                                instagram_carousel_slides = apply_field_props(
                                    ui.number(label="Carousel slide count", value=5, min=2, max=10, precision=0),
                                )
                                instagram_carousel_slides.set_visibility(False)

                                def update_carousel_slides_visibility(formats=None) -> None:
                                    selected = [str(f).strip().lower() for f in (formats if formats is not None else (instagram_formats.value or []))]
                                    instagram_carousel_slides.set_visibility("carousel" in selected)

                                instagram_formats.on_value_change(lambda event: update_carousel_slides_visibility(event.value))
                                instagram_hashtags = apply_field_props(
                                    ui.select(INSTAGRAM_HASHTAG_OPTIONS, label="Instagram hashtags (baseline)"),
                                )
                                ui.label("Hashtag counts are tuned per format automatically (Feed 10–15, Reel 3–5, Story ≤1, Carousel 5–10).").classes("mce-muted")
                                instagram_image_content_type = apply_field_props(
                                    ui.select(
                                        ["text / quote graphic", "character portrait", "abstract / mood graphic", "mixed"],
                                        value="text / quote graphic",
                                        label="Image content type",
                                    ),
                                )
                                instagram_hook = apply_field_props(
                                    ui.textarea(label="Hook / Opening line", placeholder="Opening line or tonal direction"),
                                    "outlined autogrow",
                                )
                                ui.label("Used as the caption's opening line only. The caption body is generated from your book/source content.").classes("mce-muted")
                                instagram_image_formats = apply_field_props(
                                    ui.select(list(INSTAGRAM_FORMAT_OPTIONS), multiple=True, value=list(INSTAGRAM_FORMAT_OPTIONS), label="Generate images for"),
                                    "outlined dense use-chips clearable",
                                )
                                ui.label("Toggle which formats get an image. Remove a chip to skip that format; clear all to skip image generation.").classes("mce-muted")

                                def instagram_image_format_selection() -> list[str]:
                                    return [f for f in (instagram_formats.value or []) if f in (instagram_image_formats.value or [])]
                                # BUG-IG-A: formats auto-map to aspect ratios (Feed 4:5, Reel/Story 9:16,
                                # Carousel slides), so this extra override is hidden to avoid a redundant
                                # second format picker. Kept in code for the rare edge-case override.
                                instagram_visual_formats = apply_field_props(
                                    ui.select(INSTAGRAM_VISUAL_FORMAT_OPTIONS, multiple=True, label="Extra image formats (optional override)"),
                                    "outlined dense use-chips clearable",
                                )
                                instagram_visual_formats.visible = False
                                instagram_visual_style = apply_field_props(
                                    ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Gothic", label="Visual style"),
                                )
                                instagram_hook_status = ui.label("").classes("mce-muted")

                                async def suggest_instagram_hook() -> None:
                                    set_field_status(instagram_hook_status, "loading", "Suggesting hook…")
                                    prompt = f"""Suggest one short Instagram hook / first line for Tell Tales Ink.

Topic: {topic.value}
Related book/source: {related_book.value}
Selected formats: {normalize_selected(instagram_formats.value)}
Audience: {normalize_selected(audience.value)}
Objective: {normalize_selected(social_objectives.value)}

Return only the hook. Keep it under 14 words. Make it bookish, specific, and non-generic."""
                                    try:
                                        instagram_hook.value = (await asyncio.to_thread(generate_text, prompt)).strip().splitlines()[0][:120]
                                        set_field_status(instagram_hook_status, "success", "Hook suggested.")
                                    except Exception as exc:
                                        set_field_status(instagram_hook_status, "error", f"Could not suggest hook: {exc}")

                                make_secondary_button("Suggest Hook", suggest_instagram_hook)
                                instagram_image_path = readonly_input("Instagram image", "")
                                instagram_image_status = ui.label("No image selected").classes("mce-muted")
                                instagram_image_thumb = ui.image("").classes("mce-upload-thumb")
                                instagram_image_thumb.set_visibility(False)

                                async def handle_instagram_image_upload(event) -> None:
                                    await save_social_image(event, instagram_image_path, instagram_image_status, "Instagram", instagram_image_thumb)

                                ui.upload(on_upload=handle_instagram_image_upload, auto_upload=True).props(
                                    'accept=".png,.jpg,.jpeg,.webp" max-files=1 label="Upload Instagram image"'
                                ).classes("w-full")
                                attach_asset_picker(instagram_image_path, instagram_image_status, instagram_image_thumb)

                            with ui.card().classes("mce-subcard") as linkedin_card:
                                subcard_heading("LinkedIn Settings")
                                linkedin_formats = apply_field_props(
                                    ui.select(LINKEDIN_CONTENT_FORMAT_OPTIONS, multiple=True, label="LinkedIn formats"),
                                    "outlined dense use-chips clearable",
                                )
                                linkedin_angle = apply_field_props(
                                    ui.textarea(label="LinkedIn angle", placeholder="Professional angle or framing"),
                                    "outlined autogrow",
                                )
                                linkedin_cta = apply_field_props(
                                    ui.input(label="LinkedIn CTA", placeholder="Connect, comment, click, or follow-up"),
                                )
                                linkedin_image_path = readonly_input("LinkedIn image", "")
                                linkedin_image_status = ui.label("No image selected").classes("mce-muted")
                                linkedin_image_thumb = ui.image("").classes("mce-upload-thumb")
                                linkedin_image_thumb.set_visibility(False)

                                async def handle_linkedin_image_upload(event) -> None:
                                    await save_social_image(event, linkedin_image_path, linkedin_image_status, "LinkedIn", linkedin_image_thumb)

                                ui.upload(on_upload=handle_linkedin_image_upload, auto_upload=True).props(
                                    'accept=".png,.jpg,.jpeg,.webp" max-files=1 label="Upload LinkedIn image"'
                                ).classes("w-full")
                                attach_asset_picker(linkedin_image_path, linkedin_image_status, linkedin_image_thumb)

                            with ui.card().classes("mce-subcard") as youtube_card:
                                subcard_heading("YouTube Settings")
                                youtube_formats = apply_field_props(
                                    ui.select(YOUTUBE_CONTENT_FORMAT_OPTIONS, multiple=True, label="YouTube formats"),
                                    "outlined dense use-chips clearable",
                                )
                                youtube_keywords = apply_field_props(
                                    ui.input(label="YouTube keywords", placeholder="Primary search keywords"),
                                )
                                youtube_link = apply_field_props(
                                    ui.input(label="Video or channel link", placeholder="Optional URL to include"),
                                )
                                youtube_image_path = readonly_input("YouTube image", "")
                                youtube_image_status = ui.label("No image selected").classes("mce-muted")
                                youtube_image_thumb = ui.image("").classes("mce-upload-thumb")
                                youtube_image_thumb.set_visibility(False)

                                async def handle_youtube_image_upload(event) -> None:
                                    await save_social_image(event, youtube_image_path, youtube_image_status, "YouTube", youtube_image_thumb)

                                ui.upload(on_upload=handle_youtube_image_upload, auto_upload=True).props(
                                    'accept=".png,.jpg,.jpeg,.webp" max-files=1 label="Upload YouTube image"'
                                ).classes("w-full")
                                attach_asset_picker(youtube_image_path, youtube_image_status, youtube_image_thumb)

                            with ui.card().classes("mce-subcard") as quote_card:
                                subcard_heading("Quote Post Settings")
                                quote_book = apply_field_props(
                                    ui.select(QUOTE_BOOK_OPTIONS, value=QUOTE_BOOK_OPTIONS[0] if QUOTE_BOOK_OPTIONS else None, label="Quote book/source"),
                                )
                                quote_moods = apply_field_props(
                                    ui.select(QUOTE_MOOD_TAGS, multiple=True, label="Quote mood/category tags"),
                                    "outlined dense use-chips clearable",
                                )
                                character_tags = apply_field_props(
                                    ui.select(CHARACTER_TAGS, multiple=True, label="Character tags"),
                                    "outlined dense use-chips clearable",
                                )
                                quote_visual_formats = apply_field_props(
                                    ui.select(quote_graphic_format_options(), multiple=True, label="Quote image destination formats"),
                                    "outlined dense use-chips clearable",
                                )
                                quote_visual_style = apply_field_props(
                                    ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Gothic", label="Quote image color style"),
                                )

                            with ui.card().classes("mce-subcard") as blog_card:
                                subcard_heading("Blog Post Settings")
                                blog_length = apply_field_props(
                                    ui.select(BLOG_LENGTH_OPTIONS, value=BLOG_LENGTH_OPTIONS[1] if len(BLOG_LENGTH_OPTIONS) > 1 else BLOG_LENGTH_OPTIONS[0], label="Blog length"),
                                )
                                blog_format = apply_field_props(
                                    ui.select(BLOG_FORMAT_OPTIONS, value=BLOG_FORMAT_OPTIONS[0], label="Blog format"),
                                )
                                blog_structure_options = apply_field_props(
                                    ui.select(BLOG_STRUCTURE_CHECKLIST_OPTIONS, multiple=True, label="Structure checklist"),
                                    "outlined dense use-chips clearable",
                                )
                                blog_sections = apply_field_props(
                                    ui.textarea(label="Optional structure notes", placeholder="Add any specific order, angle, or section notes"),
                                    "outlined autogrow",
                                )
                                blog_structure_status = ui.label("").classes("mce-muted")

                                async def suggest_blog_structure() -> None:
                                    set_field_status(blog_structure_status, "loading", "Suggesting structure…")
                                    prompt = f"""Suggest a concise blog structure for Tell Tales Ink.

Topic: {topic.value}
Related book/source: {related_book.value}
Blog format: {blog_format.value}
Length: {blog_length.value}
Selected checklist: {normalize_selected(blog_structure_options.value)}

Return only a practical section outline with 5-8 bullets."""
                                    try:
                                        blog_sections.value = await asyncio.to_thread(generate_text, prompt)
                                        set_field_status(blog_structure_status, "success", "Structure suggested.")
                                    except Exception as exc:
                                        set_field_status(blog_structure_status, "error", f"Could not suggest structure: {exc}")

                                make_secondary_button("Suggest Structure", suggest_blog_structure)
                                blog_seo_keywords = apply_field_props(
                                    ui.input(label="Blog SEO keywords", placeholder="Comma-separated keywords"),
                                )
                                blog_image_mode = apply_field_props(
                                    ui.select(["use uploaded image", "use selected character portrait", "text-only / quote graphic"], value="text-only / quote graphic", label="Blog image mode"),
                                )
                                blog_character = apply_field_props(
                                    ui.select(CHARACTER_TAGS, label="Selected character portrait"),
                                )
                                blog_visual_formats = apply_field_props(
                                    ui.select(BLOG_VISUAL_FORMAT_OPTIONS, multiple=True, label="Blog image formats"),
                                    "outlined dense use-chips clearable",
                                )
                                blog_visual_style = apply_field_props(
                                    ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Press", label="Blog visual style"),
                                )

                                def sync_blog_visual_defaults(event=None) -> None:
                                    if blog_format.value == "photo/gallery slideshow" and not blog_visual_formats.value:
                                        blog_visual_formats.value = BLOG_GALLERY_VISUAL_FORMAT_OPTIONS

                                blog_format.on_value_change(sync_blog_visual_defaults)
                                blog_image_path = readonly_input("Blog image", "")
                                blog_image_status = ui.label("No image selected").classes("mce-muted")
                                blog_image_thumb = ui.image("").classes("mce-upload-thumb")
                                blog_image_thumb.set_visibility(False)

                                async def handle_blog_image_upload(event) -> None:
                                    await save_social_image(event, blog_image_path, blog_image_status, "Blog", blog_image_thumb)

                                ui.upload(on_upload=handle_blog_image_upload, auto_upload=True).props(
                                    'accept=".png,.jpg,.jpeg,.webp" max-files=1 label="Upload blog image"'
                                ).classes("w-full")
                                attach_asset_picker(blog_image_path, blog_image_status, blog_image_thumb)

                            with ui.card().classes("mce-subcard") as newsletter_card:
                                subcard_heading("Newsletter Settings")
                                nl_subject = apply_field_props(
                                    ui.input(label="Newsletter subject line", placeholder="What is the subject line?"),
                                )
                                nl_preview = apply_field_props(
                                    ui.textarea(label="Preview text", placeholder="Short preview text shown in inboxes"),
                                    "outlined autogrow",
                                )
                                nl_structure = apply_field_props(
                                    ui.select(
                                        ["announcement", "roundup / digest", "personal note", "behind-the-scenes"],
                                        label="Newsletter structure",
                                    ),
                                )

                            with ui.card().classes("mce-subcard") as character_card:
                                subcard_heading("Character Spotlight Settings")
                                cs_character = apply_field_props(
                                    ui.select(CHARACTER_TAGS, label="Character spotlight subject"),
                                )
                                cs_platform_format = apply_field_props(
                                    ui.select(CHARACTER_SPOTLIGHT_PLATFORM_FORMATS, value=CHARACTER_SPOTLIGHT_PLATFORM_FORMATS[0], label="Destination format"),
                                )
                                cs_image_mode = apply_field_props(
                                    ui.select(CHARACTER_IMAGE_MODE_OPTIONS, value=CHARACTER_IMAGE_MODE_OPTIONS[1], label="Image source"),
                                )
                                cs_focus = apply_field_props(
                                    ui.textarea(label="Spotlight angle / focus", placeholder="Why this character matters, what to highlight"),
                                    "outlined autogrow",
                                )
                                cs_image_path = readonly_input("Character spotlight image", "")
                                cs_image_status = ui.label("No image selected").classes("mce-muted")
                                cs_image_thumb = ui.image("").classes("mce-upload-thumb")
                                cs_image_thumb.set_visibility(False)

                                async def handle_character_image_upload(event) -> None:
                                    await save_social_image(event, cs_image_path, cs_image_status, "Character Spotlight", cs_image_thumb)

                                ui.upload(on_upload=handle_character_image_upload, auto_upload=True).props(
                                    'accept=".png,.jpg,.jpeg,.webp" max-files=1 label="Upload character spotlight image"'
                                ).classes("w-full")
                                attach_asset_picker(cs_image_path, cs_image_status, cs_image_thumb)

                            with ui.card().classes("mce-subcard") as press_release_card:
                                subcard_heading("Press Release Settings")
                                pr_profile = apply_field_props(
                                    ui.select(press_profile_options(), label="Saved press profile"),
                                    "outlined dense clearable",
                                )
                                pr_profile_name = apply_field_props(
                                    ui.input(label="Profile name", placeholder="Author / publisher profile name"),
                                )

                                subcard_heading("Timing & Distribution")
                                pr_destination = readonly_input("Distribution", PRESS_RELEASE_DESTINATION)
                                with ui.element("div").classes("mce-two-col"):
                                    pr_timing = apply_field_props(ui.select(PRESS_RELEASE_TIMING_OPTIONS, label="Press release timing"))
                                    pr_embargo_date = apply_field_props(ui.input(label="Embargo date", placeholder="If embargoed, provide the date"))
                                    pr_release_date = apply_field_props(ui.input(label="Release date").props("type=date"))
                                    pr_city = apply_field_props(ui.input(label="Dateline city", placeholder="City for the dateline"))
                                    pr_state = apply_field_props(ui.input(label="Dateline state / region", placeholder="State or region"))

                                subcard_heading("Contact Info")
                                with ui.element("div").classes("mce-two-col"):
                                    pr_contact_name = apply_field_props(ui.input(label="Media contact name", placeholder="Contact person name"))
                                    pr_contact_title = apply_field_props(ui.input(label="Media contact title", placeholder="Contact person's title"))
                                    pr_organization = apply_field_props(ui.input(label="Organization / imprint", placeholder="Organization or imprint"))
                                    pr_contact_email = apply_field_props(ui.input(label="Media contact email", placeholder="contact@example.com"))
                                    pr_contact_phone = apply_field_props(ui.input(label="Media contact phone", placeholder="Phone number"))
                                    pr_website = apply_field_props(ui.input(label="Website", placeholder="Organization website"))

                                subcard_heading("Newswire Details")
                                pr_news_angle = apply_field_props(ui.textarea(label="News angle", placeholder="Why this is newsworthy"), "outlined autogrow")
                                pr_news_angle_status = ui.label("").classes("mce-muted")
                                pr_profile_actions = ui.row().classes("mce-actions")
                                pr_supporting_proof = apply_field_props(ui.textarea(label="Supporting proof", placeholder="Verified facts, awards, or evidence"), "outlined autogrow")
                                pr_quote_source = apply_field_props(ui.textarea(label="Quote source", placeholder="Approved spokesperson quote source"), "outlined autogrow")
                                pr_target_media = apply_field_props(ui.textarea(label="Target media", placeholder="Journalists, reviewers, outlets, or PR channels"), "outlined autogrow")

                                subcard_heading("Assets")
                                pr_required_assets = apply_field_props(
                                    ui.select(PRESS_RELEASE_COMPANION_ASSETS, multiple=True, label="Companion assets"),
                                    "outlined dense use-chips clearable",
                                )
                                ui.label("Optional assets to generate alongside the press release.").classes("mce-muted")
                                pr_asset_status = readonly_input("Generated PR assets", "")

                            with ui.card().classes("mce-subcard") as review_card:
                                subcard_heading("Review / Pull Quote Settings")
                                ui.label("Build a pull quote from a real review, an AI quote in the style of real reviews, or your own pasted quote.").classes("mce-muted")
                                review_quote_mode = apply_field_props(
                                    ui.select(
                                        ["From my quote bank", "AI-generated in style of real review", "I'll paste it manually"],
                                        value="From my quote bank",
                                        label="Quote source",
                                    ),
                                )
                                with ui.column().classes("w-full mce-stack") as review_bank_group:
                                    review_book = apply_field_props(
                                        ui.select(
                                            review_book_options(),
                                            value=review_book_options()[0] if review_book_options() else None,
                                            label="Review book/source",
                                        ),
                                    )
                                    review_source = apply_field_props(
                                        ui.select([], label="Review source / outlet"),
                                        "outlined dense clearable",
                                    )
                                    review_quote = apply_field_props(
                                        ui.select([], label="Grounded review quote"),
                                        "outlined dense clearable",
                                    )
                                review_manual_quote = apply_field_props(
                                    ui.textarea(label="Paste the quote you want to use", placeholder="Paste the exact pull-quote text here"),
                                    "outlined autogrow",
                                )
                                review_manual_quote.set_visibility(False)
                                review_mode = apply_field_props(
                                    ui.select(
                                        ["critic pull quote", "reader praise", "short testimonial", "media kit blurb"],
                                        value="critic pull quote",
                                        label="Quote style",
                                    ),
                                )
                                review_attribution = apply_field_props(
                                    ui.input(label="Reviewer / source attribution", placeholder="e.g. Goodreads reviewer, Literary Titan"),
                                )
                                review_image_type = apply_field_props(
                                    ui.select(
                                        ["Text-only card", "Character image + quote overlay", "Book cover + quote overlay"],
                                        value="Text-only card",
                                        label="Image type",
                                    ),
                                )
                                review_instagram_formats = apply_field_props(
                                    ui.select(["Feed post", "Story", "Reel"], multiple=True, value=["Feed post"], label="Instagram format"),
                                    "outlined dense use-chips clearable",
                                )
                                review_card_style = apply_field_props(
                                    ui.select(["Dark gothic", "Light minimal", "Brand default"], value="Dark gothic", label="Card style"),
                                )
                                review_promo_line = apply_field_props(
                                    ui.input(label="Brand promo line", placeholder="Optional tagline shown under the quote"),
                                )
                                review_visual_formats = apply_field_props(
                                    ui.select(quote_graphic_format_options(), multiple=True, label="Extra image formats (optional override)"),
                                    "outlined dense use-chips clearable",
                                )
                                review_visual_style = apply_field_props(
                                    ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Press", label="Visual style"),
                                )
                                review_image_path = readonly_input("Review pull quote image", "")
                                review_image_status = ui.label("No image selected").classes("mce-muted")
                                review_image_thumb = ui.image("").classes("mce-upload-thumb")
                                review_image_thumb.set_visibility(False)

                                async def handle_review_image_upload(event) -> None:
                                    await save_social_image(event, review_image_path, review_image_status, "Review pull quote", review_image_thumb)

                                ui.upload(on_upload=handle_review_image_upload, auto_upload=True).props(
                                    'accept=".png,.jpg,.jpeg,.webp" max-files=1 label="Upload pull quote image"'
                                ).classes("w-full")
                                attach_asset_picker(review_image_path, review_image_status, review_image_thumb)

                                def trigger_review_preview() -> None:
                                    refresh = preview_refresh.get("fn")
                                    if callable(refresh):
                                        refresh()

                                def update_review_quotes() -> None:
                                    options = review_quotes_for_source(review_book.value or "", review_source.value or "")
                                    review_quote.options = options
                                    if review_quote.value not in options:
                                        review_quote.value = options[0] if options else None
                                    review_quote.update()
                                    if review_source.value and not (review_attribution.value or "").strip():
                                        review_attribution.value = review_source.value
                                    trigger_review_preview()

                                def update_review_sources() -> None:
                                    options = review_source_options(review_book.value or "")
                                    review_source.options = options
                                    review_source.value = options[0] if options else None
                                    review_source.update()
                                    update_review_quotes()

                                def update_review_mode_visibility() -> None:
                                    mode = review_quote_mode.value or ""
                                    review_bank_group.set_visibility(mode != "I'll paste it manually")
                                    review_manual_quote.set_visibility(mode == "I'll paste it manually")
                                    trigger_review_preview()

                                review_book.on_value_change(lambda _e: update_review_sources())
                                review_source.on_value_change(lambda _e: update_review_quotes())
                                review_quote.on_value_change(lambda _e: trigger_review_preview())
                                review_quote_mode.on_value_change(lambda _e: update_review_mode_visibility())
                                review_manual_quote.on_value_change(lambda _e: trigger_review_preview())
                                review_instagram_formats.on_value_change(lambda _e: trigger_review_preview())
                                review_image_type.on_value_change(lambda _e: trigger_review_preview())
                                review_card_style.on_value_change(lambda _e: trigger_review_preview())
                                update_review_sources()
                                update_review_mode_visibility()

                            generator_actions = ui.row().classes("mce-actions")

                    with ui.card().classes("mce-card mce-sticky"):
                        section_heading("Output Studio", "Preview the draft, trace artifact paths, and compare against a clean baseline.")
                        with ui.column().classes("mce-stack w-full"):
                            status = readonly_input("Status", "Ready", mono=False)
                            generation_progress = ui.linear_progress(value=0, show_value=False).classes("w-full")
                            generation_progress.visible = False
                            generation_progress_label = ui.label("Ready").classes("mce-muted")
                            platform_preview = ui.html(
                                build_preview_html(content_type.value, {"draft": "Select a content type to see a platform-style preview."}),
                                sanitize=False,
                            ).classes("mce-preview")
                            generated_output = ui.textarea(label="Generated content (editable)", value="").props("outlined autogrow").classes("w-full mce-output-textarea")

                            # Registry of generator form fields, used to snapshot the form into
                            # a saved draft and restore it when the draft is re-opened in the
                            # builder. Built from locals() so any field not defined for a content
                            # type is simply skipped (never a NameError).
                            # Every per-type widget that should survive a draft save/reopen.
                            # Order matters where a widget's on_value_change resets others:
                            #  - blog_format resets blog visual defaults (sync_blog_visual_defaults),
                            #    so it precedes blog_visual_*.
                            #  - review_book -> review_source -> review_quote cascade, so review_quote
                            #    is restored last; review_attribution after review_source (which can
                            #    auto-fill it). Names that aren't live widgets are skipped safely.
                            _generator_field_names = [
                                "topic", "related_book", "platform", "social_objectives", "audience",
                                "cta", "constraints",
                                # Instagram
                                "instagram_formats", "instagram_carousel_slides", "instagram_hashtags",
                                "instagram_image_content_type", "instagram_hook",
                                "instagram_image_formats", "instagram_visual_formats", "instagram_visual_style",
                                # LinkedIn
                                "linkedin_formats", "linkedin_angle", "linkedin_cta",
                                # YouTube
                                "youtube_formats", "youtube_keywords", "youtube_link",
                                # Quote post
                                "quote_book", "quote_moods", "character_tags",
                                "quote_visual_formats", "quote_visual_style",
                                # Blog (blog_format before blog_visual_* due to its reset cascade)
                                "blog_length", "blog_format", "blog_structure_options", "blog_sections",
                                "blog_seo_keywords", "blog_image_mode",
                                "blog_character", "blog_visual_formats", "blog_visual_style",
                                # Character spotlight
                                "cs_character", "cs_focus", "cs_platform_format", "cs_image_mode",
                                # Review / pull quote (review_book->source->quote cascade order)
                                "review_mode", "review_quote_mode", "review_book", "review_source",
                                "review_attribution", "review_promo_line", "review_image_type",
                                "review_instagram_formats", "review_card_style",
                                "review_visual_formats", "review_visual_style",
                                "review_manual_quote", "review_quote",
                                # Newsletter (real widget names are nl_*, not newsletter_*)
                                "nl_subject", "nl_preview", "nl_structure",
                                # Press release
                                "pr_timing", "pr_embargo_date", "pr_release_date", "pr_city", "pr_state",
                                "pr_contact_name", "pr_contact_title", "pr_organization",
                                "pr_contact_email", "pr_contact_phone", "pr_website",
                                "pr_news_angle", "pr_supporting_proof", "pr_quote_source",
                                "pr_target_media", "pr_required_assets",
                            ]
                            _generator_local_scope = locals()
                            generator_fields = {
                                name: _generator_local_scope[name]
                                for name in _generator_field_names
                                if name in _generator_local_scope
                                and hasattr(_generator_local_scope[name], "value")
                            }
                            generator_fields_snapshot = lambda: {
                                name: widget.value for name, widget in generator_fields.items()
                            }
                            with ui.card().classes("mce-subcard w-full"):
                                subcard_heading("Rendered preview")
                                generated_output_rendered = ui.markdown("").classes("w-full mce-rendered-output")
                                generated_output_rendered.bind_content_from(generated_output, "value")
                            result_actions = ui.row().classes("mce-actions")
                            with ui.card().classes("mce-subcard") as generated_images_card:
                                subcard_heading("Generated images")
                                ui.label("Rendered image assets for this draft. Use Download Images / Assets for the full package.").classes("mce-muted")
                                generated_images_gallery = ui.row().classes("w-full mce-gallery")
                            generated_images_card.visible = False
                            instagram_format_images: dict[str, list[str]] = {}
                            with ui.card().classes("mce-subcard") as instagram_regen_card:
                                subcard_heading("Per-format images")
                                ui.label("Regenerate the image for a single Instagram format without rebuilding the draft.").classes("mce-muted")
                                instagram_regen_actions = ui.row().classes("mce-actions")
                            instagram_regen_card.visible = False
                            with ui.card().classes("mce-subcard") as quote_image_builder:
                                subcard_heading("Quote Image Builder")
                                ui.label("Select the quote cards you want to render, then generate platform-ready image mockups.").classes("mce-muted")
                                quote_candidate_select = apply_field_props(
                                    ui.select([], multiple=True, label="Quote cards to render"),
                                    "outlined dense use-chips clearable",
                                )
                                quote_candidate_select.visible = False
                                quote_candidate_cards = ui.column().classes("w-full mce-stack")
                                quote_candidate_actions = ui.row().classes("mce-actions")
                                quote_image_status = readonly_input("Quote image status", "Generate a quote draft first.")
                                quote_platform_mockup = ui.html(quote_platform_mockup_html([]), sanitize=False).classes("w-full")
                                quote_image_actions = ui.row().classes("mce-actions")

                            with ui.expansion("Artifact paths and structured brief", icon="inventory_2").classes("mce-expansion"):
                                with ui.column().classes("mce-stack w-full"):
                                    filtered_context_path = readonly_input("Filtered context", "")
                                    prompt_path = readonly_input("Generation prompt", "")
                                    draft_path = readonly_input("Draft output", "")
                                    saved_draft_path = readonly_input("Saved draft", "")
                                    generated_visual_path = readonly_input("Generated image preview", "")
                                    generated_visual_package_path = readonly_input("Generated image package", "")
                                    structured_brief_output = readonly_textarea("Structured brief used for comparison", "")

                            with ui.expansion("Compare Tell Tales Ink vs ChatGPT", icon="compare_arrows").classes("mce-expansion"):
                                with ui.column().classes("mce-stack w-full"):
                                    ui.label("Charisma. Uniqueness. Nerve and Talent.").classes("mce-section-title")
                                    ui.label("Generate a Tell Tales Ink draft first, then create a fresh baseline from the same selections.").classes("mce-muted")
                                    chatgpt_model = apply_field_props(
                                        ui.select(
                                            CHATGPT_COMPARISON_MODEL_OPTIONS,
                                            value=CHATGPT_COMPARISON_MODEL_OPTIONS[0],
                                            label="ChatGPT model",
                                        )
                                    )
                                    comparison_status = readonly_input("Comparison status", "Generate a Tell Tales Ink draft first.", mono=False)
                                    comparison_view = ui.html(comparison_side_by_side_html(), sanitize=False).classes("w-full")
                                    chatgpt_output = readonly_textarea("ChatGPT baseline output", "")
                                    chatgpt_prompt_path = readonly_input("ChatGPT prompt path", "")
                                    chatgpt_draft_path = readonly_input("ChatGPT draft path", "")
                                    comparison_preference = apply_field_props(
                                        ui.select(["Tell Tales Ink", "ChatGPT", "Tie / needs revision"], label="Which result do you prefer?")
                                    )
                                    comparison_notes = apply_field_props(
                                        ui.textarea(
                                            label="Judge notes",
                                            placeholder="Example: Tell Tales Ink used the brand world more specifically; ChatGPT was clean but generic.",
                                        ),
                                        "outlined autogrow",
                                    )
                                    comparison_save_status = readonly_input("Saved judgment status", "")
                                    comparison_vote_path = readonly_input("Saved judgment path", "")
                                    comparison_actions = ui.row().classes("mce-actions")

                    def current_uploaded_image_path() -> str:
                        if content_type.value in {"instagram_caption", "blog_post", "character_spotlight", "quote_post"} and generated_visual_path.value:
                            generated_path = Path(str(generated_visual_path.value))
                            if generated_path.exists():
                                return str(generated_path)
                        if (
                            content_type.value == "blog_post"
                            and "selected character portrait" in str(blog_image_mode.value or "").lower()
                            and blog_character.value
                        ):
                            portrait = resolve_character_portrait_asset(blog_character.value, book=related_book.value)
                            if portrait:
                                return str(portrait)
                        return str(
                            {
                                "instagram_caption": instagram_image_path.value,
                                "linkedin_content": linkedin_image_path.value,
                                "youtube_content": youtube_image_path.value,
                                "blog_post": blog_image_path.value,
                                "character_spotlight": cs_image_path.value,
                            }.get(content_type.value, "")
                            or ""
                        )

                    def preview_fields() -> dict[str, object]:
                        return {
                            "caption": topic.value,
                            "draft": generated_output.value or topic.value,
                            "instagram_formats": instagram_formats.value,
                            "instagram_hashtags": instagram_hashtags.value,
                            "instagram_hook": instagram_hook.value,
                            "instagram_visual_formats": instagram_visual_formats.value,
                            "instagram_visual_style": instagram_visual_style.value,
                            "linkedin_formats": linkedin_formats.value,
                            "linkedin_angle": linkedin_angle.value,
                            "linkedin_cta": linkedin_cta.value,
                            "youtube_formats": youtube_formats.value,
                            "youtube_keywords": youtube_keywords.value,
                            "youtube_link": youtube_link.value,
                            "quote_moods": quote_moods.value,
                            "quote_characters": character_tags.value,
                            "quote_source": quote_book.value,
                            "generated_visual_package_path": generated_visual_package_path.value,
                            "blog_format": blog_format.value,
                            "blog_length": blog_length.value,
                            "blog_structure": blog_sections.value,
                            "blog_structure_checklist": blog_structure_options.value,
                            "blog_image_mode": blog_image_mode.value,
                            "blog_character": blog_character.value,
                            "blog_visual_formats": blog_visual_formats.value,
                            "blog_visual_style": blog_visual_style.value,
                            "character_subject": cs_character.value,
                            "character_focus": cs_focus.value,
                            "character_format": cs_platform_format.value,
                            "character_image_mode": cs_image_mode.value,
                            "newsletter_subject": nl_subject.value,
                            "newsletter_preview": nl_preview.value,
                            "newsletter_structure": nl_structure.value,
                            "review_mode": review_mode.value,
                            "review_quote_text": (
                                review_manual_quote.value
                                if (review_quote_mode.value == "I'll paste it manually" and str(review_manual_quote.value or "").strip())
                                else review_quote.value or generated_output.value or topic.value
                            ),
                            "review_source": review_source.value,
                            "review_attribution": review_attribution.value or review_source.value,
                            "review_promo_line": review_promo_line.value,
                            "review_instagram_formats": review_instagram_formats.value,
                            "review_card_style": review_card_style.value,
                            "review_image_type": review_image_type.value,
                            "pr_city": pr_city.value,
                            "pr_state": pr_state.value,
                            "pr_release_date": pr_release_date.value,
                            "pr_contact_name": pr_contact_name.value,
                            "pr_website": pr_website.value,
                            "pr_news_angle": pr_news_angle.value,
                            "pr_required_assets": pr_required_assets.value,
                            "image_path": current_uploaded_image_path(),
                            "instagram_format_images": dict(instagram_format_images),
                        }

                    def refresh_platform_preview() -> None:
                        platform_preview.content = build_preview_html(content_type.value, preview_fields())

                    def refresh_generated_images_gallery() -> None:
                        """Show thumbnails of the image files rendered for this draft (BUG-IMG-02)."""
                        generated_images_gallery.clear()
                        primary = str(generated_visual_path.value or "").strip()
                        image_paths: list[str] = []
                        if primary and Path(primary).exists():
                            # All images of one generation land in a single timestamped folder.
                            folder = Path(primary).parent
                            image_paths = sorted(
                                str(p) for p in folder.glob("*")
                                if p.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
                            )
                        generated_images_card.visible = bool(image_paths)
                        if not image_paths:
                            return
                        with generated_images_gallery:
                            for path in image_paths[:12]:
                                ui.image(path).classes("mce-gallery-thumb")

                    async def regenerate_instagram_format(fmt: str) -> None:
                        labels = instagram_visual_formats_for_post([fmt], instagram_visual_formats.value, int(instagram_carousel_slides.value or 5))
                        status.value = f"Regenerating {fmt} image…"
                        try:
                            package = await asyncio.to_thread(
                                render_generated_visual_package,
                                content_type="instagram_caption",
                                content=generated_output.value or topic.value,
                                topic=topic.value,
                                format_labels=labels,
                                theme_name=instagram_visual_style.value or "Gothic",
                                image_content_type=instagram_image_content_type.value,
                            )
                        except Exception as exc:
                            status.value = f"Regeneration failed: {exc}"
                            ui.notify(status.value, type="negative")
                            return
                        paths = [str(p) for p in package.get("paths", [])] if isinstance(package, dict) else []
                        if paths:
                            instagram_format_images[fmt.lower()] = paths
                            status.value = f"Regenerated {fmt} ({len(paths)} image file(s))."
                            ui.notify(status.value, type="positive")
                        else:
                            error = str(package.get("error", "")) if isinstance(package, dict) else ""
                            status.value = error or f"No image generated for {fmt}."
                            ui.notify(status.value, type="warning")
                        refresh_platform_preview()
                        refresh_generated_images_gallery()

                    def refresh_instagram_regen_actions() -> None:
                        instagram_regen_actions.clear()
                        has_draft = bool(str(generated_output.value or "").strip())
                        formats = (
                            instagram_image_format_selection()
                            if (content_type.value == "instagram_caption" and has_draft)
                            else []
                        )
                        instagram_regen_card.visible = bool(formats)
                        if not formats:
                            return
                        with instagram_regen_actions:
                            for fmt in formats:
                                make_secondary_button(
                                    f"Regenerate {fmt}",
                                    lambda f=fmt: regenerate_instagram_format(f),
                                )

                    def refresh_press_profile_options() -> None:
                        pr_profile.set_options(press_profile_options(), value=pr_profile.value if pr_profile.value in press_profile_options() else None)

                    def apply_press_profile() -> None:
                        profile_name = pr_profile.value
                        profiles = load_press_profiles()
                        data = profiles.get(profile_name or "")
                        if not data:
                            ui.notify("Select a saved press profile first.", type="warning")
                            return
                        pr_profile_name.value = profile_name
                        pr_contact_name.value = data.get("contact_name", "")
                        pr_contact_title.value = data.get("contact_title", "")
                        pr_organization.value = data.get("organization", "")
                        pr_contact_email.value = data.get("email", "")
                        pr_contact_phone.value = data.get("phone", "")
                        pr_website.value = data.get("website", "")
                        ui.notify("Press profile loaded.", type="positive")
                        refresh_platform_preview()

                    def save_press_profile_from_fields() -> None:
                        name = str(pr_profile_name.value or pr_contact_name.value or pr_organization.value or "").strip()
                        if not name:
                            ui.notify("Add a profile name before saving.", type="warning")
                            return
                        profiles = load_press_profiles()
                        profiles[name] = {
                            "contact_name": pr_contact_name.value,
                            "contact_title": pr_contact_title.value,
                            "organization": pr_organization.value,
                            "email": pr_contact_email.value,
                            "phone": pr_contact_phone.value,
                            "website": pr_website.value,
                        }
                        save_press_profiles(profiles)
                        refresh_press_profile_options()
                        pr_profile.value = name
                        ui.notify("Press profile saved.", type="positive")

                    async def suggest_press_news_angle() -> None:
                        set_field_status(pr_news_angle_status, "loading", "Suggesting news angle…")
                        prompt = f"""Suggest 3 strong, journalist-friendly news angles for this press release.

Topic: {topic.value}
Related book/source: {related_book.value}
Supporting proof: {pr_supporting_proof.value}
Target media: {pr_target_media.value}

Return concise angle options with why each is newsworthy."""
                        try:
                            pr_news_angle.value = await asyncio.to_thread(generate_text, prompt)
                            set_field_status(pr_news_angle_status, "success", "News angle suggested.")
                            ui.notify("News angle suggested.", type="positive")
                        except Exception as exc:
                            set_field_status(pr_news_angle_status, "error", f"Could not suggest news angle: {exc}")
                            ui.notify(f"Could not suggest news angle: {exc}", type="negative")
                        refresh_platform_preview()

                    def quote_attribution_value() -> str:
                        selected = normalize_selected(character_tags.value)
                        return "" if selected == "Not specified" else selected

                    quote_candidate_checks = []

                    def selected_quote_card_values() -> list[str]:
                        selected = [quote for quote, checkbox in quote_candidate_checks if checkbox.value]
                        if selected:
                            return selected
                        fallback = quote_candidate_select.value or []
                        return [fallback] if isinstance(fallback, str) else list(fallback)

                    def sync_quote_select_from_cards() -> None:
                        selected = [quote for quote, checkbox in quote_candidate_checks if checkbox.value]
                        quote_candidate_select.value = selected
                        quote_image_status.value = f"Selected {len(selected)} quote card(s)."

                    def render_quote_candidate_cards(candidates: list[str], selected: list[str]) -> None:
                        quote_candidate_cards.clear()
                        quote_candidate_checks.clear()
                        selected_set = set(selected)
                        with quote_candidate_cards:
                            for index, candidate in enumerate(candidates, start=1):
                                with ui.card().classes("mce-subcard w-full"):
                                    checkbox = ui.checkbox(f"Quote card {index}", value=candidate in selected_set)
                                    ui.label(candidate).classes("mce-output-text")
                                    quote_candidate_checks.append((candidate, checkbox))
                                    checkbox.on_value_change(lambda _event: sync_quote_select_from_cards())

                    def refresh_quote_candidates(select_all: bool = True) -> None:
                        candidates = extract_quote_candidates(generated_output.value)
                        selected = candidates if select_all else []
                        quote_candidate_select.set_options(candidates, value=selected)
                        render_quote_candidate_cards(candidates, selected)
                        if candidates:
                            quote_image_status.value = f"{len(candidates)} quote card option(s) ready."
                            quote_platform_mockup.content = quote_platform_mockup_html([])
                        else:
                            quote_image_status.value = "No quote card candidates found yet."
                            quote_platform_mockup.content = quote_platform_mockup_html(
                                [],
                                error="No selectable quote cards were found in the generated draft. Regenerate with clearer quote options.",
                            )

                    def select_all_quote_candidates() -> None:
                        options = list(quote_candidate_select.options or [])
                        quote_candidate_select.value = options
                        for _, checkbox in quote_candidate_checks:
                            checkbox.value = True
                        quote_image_status.value = f"Selected {len(options)} quote card(s)."

                    def select_no_quote_candidates() -> None:
                        quote_candidate_select.value = []
                        for _, checkbox in quote_candidate_checks:
                            checkbox.value = False
                        quote_image_status.value = "No quote cards selected."

                    async def generate_selected_quote_images() -> None:
                        selected_quotes = selected_quote_card_values()
                        if not selected_quotes:
                            quote_image_status.value = "Select at least one quote card first."
                            ui.notify(quote_image_status.value, type="warning")
                            return
                        quote_image_status.value = "Generating quote card images..."
                        is_review = content_type.value == "review_pull_quote"
                        if is_review:
                            format_labels = review_visual_formats_for_instagram(review_instagram_formats.value, review_visual_formats.value)
                            theme = review_visual_style.value or "Press"
                            attribution = review_attribution.value or quote_attribution_value()
                        else:
                            format_labels = quote_visual_formats.value or ["Instagram Post (4:5)"]
                            theme = quote_visual_style.value or "Gothic"
                            attribution = quote_attribution_value()
                        try:
                            package = await asyncio.to_thread(
                                render_selected_quote_visual_package,
                                quotes=list(selected_quotes),
                                attribution=attribution,
                                format_labels=format_labels,
                                theme_name=theme,
                                book=(review_book.value if is_review else quote_book.value),
                            )
                        except Exception as exc:
                            quote_image_status.value = f"Quote image generation failed: {exc}"
                            ui.notify(quote_image_status.value, type="negative")
                            return

                        paths = [str(path) for path in package.get("paths", [])] if isinstance(package, dict) else []
                        error = str(package.get("error", "")) if isinstance(package, dict) else ""
                        package_path = str(package.get("zip_path", "")) if isinstance(package, dict) else ""
                        generated_visual_path.value = paths[0] if paths else error
                        generated_visual_package_path.value = package_path
                        quote_platform_mockup.content = quote_platform_mockup_html(paths, package_path, error)
                        quote_image_status.value = f"Generated {len(paths)} image file(s)." if paths else error or "No images generated."
                        refresh_platform_preview()
                        ui.notify(quote_image_status.value, type="positive" if paths else "warning")

                    preview_refresh["fn"] = refresh_platform_preview
                    instagram_image_formats.on_value_change(lambda _event: refresh_instagram_regen_actions())
                    # BUG-EDIT-01: edits to the generated content re-render the preview/mock live.
                    generated_output.on_value_change(lambda _event: refresh_platform_preview())

                    def apply_visibility(selected_type: str) -> None:
                        is_press_release = selected_type == "press_release"
                        instagram_card.visible = selected_type == "instagram_caption"
                        linkedin_card.visible = selected_type == "linkedin_content"
                        youtube_card.visible = selected_type == "youtube_content"
                        quote_card.visible = selected_type == "quote_post"
                        blog_card.visible = selected_type == "blog_post"
                        newsletter_card.visible = selected_type == "newsletter_blurb"
                        character_card.visible = selected_type == "character_spotlight"
                        review_card.visible = selected_type == "review_pull_quote"
                        press_release_card.visible = is_press_release
                        quote_image_builder.visible = selected_type in {"quote_post", "review_pull_quote"}
                        # BUG-QP-01: quote_post has its own "Quote book/source" field, so hide
                        # the duplicate top-level Related book/source for it.
                        related_book.visible = selected_type != "quote_post"
                        platform.visible = selected_type not in PLATFORM_IMPLIED_CONTENT_TYPES
                        hide_social_audience = selected_type in HIDE_SOCIAL_AUDIENCE_CONTENT_TYPES
                        social_objectives.visible = not hide_social_audience
                        audience.visible = not hide_social_audience
                        is_newsletter = selected_type == "newsletter_blurb"
                        objective_options = NEWSLETTER_OBJECTIVE_OPTIONS if is_newsletter else SOCIAL_OBJECTIVES
                        if list(social_objectives.options) != list(objective_options):
                            social_objectives.options = list(objective_options)
                            social_objectives.value = [v for v in (social_objectives.value or []) if v in objective_options]
                            social_objectives.props(f'label="{"Email objectives" if is_newsletter else "Social objectives"}"')
                            social_objectives.update()
                        if is_press_release:
                            platform.value = PRESS_RELEASE_DESTINATION
                            social_objectives.value = []
                            audience.value = []
                        if not is_press_release and platform.value not in PLATFORM_OPTIONS:
                            platform.value = PLATFORM_OPTIONS[0]
                        platform_preview.content = build_preview_html(selected_type, preview_fields())
                        refresh_instagram_regen_actions()
                        refresh_generated_images_gallery()

                    def sync_platform_default(selected_type: str) -> None:
                        default_platform = PLATFORM_BY_CONTENT_TYPE.get(selected_type)
                        if default_platform:
                            platform.value = default_platform

                    def on_content_type_change(event) -> None:
                        sync_platform_default(event.value)
                        apply_visibility(event.value)

                    content_type.on_value_change(on_content_type_change)
                    sync_platform_default(content_type.value)
                    apply_visibility(content_type.value)

                    async def generate_content() -> None:
                        started_at = datetime.now()
                        instagram_format_images.clear()
                        content_asset_total = 1
                        if content_type.value == "press_release" and pr_required_assets.value:
                            content_asset_total += len(pr_required_assets.value)
                        elif content_type.value == "quote_post":
                            content_asset_total += len(quote_visual_formats.value or [])
                        elif content_type.value == "review_pull_quote":
                            content_asset_total += len(review_visual_formats_for_instagram(review_instagram_formats.value, review_visual_formats.value))
                        elif content_type.value == "instagram_caption" and instagram_image_format_selection():
                            content_asset_total += len(instagram_visual_formats_for_post(instagram_image_format_selection(), instagram_visual_formats.value, instagram_carousel_slides.value))
                        elif content_type.value == "blog_post":
                            content_asset_total += len(
                                blog_visual_formats.value
                                or (
                                    BLOG_GALLERY_VISUAL_FORMAT_OPTIONS
                                    if blog_format.value == "photo/gallery slideshow"
                                    else BLOG_VISUAL_FORMAT_OPTIONS
                                )
                            )
                        elif content_type.value == "character_spotlight" and cs_image_mode.value != "text-only highlight card":
                            content_asset_total += len(formats_for_character_spotlight(cs_platform_format.value))
                        generation_progress.visible = True
                        generation_progress.value = 0
                        status.value = "Generating..."
                        try:
                            await set_generation_progress(
                                progress=generation_progress,
                                label=generation_progress_label,
                                status=status,
                                started_at=started_at,
                                completed_steps=0,
                                total_steps=4,
                                completed_assets=0,
                                total_assets=content_asset_total,
                                message="Preparing content brief...",
                            )
                            apply_visibility(content_type.value)
                            if content_type.value == "instagram_caption" and not str(instagram_hook.value or "").strip():
                                await set_generation_progress(
                                    progress=generation_progress,
                                    label=generation_progress_label,
                                    status=status,
                                    started_at=started_at,
                                    completed_steps=1,
                                    total_steps=4,
                                    completed_assets=0,
                                    total_assets=content_asset_total,
                                    message="Suggesting Instagram hook...",
                                )
                                await suggest_instagram_hook()
                            selected_blog_visual_formats = blog_visual_formats.value or (
                                BLOG_GALLERY_VISUAL_FORMAT_OPTIONS
                                if blog_format.value == "photo/gallery slideshow"
                                else BLOG_VISUAL_FORMAT_OPTIONS
                            )
                            await set_generation_progress(
                                progress=generation_progress,
                                label=generation_progress_label,
                                status=status,
                                started_at=started_at,
                                completed_steps=1,
                                total_steps=4,
                                completed_assets=0,
                                total_assets=content_asset_total,
                                message="Generating draft and selected media assets...",
                            )
                            await generate_draft_from_fields(
                                form_fields_snapshot=generator_fields_snapshot(),
                                content_type=content_type.value,
                                topic=topic.value,
                                # BUG-QP-01: ground quote_post on the dedicated Quote book/source only,
                                # so the hidden top-level Related book/source can't contaminate retrieval.
                                related_book=(quote_book.value if content_type.value == "quote_post" else related_book.value),
                                platform=platform.value,
                                social_objectives=social_objectives.value,
                                audience=audience.value,
                                cta=cta.value,
                                constraints=constraints.value,
                                quote_book=quote_book.value,
                                quote_moods=quote_moods.value,
                                character_tags=character_tags.value,
                                podcast_format="",
                                podcast_speakers="",
                                podcast_roles="",
                                podcast_tone="",
                                podcast_length="",
                                elevenlabs_model="",
                                elevenlabs_voice_ids="",
                                podcast_show_title="",
                                podcast_episode_title="",
                                podcast_episode_number="",
                                blog_length=blog_length.value,
                                blog_format=blog_format.value,
                                blog_sections=blog_sections.value,
                                blog_structure_options=blog_structure_options.value,
                                blog_seo_keywords=blog_seo_keywords.value,
                                blog_image_mode=blog_image_mode.value,
                                blog_character=blog_character.value,
                                instagram_formats=instagram_formats.value,
                                instagram_hashtags=instagram_hashtags.value,
                                instagram_hook=instagram_hook.value,
                                instagram_image_content_type=instagram_image_content_type.value,
                                instagram_carousel_slides=instagram_carousel_slides.value,
                                cs_character=cs_character.value,
                                cs_focus=cs_focus.value,
                                cs_platform_format=cs_platform_format.value,
                                cs_image_mode=cs_image_mode.value,
                                nl_subject=nl_subject.value,
                                nl_preview=nl_preview.value,
                                nl_structure=nl_structure.value,
                                pr_timing=pr_timing.value,
                                pr_embargo_date=pr_embargo_date.value,
                                pr_city=pr_city.value,
                                pr_state=pr_state.value,
                                pr_release_date=pr_release_date.value,
                                pr_contact_name=pr_contact_name.value,
                                pr_contact_title=pr_contact_title.value,
                                pr_organization=pr_organization.value,
                                pr_contact_email=pr_contact_email.value,
                                pr_contact_phone=pr_contact_phone.value,
                                pr_website=pr_website.value,
                                pr_news_angle=pr_news_angle.value,
                                pr_release_goal="",
                                pr_primary_announcement="",
                                pr_supporting_proof=pr_supporting_proof.value,
                                pr_quote_source=pr_quote_source.value,
                                pr_target_media=pr_target_media.value,
                                pr_required_assets=pr_required_assets.value,
                                review_book=review_book.value,
                                review_source=review_source.value,
                                review_quote=review_quote.value,
                                review_mode=review_mode.value,
                                review_attribution=review_attribution.value,
                                review_promo_line=review_promo_line.value,
                                review_graphic_formats=review_visual_formats.value,
                                review_visual_style=review_visual_style.value,
                                review_quote_mode=review_quote_mode.value,
                                review_manual_quote=review_manual_quote.value,
                                review_image_type=review_image_type.value,
                                review_instagram_formats=review_instagram_formats.value,
                                review_card_style=review_card_style.value,
                                status=status,
                                filtered_context_path=filtered_context_path,
                                prompt_path=prompt_path,
                                draft_path=draft_path,
                                generated_output=generated_output,
                                visual_format_labels=(
                                    quote_visual_formats.value
                                    if content_type.value == "quote_post"
                                    else review_visual_formats_for_instagram(review_instagram_formats.value, review_visual_formats.value)
                                    if content_type.value == "review_pull_quote"
                                    else (
                                        instagram_visual_formats_for_post(instagram_image_format_selection(), instagram_visual_formats.value, instagram_carousel_slides.value)
                                        if instagram_image_format_selection() else []
                                    )
                                    if content_type.value == "instagram_caption"
                                    else formats_for_character_spotlight(cs_platform_format.value)
                                    if content_type.value == "character_spotlight"
                                    else selected_blog_visual_formats
                                    if content_type.value == "blog_post"
                                    else None
                                ),
                                visual_theme_name=(
                                    quote_visual_style.value
                                    if content_type.value == "quote_post"
                                    else review_visual_style.value
                                    if content_type.value == "review_pull_quote"
                                    else instagram_visual_style.value
                                    if content_type.value == "instagram_caption"
                                    else blog_visual_style.value
                                    if content_type.value == "blog_post"
                                    else "Gothic"
                                ),
                                saved_draft_path=saved_draft_path,
                                structured_brief_output=structured_brief_output,
                                uploaded_image_path=current_uploaded_image_path(),
                                generated_visual_path=generated_visual_path,
                                generated_visual_package_path=generated_visual_package_path,
                            )
                            record_session_draft(content_type.value, str(topic.value or ""))
                            refresh_instagram_regen_actions()
                            refresh_generated_images_gallery()
                            completed_assets = 1
                            generated_package_value = str(generated_visual_package_path.value or "").strip()
                            generated_package = Path(generated_package_value) if generated_package_value else None
                            if generated_package and generated_package.exists() and content_asset_total > 1 and content_type.value != "press_release":
                                completed_assets = content_asset_total
                            await set_generation_progress(
                                progress=generation_progress,
                                label=generation_progress_label,
                                status=status,
                                started_at=started_at,
                                completed_steps=3,
                                total_steps=4,
                                completed_assets=completed_assets,
                                total_assets=content_asset_total,
                                message="Draft ready. Finalizing assets...",
                            )
                            if content_type.value == "press_release" and pr_required_assets.value:
                                await set_generation_progress(
                                    progress=generation_progress,
                                    label=generation_progress_label,
                                    status=status,
                                    started_at=started_at,
                                    completed_steps=3,
                                    total_steps=4,
                                    completed_assets=1,
                                    total_assets=content_asset_total,
                                    message="Generating selected PR assets...",
                                )
                                assets = await asyncio.to_thread(
                                    generate_press_release_assets,
                                    selected_assets=pr_required_assets.value,
                                    topic=topic.value,
                                    draft=generated_output.value,
                                    contact={
                                        "name": pr_contact_name.value,
                                        "title": pr_contact_title.value,
                                        "organization": pr_organization.value,
                                        "email": pr_contact_email.value,
                                        "phone": pr_contact_phone.value,
                                        "website": pr_website.value,
                                    },
                                    news_angle=pr_news_angle.value,
                                )
                                if assets:
                                    pr_asset_status.value = str(assets.get("zip_path") or assets.get("error") or "")
                                    generated_visual_package_path.value = str(assets.get("zip_path") or generated_visual_package_path.value)
                                    completed_assets = content_asset_total if assets.get("paths") else 1
                                    # Show the actual companion-asset content in the editable results.
                                    rendered_assets = assets.get("assets") if isinstance(assets.get("assets"), list) else []
                                    if rendered_assets:
                                        blocks = [str(generated_output.value or "").rstrip(), "", "---", "", "# Companion Assets", ""]
                                        for item in rendered_assets:
                                            blocks.append(f"## {item.get('asset', 'Asset')}")
                                            blocks.append("")
                                            blocks.append(str(item.get("content") or "").strip())
                                            blocks.append("")
                                        generated_output.value = "\n".join(blocks).strip()
                                    ui.notify("PR assets generated.", type="positive" if assets.get("paths") else "warning")
                            if content_type.value in {"quote_post", "review_pull_quote"}:
                                refresh_quote_candidates(select_all=True)
                            refresh_platform_preview()
                            await finish_generation_progress(
                                progress=generation_progress,
                                label=generation_progress_label,
                                status=status,
                                total_assets=content_asset_total,
                                message="Draft ready.",
                            )
                        finally:
                            pass

                    async def run_comparison() -> None:
                        await generate_chatgpt_comparison(
                            content_type=content_type.value,
                            structured_brief=structured_brief_output.value,
                            mythos_content=generated_output.value,
                            model=chatgpt_model.value,
                            comparison_status=comparison_status,
                            comparison_view=comparison_view,
                            chatgpt_output=chatgpt_output,
                            chatgpt_prompt_path=chatgpt_prompt_path,
                            chatgpt_draft_path=chatgpt_draft_path,
                        )

                    def save_preference() -> None:
                        save_comparison_preference(
                            preference=comparison_preference.value,
                            notes=comparison_notes.value,
                            content_type=content_type.value,
                            structured_brief=structured_brief_output.value,
                            mythos_content=generated_output.value,
                            chatgpt_model=chatgpt_model.value,
                            chatgpt_content=chatgpt_output.value,
                            chatgpt_draft_path=chatgpt_draft_path.value,
                            chatgpt_prompt_path=chatgpt_prompt_path.value,
                            comparison_save_status=comparison_save_status,
                            comparison_vote_path=comparison_vote_path,
                        )

                    def save_current_draft() -> None:
                        if not str(generated_output.value or "").strip():
                            status.value = "Generate content before saving a draft."
                            ui.notify(status.value, type="warning")
                            return
                        title_source = (
                            first_nonempty_line(generated_output.value)
                            or first_nonempty_line(topic.value)
                            or f"{content_type.value} draft"
                        )
                        draft_record = save_draft(
                            title=title_source[:90],
                            content_type=content_type.value,
                            content=generated_output.value,
                            source_path=draft_path.value,
                            metadata={
                                "topic": str(topic.value or "").strip(),
                                "related_book": normalize_selected(related_book.value),
                                "platform": normalize_selected(platform.value),
                                "structured_brief": structured_brief_output.value,
                                "generated_visual_path": generated_visual_path.value,
                                "generated_visual_package_path": generated_visual_package_path.value,
                                "manual_save": True,
                                "form_fields": generator_fields_snapshot(),
                            },
                        )
                        saved_draft_path.value = draft_record["path"]
                        status.value = "Draft saved."
                        ui.notify("Draft saved.", type="positive")

                    def download_generated_draft() -> None:
                        download_docx(generated_output.value, content_type_label(content_type.value))

                    def download_generated_media_package() -> None:
                        package = str(generated_visual_package_path.value or "").strip()
                        if package and Path(package).exists():
                            ui.download(package)
                            return
                        preview = str(generated_visual_path.value or "").strip()
                        if preview and Path(preview).exists():
                            ui.download(preview)
                            return
                        ui.notify("No generated image or asset package is available yet.", type="warning")

                    with quote_candidate_actions:
                        make_secondary_button("Select All", select_all_quote_candidates)
                        make_secondary_button("Select None", select_no_quote_candidates)
                    with quote_image_actions:
                        make_primary_button("Generate Selected Quote Images", generate_selected_quote_images)
                    with pr_profile_actions:
                        make_secondary_button("Load Profile", apply_press_profile)
                        make_secondary_button("Save Profile", save_press_profile_from_fields)
                        make_secondary_button("Suggest News Angle", suggest_press_news_angle)
                    brief_fields = [topic, related_book, platform, social_objectives, audience, cta, constraints]
                    brief_history: list[list] = []
                    brief_restoring = {"active": False}

                    def snapshot_brief() -> list:
                        return [list(f.value) if isinstance(f.value, list) else f.value for f in brief_fields]

                    def push_brief_history() -> None:
                        if brief_restoring["active"]:
                            return
                        snap = snapshot_brief()
                        if not brief_history or brief_history[-1] != snap:
                            brief_history.append(snap)
                            if len(brief_history) > 40:
                                brief_history.pop(0)

                    def undo_brief() -> None:
                        if len(brief_history) < 2:
                            ui.notify("Nothing to undo.", type="warning")
                            return
                        brief_restoring["active"] = True
                        brief_history.pop()
                        prev = brief_history[-1]
                        for field, val in zip(brief_fields, prev):
                            field.value = list(val) if isinstance(val, list) else val
                        brief_restoring["active"] = False
                        refresh_platform_preview()
                        ui.notify("Brief restored to previous state.", type="positive")

                    for _brief_field in brief_fields:
                        _brief_field.on_value_change(lambda _event: push_brief_history())
                    push_brief_history()

                    with generator_actions:
                        make_primary_button("Generate Draft", generate_content)
                        make_secondary_button("Save Draft", save_current_draft)
                        make_secondary_button("Undo Brief Change", undo_brief)
                    with result_actions:
                        make_secondary_button("Retry / Regenerate", generate_content)
                        make_secondary_button("Download Draft", download_generated_draft)
                        make_secondary_button("Download Images / Assets", download_generated_media_package)
                        make_secondary_button("Save Draft", save_current_draft)
                    with comparison_actions:
                        make_secondary_button("Generate ChatGPT Baseline", run_comparison)
                        make_secondary_button("Save Judgment", save_preference)

            with ui.tab_panel(campaign_tab).classes("mce-panel"):
                with ui.element("section").classes("mce-grid mce-grid-single") as campaign_section:
                    with ui.card().classes("mce-card"):
                        section_heading(
                            "Campaign Mode",
                            "One brief, multiple configured assets. Select outputs first, then tune each content type.",
                        )
                        with ui.column().classes("mce-stack w-full"):
                            campaign_name = apply_field_props(
                                ui.input(label="Campaign name", placeholder="e.g. Mortal Vengeance Award Launch"),
                            )
                            campaign_topic = apply_field_props(
                                ui.textarea(
                                    label="Campaign topic",
                                    placeholder="Example: Announce Mortal Vengeance winning an award and drive readers to the book page.",
                                ),
                                "outlined autogrow",
                            )
                            with ui.element("div").classes("mce-two-col"):
                                campaign_start_date = apply_field_props(ui.input(label="Start date (optional)").props('type=date hint="YYYY-MM-DD"'))
                                campaign_cadence = apply_field_props(
                                    ui.select(
                                        ["daily", "every 2 days", "twice a week", "weekly"],
                                        value=None,
                                        label="Posting cadence (optional)",
                                    ).props("clearable"),
                                )
                                campaign_duration = apply_field_props(
                                    ui.select(CAMPAIGN_DURATION_OPTIONS, value=None, label="Campaign duration (optional)").props("clearable"),
                                )
                            campaign_book = apply_field_props(
                                ui.select(
                                    QUOTE_BOOK_OPTIONS,
                                    value=QUOTE_BOOK_OPTIONS[0] if QUOTE_BOOK_OPTIONS else None,
                                    label="Campaign related book/source (optional)",
                                ),
                            )
                            campaign_tone = apply_field_props(
                                ui.select(
                                    PODCAST_TONE_OPTIONS,
                                    value=None,
                                    label="Campaign tone",
                                ).props("use-input new-value-mode=add-unique clearable"),
                            )
                            campaign_objectives = apply_field_props(
                                ui.select(SOCIAL_OBJECTIVES, multiple=True, label="Campaign objectives"),
                                "outlined dense use-chips clearable",
                            )
                            campaign_audience = apply_field_props(
                                ui.select(AUDIENCE_OPTIONS, multiple=True, label="Campaign audience"),
                                "outlined dense use-chips clearable",
                            )
                            campaign_cta = apply_field_props(
                                ui.input(label="Campaign CTA", placeholder="e.g. 'Get the book at [link]' — applied across all generated assets"),
                            )
                            campaign_constraints = apply_field_props(
                                ui.select(CONSTRAINT_OPTIONS, multiple=True, label="Campaign constraints"),
                                "outlined dense use-chips clearable",
                            )
                            ui.label("e.g. spoiler-free, make it punchy, no hashtags").classes("mce-muted")
                            campaign_formats = apply_field_props(
                                ui.select({key: label for key, label in CAMPAIGN_FORMATS}, multiple=True, label="Campaign content types"),
                                "outlined dense use-chips clearable",
                            )
                            # Registry of the shared campaign fields, used to snapshot the
                            # form into a saved draft and restore it when re-opened.
                            campaign_shared_fields = {
                                "campaign_name": campaign_name,
                                "campaign_topic": campaign_topic,
                                "campaign_start_date": campaign_start_date,
                                "campaign_cadence": campaign_cadence,
                                "campaign_duration": campaign_duration,
                                "campaign_book": campaign_book,
                                "campaign_tone": campaign_tone,
                                "campaign_objectives": campaign_objectives,
                                "campaign_audience": campaign_audience,
                                "campaign_cta": campaign_cta,
                                "campaign_constraints": campaign_constraints,
                                "campaign_formats": campaign_formats,
                            }
                            ui.label("Pick the formats to generate, e.g. Instagram, Blog Post, Newsletter").classes("mce-muted")

                            ui.label("Selected Content Type Settings").classes("mce-section-title")
                            ui.label("Only settings for selected formats appear below.").classes("mce-muted")

                            campaign_widgets: dict[str, dict[str, object]] = {}
                            campaign_cards: dict[str, object] = {}

                            def add_base_campaign_fields(content_type: str) -> dict[str, object]:
                                return {
                                    "style": apply_field_props(
                                        ui.select(style_choices(content_type), value=default_style(content_type), label="Style / variant"),
                                    ),
                                    "quantity": apply_field_props(
                                        ui.number(
                                            label="Quantity (manual · auto-fills when cadence + duration are set)",
                                            value=1, min=1, format="%.0f",
                                        ),
                                    ),
                                }

                            with ui.card().classes("mce-subcard") as campaign_podcast_card:
                                subcard_heading("Podcast")
                                campaign_widgets["podcast"] = add_base_campaign_fields("podcast")
                                campaign_widgets["podcast"]["format"] = apply_field_props(ui.select(PODCAST_FORMAT_OPTIONS, value=PODCAST_FORMAT_OPTIONS[0], label="Podcast format"))
                                campaign_widgets["podcast"]["length"] = apply_field_props(ui.select(PODCAST_LENGTH_OPTIONS, value=PODCAST_LENGTH_OPTIONS[0], label="Podcast length"))
                                campaign_widgets["podcast"]["tone"] = apply_field_props(ui.select(PODCAST_TONE_OPTIONS, value=PODCAST_TONE_OPTIONS[0], label="Podcast delivery tone"))
                                campaign_widgets["podcast"]["speakers"] = apply_field_props(ui.number(label="Podcast speaker count", value=2, min=1, format="%.0f"))
                                campaign_widgets["podcast"]["roles"] = apply_field_props(ui.input(label="Speaker roles / names", placeholder="e.g. Host: Maya, Guest: Carlos"))
                                with ui.expansion("Podcast voice settings (audio)", icon="record_voice_over").classes("mce-expansion"):
                                    with ui.column().classes("mce-stack w-full"):
                                        campaign_widgets["podcast"]["model"] = apply_field_props(ui.select(ELEVENLABS_MODEL_OPTIONS, value=ELEVENLABS_MODEL_OPTIONS[0], label="ElevenLabs model"))
                                        campaign_podcast_voice_status = readonly_input("Voice library status", "Loading ElevenLabs voices...", mono=False)
                                        campaign_widgets["podcast"]["host_voice"] = apply_field_props(ui.select({}, label="Host voice"))
                                        campaign_widgets["podcast"]["guest_voice"] = apply_field_props(ui.select({}, label="Guest voice"))
                                        campaign_widgets["podcast"]["guest_2_voice"] = apply_field_props(ui.select({}, label="Guest 2 / co-host voice"))
                                        ui.label("Voice sample preview (click a voice above)").classes("mce-muted")
                                        campaign_podcast_voice_preview = ui.audio("", controls=True).classes("mce-audio")

                                        def preview_campaign_voice(voice_id) -> None:
                                            url = voice_preview_url(voice_id or "", campaign_podcast_voices_state)
                                            campaign_podcast_voice_preview.set_source(url or "")

                                        campaign_widgets["podcast"]["host_voice"].on_value_change(lambda e: preview_campaign_voice(e.value))
                                        campaign_widgets["podcast"]["guest_voice"].on_value_change(lambda e: preview_campaign_voice(e.value))
                                        campaign_widgets["podcast"]["guest_2_voice"].on_value_change(lambda e: preview_campaign_voice(e.value))
                                        ui.label("Stability").classes("mce-muted")
                                        campaign_widgets["podcast"]["stability"] = ui.slider(min=0, max=1, value=0.5, step=0.05).props("label-always").classes("w-full")
                                        ui.label("Similarity").classes("mce-muted")
                                        campaign_widgets["podcast"]["similarity"] = ui.slider(min=0, max=1, value=0.75, step=0.05).props("label-always").classes("w-full")
                                        ui.label("Style").classes("mce-muted")
                                        campaign_widgets["podcast"]["style_slider"] = ui.slider(min=0, max=1, value=0.0, step=0.05).props("label-always").classes("w-full")
                                        ui.label("Speed").classes("mce-muted")
                                        campaign_widgets["podcast"]["speed"] = ui.slider(min=0.7, max=1.2, value=1.0, step=0.05).props("label-always").classes("w-full")
                                        campaign_widgets["podcast"]["speaker_boost"] = ui.switch("Speaker boost", value=True)
                            campaign_cards["podcast"] = campaign_podcast_card

                            campaign_podcast_voices_state: list[dict] = []

                            async def load_campaign_podcast_voices() -> None:
                                nonlocal campaign_podcast_voices_state
                                try:
                                    voices = await asyncio.to_thread(list_voices)
                                except Exception as exc:
                                    campaign_podcast_voice_status.value = f"Could not load voices: {exc}"
                                    return
                                campaign_podcast_voices_state = voices
                                options = voice_options(voices)
                                ids = list(options.keys())
                                campaign_widgets["podcast"]["host_voice"].set_options(options, value=ids[0] if ids else None)
                                campaign_widgets["podcast"]["guest_voice"].set_options(options, value=ids[1] if len(ids) > 1 else (ids[0] if ids else None))
                                campaign_widgets["podcast"]["guest_2_voice"].set_options(options, value=ids[2] if len(ids) > 2 else (ids[0] if ids else None))
                                preview_campaign_voice(campaign_widgets["podcast"]["host_voice"].value)
                                campaign_podcast_voice_status.value = f"Loaded {len(voices)} ElevenLabs voice(s)."

                            ui.timer(0.3, load_campaign_podcast_voices, once=True)

                            with ui.card().classes("mce-subcard") as campaign_instagram_card:
                                subcard_heading("Instagram")
                                campaign_widgets["instagram_caption"] = add_base_campaign_fields("instagram_caption")
                                campaign_widgets["instagram_caption"]["formats"] = apply_field_props(ui.select(INSTAGRAM_FORMAT_OPTIONS, multiple=True, label="Instagram formats"), "outlined dense use-chips clearable")
                                campaign_widgets["instagram_caption"]["hook"] = apply_field_props(ui.input(label="Hook / Opening line", placeholder="Optional opening line / tone"))
                                campaign_widgets["instagram_caption"]["generate_images"] = apply_field_props(ui.select(INSTAGRAM_FORMAT_OPTIONS, multiple=True, label="Generate images for"), "outlined dense use-chips clearable")
                                campaign_widgets["instagram_caption"]["image_content_type"] = apply_field_props(ui.select(["text / quote graphic", "character portrait", "abstract / mood graphic", "mixed"], value="text / quote graphic", label="Image content type"))
                                campaign_widgets["instagram_caption"]["visual_style"] = apply_field_props(ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Gothic", label="Visual style"))
                                campaign_widgets["instagram_caption"]["base_image"] = base_image_picker()
                            campaign_cards["instagram_caption"] = campaign_instagram_card

                            with ui.card().classes("mce-subcard") as campaign_youtube_card:
                                subcard_heading("YouTube")
                                campaign_widgets["youtube_content"] = add_base_campaign_fields("youtube_content")
                                campaign_widgets["youtube_content"]["formats"] = apply_field_props(ui.select(YOUTUBE_CONTENT_FORMAT_OPTIONS, multiple=True, label="YouTube deliverables"), "outlined dense use-chips clearable")
                                campaign_widgets["youtube_content"]["generate_images"] = apply_field_props(ui.select(["YouTube Image Cover (16:9)", "YouTube Post (1:1)"], multiple=True, label="Generate images for"), "outlined dense use-chips clearable")
                                campaign_widgets["youtube_content"]["image_content_type"] = apply_field_props(ui.select(["text / quote graphic", "character portrait", "abstract / mood graphic", "mixed"], value="text / quote graphic", label="Image content type"))
                                campaign_widgets["youtube_content"]["visual_style"] = apply_field_props(ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Gothic", label="Visual style"))
                                campaign_widgets["youtube_content"]["base_image"] = base_image_picker()
                            campaign_cards["youtube_content"] = campaign_youtube_card

                            with ui.card().classes("mce-subcard") as campaign_linkedin_card:
                                subcard_heading("LinkedIn")
                                campaign_widgets["linkedin_content"] = add_base_campaign_fields("linkedin_content")
                                campaign_widgets["linkedin_content"]["formats"] = apply_field_props(ui.select(LINKEDIN_CONTENT_FORMAT_OPTIONS, multiple=True, label="LinkedIn deliverables"), "outlined dense use-chips clearable")
                                campaign_widgets["linkedin_content"]["generate_images"] = apply_field_props(ui.select(["LinkedIn Post Horizontal (1.91:1)", "LinkedIn Post Square (1:1)"], multiple=True, label="Generate images for"), "outlined dense use-chips clearable")
                                campaign_widgets["linkedin_content"]["image_content_type"] = apply_field_props(ui.select(["text / quote graphic", "character portrait", "abstract / mood graphic", "mixed"], value="text / quote graphic", label="Image content type"))
                                campaign_widgets["linkedin_content"]["visual_style"] = apply_field_props(ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Press", label="Visual style"))
                                campaign_widgets["linkedin_content"]["base_image"] = base_image_picker()
                            campaign_cards["linkedin_content"] = campaign_linkedin_card

                            with ui.card().classes("mce-subcard") as campaign_blog_card:
                                subcard_heading("Blog Post")
                                campaign_widgets["blog_post"] = add_base_campaign_fields("blog_post")
                                campaign_widgets["blog_post"]["format"] = apply_field_props(ui.select(BLOG_FORMAT_OPTIONS, value=BLOG_FORMAT_OPTIONS[0], label="Blog format"))
                                campaign_widgets["blog_post"]["length"] = apply_field_props(ui.select(BLOG_LENGTH_OPTIONS, value=BLOG_LENGTH_OPTIONS[1] if len(BLOG_LENGTH_OPTIONS) > 1 else BLOG_LENGTH_OPTIONS[0], label="Blog length"))
                                campaign_widgets["blog_post"]["structure"] = apply_field_props(ui.select(BLOG_STRUCTURE_CHECKLIST_OPTIONS, multiple=True, label="Structure checklist"), "outlined dense use-chips clearable")
                                campaign_widgets["blog_post"]["sections"] = apply_field_props(ui.textarea(label="Desired sections / outline", placeholder="Optional section plan for campaign blog assets"), "outlined autogrow")
                                campaign_widgets["blog_post"]["seo"] = apply_field_props(ui.input(label="SEO keywords", placeholder="Comma-separated keywords"))
                                campaign_widgets["blog_post"]["image_mode"] = apply_field_props(ui.select(["use selected character portrait", "text-only / quote graphic", "repurpose uploaded assets"], value="text-only / quote graphic", label="Blog image mode"))
                            campaign_cards["blog_post"] = campaign_blog_card

                            with ui.card().classes("mce-subcard") as campaign_newsletter_card:
                                subcard_heading("Newsletter")
                                campaign_widgets["newsletter_blurb"] = add_base_campaign_fields("newsletter_blurb")
                                campaign_widgets["newsletter_blurb"]["structure"] = apply_field_props(ui.select(NEWSLETTER_STRUCTURE_OPTIONS, value=NEWSLETTER_STRUCTURE_OPTIONS[0], label="Newsletter structure"))
                                campaign_widgets["newsletter_blurb"]["generate_images"] = apply_field_props(ui.select(["Link Preview (1.91:1)", "Square Post (1:1)"], multiple=True, label="Generate images for"), "outlined dense use-chips clearable")
                                campaign_widgets["newsletter_blurb"]["image_content_type"] = apply_field_props(ui.select(["text / quote graphic", "character portrait", "abstract / mood graphic", "mixed"], value="text / quote graphic", label="Image content type"))
                                campaign_widgets["newsletter_blurb"]["visual_style"] = apply_field_props(ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Gothic", label="Visual style"))
                                campaign_widgets["newsletter_blurb"]["base_image"] = base_image_picker()
                            campaign_cards["newsletter_blurb"] = campaign_newsletter_card

                            with ui.card().classes("mce-subcard") as campaign_quote_card:
                                subcard_heading("Quote Post")
                                campaign_widgets["quote_post"] = add_base_campaign_fields("quote_post")
                                campaign_widgets["quote_post"]["book"] = apply_field_props(ui.select(QUOTE_BOOK_OPTIONS, value=QUOTE_BOOK_OPTIONS[0] if QUOTE_BOOK_OPTIONS else None, label="Quote book/source"))
                                campaign_widgets["quote_post"]["moods"] = apply_field_props(ui.select(QUOTE_MOOD_TAGS, multiple=True, label="Quote mood/category tags"), "outlined dense use-chips clearable")
                                campaign_widgets["quote_post"]["characters"] = apply_field_props(ui.select(CHARACTER_TAGS, multiple=True, label="Character tags"), "outlined dense use-chips clearable")
                                campaign_widgets["quote_post"]["graphic_formats"] = apply_field_props(ui.select(quote_graphic_format_options(), multiple=True, label="Quote image destination formats"), "outlined dense use-chips clearable")
                                campaign_widgets["quote_post"]["graphic_theme"] = apply_field_props(ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Gothic", label="Quote image color style"))
                            campaign_cards["quote_post"] = campaign_quote_card

                            with ui.card().classes("mce-subcard") as campaign_pull_quote_card:
                                subcard_heading("Review Pull Quote")
                                campaign_widgets["review_pull_quote"] = add_base_campaign_fields("review_pull_quote")
                                ui.label("Grounding rule: use only real review excerpts from the knowledge base.").classes("mce-muted")
                                campaign_widgets["review_pull_quote"]["graphic_formats"] = apply_field_props(ui.select(quote_graphic_format_options(), multiple=True, label="Pull quote image destination formats"), "outlined dense use-chips clearable")
                                campaign_widgets["review_pull_quote"]["graphic_theme"] = apply_field_props(ui.select(QUOTE_IMAGE_STYLE_OPTIONS, value="Press", label="Pull quote image color style"))
                            campaign_cards["review_pull_quote"] = campaign_pull_quote_card

                            with ui.card().classes("mce-subcard") as campaign_character_card:
                                subcard_heading("Character Spotlight")
                                campaign_widgets["character_spotlight"] = add_base_campaign_fields("character_spotlight")
                                campaign_widgets["character_spotlight"]["character"] = apply_field_props(ui.select(CHARACTER_TAGS, label="Character subject"))
                                campaign_widgets["character_spotlight"]["format"] = apply_field_props(ui.select(CHARACTER_SPOTLIGHT_PLATFORM_FORMATS, value=CHARACTER_SPOTLIGHT_PLATFORM_FORMATS[0], label="Destination format"))
                                campaign_widgets["character_spotlight"]["image_mode"] = apply_field_props(ui.select(CHARACTER_IMAGE_MODE_OPTIONS, value=CHARACTER_IMAGE_MODE_OPTIONS[1], label="Image source"))
                            campaign_cards["character_spotlight"] = campaign_character_card

                            with ui.card().classes("mce-subcard") as campaign_press_card:
                                subcard_heading("Press Release")
                                campaign_widgets["press_release"] = add_base_campaign_fields("press_release")
                                campaign_widgets["press_release"]["timing"] = apply_field_props(ui.select(PRESS_RELEASE_TIMING_OPTIONS, value=PRESS_RELEASE_TIMING_OPTIONS[0], label="Release timing"))
                                campaign_widgets["press_release"]["news_angle"] = apply_field_props(ui.textarea(label="News angle", placeholder="Why this announcement is newsworthy"), "outlined autogrow")
                            campaign_cards["press_release"] = campaign_press_card

                            campaign_actions = ui.row().classes("mce-actions")

                    with ui.card().classes("mce-card") as campaign_results_card:
                        section_heading("Campaign Results", "Each selected content type appears below in its own section with its posts, images, and download.")
                        with ui.column().classes("mce-stack w-full"):
                            campaign_status = readonly_input("Campaign status", "Ready", mono=False)
                            campaign_progress = ui.linear_progress(value=0, show_value=False).classes("w-full")
                            campaign_progress.visible = False
                            campaign_progress_label = ui.label("Ready").classes("mce-muted")
                            campaign_preview = ui.html(build_campaign_preview_html([]), sanitize=False).classes("mce-preview")
                            ui.label("Assets by content type").classes("mce-section-title")
                            campaign_sections = ui.column().classes("w-full mce-stack")
                            campaign_output = ui.textarea(label="Generated campaign (combined, editable)", value="").props("outlined autogrow").classes("w-full mce-output-textarea")
                            campaign_path = readonly_input("Combined campaign path", "")
                            campaign_saved_draft_path = readonly_input("Saved campaign draft", "")
                            campaign_result_actions = ui.row().classes("mce-actions")

                            # Podcast audio for campaign podcast assets (uses the campaign podcast voice settings).
                            campaign_podcast_scripts: list = []
                            with ui.card().classes("mce-subcard") as campaign_podcast_audio_card:
                                subcard_heading("Podcast audio")
                                ui.label("Render ElevenLabs audio from a generated campaign podcast script using the voice settings above.").classes("mce-muted")
                                campaign_podcast_script_select = apply_field_props(ui.select({}, label="Podcast script to render"))
                                campaign_podcast_script_editor = ui.textarea(
                                    label="Podcast script (editable — insert cues / effects below)", value="",
                                ).props("outlined autogrow").classes("w-full mce-script-textarea")
                                attach_podcast_script_toolbox(campaign_podcast_script_editor, "mce-podcast-script-campaign")

                                def sync_campaign_podcast_script(_event=None) -> None:
                                    try:
                                        i = int(campaign_podcast_script_select.value or "0")
                                    except (TypeError, ValueError):
                                        i = 0
                                    if 0 <= i < len(campaign_podcast_scripts):
                                        campaign_podcast_script_editor.value = campaign_podcast_scripts[i][1]

                                campaign_podcast_script_select.on_value_change(sync_campaign_podcast_script)
                                campaign_podcast_audio_status = readonly_input("Audio status", "Generate a campaign with a Podcast asset first.", mono=False)
                                campaign_podcast_audio_progress = ui.linear_progress(value=0, show_value=False).classes("w-full")
                                campaign_podcast_audio_progress.visible = False
                                campaign_podcast_audio_progress_label = ui.label("Ready").classes("mce-muted")
                                campaign_podcast_audio_player = ui.audio("", controls=True).classes("mce-audio")
                                campaign_podcast_audio_path = readonly_input("Podcast MP3 path", "Not generated yet")
                                campaign_podcast_audio_package_path = readonly_input("Audio package path", "Not generated yet")
                                campaign_podcast_audio_actions = ui.row().classes("mce-actions")
                            campaign_podcast_audio_card.set_visibility(False)
                            with ui.expansion("Structured campaign brief", icon="subject").classes("mce-expansion"):
                                campaign_brief_output = readonly_textarea("Structured campaign brief used for comparison", "")

                            with ui.expansion("Compare campaign vs ChatGPT", icon="compare_arrows").classes("mce-expansion"):
                                with ui.column().classes("mce-stack w-full"):
                                    campaign_chatgpt_model = apply_field_props(
                                        ui.select(CHATGPT_COMPARISON_MODEL_OPTIONS, value=CHATGPT_COMPARISON_MODEL_OPTIONS[0], label="ChatGPT model")
                                    )
                                    campaign_comparison_status = readonly_input("Comparison status", "Generate a Tell Tales Ink campaign first.", mono=False)
                                    campaign_comparison_view = ui.html(comparison_side_by_side_html(), sanitize=False).classes("w-full")
                                    campaign_chatgpt_output = readonly_textarea("ChatGPT campaign baseline output", "")
                                    campaign_chatgpt_prompt_path = readonly_input("ChatGPT prompt path", "")
                                    campaign_chatgpt_draft_path = readonly_input("ChatGPT draft path", "")
                                    campaign_preference = apply_field_props(ui.select(["Tell Tales Ink", "ChatGPT", "Tie / needs revision"], label="Which campaign do you prefer?"))
                                    campaign_notes = apply_field_props(ui.textarea(label="Judge notes", placeholder="What made one campaign stronger, more specific, or less generic?"), "outlined autogrow")
                                    campaign_save_status = readonly_input("Saved judgment status", "")
                                    campaign_vote_path = readonly_input("Saved judgment path", "")
                                    campaign_comparison_actions = ui.row().classes("mce-actions")

                    def apply_campaign_visibility(selected_formats=None) -> None:
                        selected_set = set(selected_formats or [])
                        for campaign_type, card in campaign_cards.items():
                            card.visible = campaign_type in selected_set

                    apply_campaign_visibility(campaign_formats.value)
                    campaign_formats.on_value_change(lambda event: apply_campaign_visibility(event.value))
                    campaign_results_card.set_visibility(False)

                    # Posting cadence × campaign duration drives how many posts are generated
                    # per content type. Auto-fill every quantity field when either changes; the
                    # number field stays editable so the user can override it. Suppressed while a
                    # saved draft is being restored so saved quantities aren't clobbered.
                    campaign_qty_autofill = {"suppress": False}

                    def recompute_campaign_quantities(*_args) -> None:
                        if campaign_qty_autofill["suppress"]:
                            return
                        cadence = str(campaign_cadence.value or "").strip()
                        duration = str(campaign_duration.value or "").strip()
                        # Cadence and duration are optional. Only auto-calculate when BOTH are
                        # set; otherwise leave the quantity fields untouched so they stay manual.
                        if not cadence or not duration:
                            return
                        qty = campaign_post_quantity(cadence, duration)
                        for _ct_fields in campaign_widgets.values():
                            quantity_field = _ct_fields.get("quantity")
                            if quantity_field is not None:
                                quantity_field.value = qty

                    campaign_cadence.on_value_change(recompute_campaign_quantities)
                    campaign_duration.on_value_change(recompute_campaign_quantities)
                    recompute_campaign_quantities()

                    async def generate_campaign() -> None:
                        # Keep the campaign in a single column so results render full-width
                        # BELOW the form (not crammed into the narrow right sidebar).
                        campaign_results_card.set_visibility(True)
                        campaign_progress.visible = True
                        campaign_status.value = "Generating campaign..."
                        campaign_podcast_scripts.clear()
                        # Snapshot every campaign field (raw values) so the draft can be
                        # re-opened in this form later to reproduce or modify the campaign.
                        campaign_field_snapshot = {
                            "shared": {name: widget.value for name, widget in campaign_shared_fields.items()},
                            "widgets": {
                                ct: {f: w.value for f, w in fields.items()}
                                for ct, fields in campaign_widgets.items()
                            },
                        }
                        try:
                            await generate_campaign_from_fields(
                                campaign_form_fields=campaign_field_snapshot,
                                campaign_topic=campaign_topic.value,
                                campaign_name=campaign_name.value,
                                campaign_start_date=campaign_start_date.value,
                                campaign_cadence=campaign_cadence.value,
                                campaign_book=campaign_book.value,
                                campaign_tone=campaign_tone.value,
                                campaign_objectives=campaign_objectives.value,
                                campaign_audience=campaign_audience.value,
                                campaign_cta=campaign_cta.value,
                                campaign_constraints=campaign_constraints.value,
                                campaign_formats=campaign_formats.value,
                                campaign_widgets=campaign_widgets,
                                status=campaign_status,
                                campaign_output=campaign_output,
                                campaign_preview=campaign_preview,
                                campaign_sections=campaign_sections,
                                campaign_brief_output=campaign_brief_output,
                                campaign_path=campaign_path,
                                saved_draft_path=campaign_saved_draft_path,
                                campaign_progress=campaign_progress,
                                campaign_progress_label=campaign_progress_label,
                                podcast_scripts_out=campaign_podcast_scripts,
                            )
                            if campaign_podcast_scripts:
                                options = {str(i): label for i, (label, _script) in enumerate(campaign_podcast_scripts)}
                                campaign_podcast_script_select.set_options(options, value="0")
                                campaign_podcast_script_editor.value = campaign_podcast_scripts[0][1]
                                campaign_podcast_audio_status.value = f"{len(campaign_podcast_scripts)} podcast script(s) ready to render."
                                campaign_podcast_audio_card.set_visibility(True)
                            else:
                                campaign_podcast_audio_card.set_visibility(False)
                            if str(campaign_output.value or "").strip():
                                record_session_draft("campaign_mode", str(campaign_topic.value or ""))
                        finally:
                            pass

                    async def campaign_generate_podcast_audio(preview: bool) -> None:
                        if not campaign_podcast_scripts:
                            campaign_podcast_audio_status.value = "Generate a campaign with a Podcast asset first."
                            ui.notify(campaign_podcast_audio_status.value, type="warning")
                            return
                        try:
                            idx = int(campaign_podcast_script_select.value or "0")
                        except (TypeError, ValueError):
                            idx = 0
                        idx = max(0, min(idx, len(campaign_podcast_scripts) - 1))
                        # Render the edited script (with any inserted cues/effects), falling
                        # back to the originally generated script if the editor is empty.
                        script = str(campaign_podcast_script_editor.value or "").strip() or campaign_podcast_scripts[idx][1]
                        pod = campaign_widgets["podcast"]
                        campaign_podcast_audio_progress.visible = True
                        try:
                            await render_podcast_audio_native(
                                script=script,
                                document_name=(str(campaign_name.value or "").strip() or "campaign_podcast"),
                                voices=campaign_podcast_voices_state,
                                host_voice=pod["host_voice"].value,
                                guest_voice=pod["guest_voice"].value,
                                guest_2_voice=pod["guest_2_voice"].value,
                                model_id=pod["model"].value,
                                stability=pod["stability"].value or 0.5,
                                similarity_boost=pod["similarity"].value or 0.75,
                                style=pod["style_slider"].value or 0.0,
                                speed=pod["speed"].value or 1.0,
                                speaker_boost=bool(pod["speaker_boost"].value),
                                preview_only=preview,
                                status=campaign_podcast_audio_status,
                                audio_player=campaign_podcast_audio_player,
                                full_audio_path=campaign_podcast_audio_path,
                                package_path=campaign_podcast_audio_package_path,
                                progress=campaign_podcast_audio_progress,
                                progress_label=campaign_podcast_audio_progress_label,
                            )
                        except Exception as exc:
                            campaign_podcast_audio_status.value = f"Podcast audio failed: {exc}"
                            ui.notify(campaign_podcast_audio_status.value, type="negative")
                            campaign_podcast_audio_progress.visible = False

                    def send_campaign_podcast_to_studio() -> None:
                        script = str(campaign_podcast_script_editor.value or "").strip()
                        if not script:
                            ui.notify("No campaign podcast script yet.", type="warning")
                            return
                        podcast_generated_output.value = script
                        podcast_status.value = "Imported from campaign — review and render audio."
                        open_podcast()
                        ui.notify("Script sent to Podcast Studio.", type="positive")

                    with campaign_podcast_audio_actions:
                        ui.label(
                            "Preview renders a short sample; Full Podcast renders the whole script (slower, more ElevenLabs credits)."
                        ).classes("mce-muted w-full")
                        make_secondary_button("Generate Audio Preview", lambda: campaign_generate_podcast_audio(True))
                        make_secondary_button("Generate Full Podcast", lambda: campaign_generate_podcast_audio(False))
                        make_secondary_button("Send to Podcast Studio", send_campaign_podcast_to_studio)
                        make_secondary_button(
                            "Download MP3",
                            lambda: ui.download(campaign_podcast_audio_path.value) if campaign_podcast_audio_path.value else ui.notify("Generate audio first.", type="warning"),
                        )

                    async def run_campaign_comparison() -> None:
                        await generate_chatgpt_comparison(
                            content_type="campaign_mode",
                            structured_brief=campaign_brief_output.value,
                            mythos_content=campaign_output.value,
                            model=campaign_chatgpt_model.value,
                            comparison_status=campaign_comparison_status,
                            comparison_view=campaign_comparison_view,
                            chatgpt_output=campaign_chatgpt_output,
                            chatgpt_prompt_path=campaign_chatgpt_prompt_path,
                            chatgpt_draft_path=campaign_chatgpt_draft_path,
                        )

                    def save_campaign_preference() -> None:
                        save_comparison_preference(
                            preference=campaign_preference.value,
                            notes=campaign_notes.value,
                            content_type="campaign_mode",
                            structured_brief=campaign_brief_output.value,
                            mythos_content=campaign_output.value,
                            chatgpt_model=campaign_chatgpt_model.value,
                            chatgpt_content=campaign_chatgpt_output.value,
                            chatgpt_draft_path=campaign_chatgpt_draft_path.value,
                            chatgpt_prompt_path=campaign_chatgpt_prompt_path.value,
                            comparison_save_status=campaign_save_status,
                            comparison_vote_path=campaign_vote_path,
                        )

                    def download_campaign_bundle() -> None:
                        download_docx(campaign_output.value, (str(campaign_name.value or "").strip() or "campaign_bundle"))

                    with campaign_actions:
                        make_primary_button("Generate Campaign", generate_campaign)
                    with campaign_result_actions:
                        make_secondary_button("Retry / Regenerate Campaign", generate_campaign)
                        make_secondary_button("Download Campaign Bundle", download_campaign_bundle)
                    with campaign_comparison_actions:
                        make_secondary_button("Generate ChatGPT Campaign", run_campaign_comparison)
                        make_secondary_button("Save Judgment", save_campaign_preference)

            with ui.tab_panel(podcast_tab).classes("mce-panel"):
                with ui.element("section").classes("mce-grid"):
                    with ui.card().classes("mce-card"):
                        section_heading("Podcast Studio", "Create a podcast brief, generate a script, and render audio directly in NiceGUI.")
                        with ui.column().classes("mce-stack w-full"):
                            podcast_topic = apply_field_props(
                                ui.textarea(label="Podcast topic", placeholder="What should the episode cover?"),
                                "outlined autogrow",
                            )
                            podcast_platform = apply_field_props(
                                ui.select(PODCAST_DESTINATION_OPTIONS, value=PODCAST_DESTINATION_OPTIONS[0], label="Target platform"),
                            )
                            podcast_platform.tooltip("This influences tone and format, not actual upload.")
                            podcast_related_book = apply_field_props(
                                ui.select(
                                    QUOTE_BOOK_OPTIONS,
                                    value=QUOTE_BOOK_OPTIONS[0] if QUOTE_BOOK_OPTIONS else None,
                                    label="Related book/source (optional)",
                                ),
                            )
                            podcast_social_objectives = apply_field_props(
                                ui.select(SOCIAL_OBJECTIVES, multiple=True, label="Social objectives"),
                                "outlined dense use-chips clearable",
                            )
                            podcast_audience = apply_field_props(
                                ui.select(AUDIENCE_OPTIONS, multiple=True, label="Audience"),
                                "outlined dense use-chips clearable",
                            )
                            podcast_cta = apply_field_props(
                                ui.input(label="CTA / next step", placeholder="What should the listener do next?"),
                            )
                            podcast_constraints = apply_field_props(
                                ui.select(CONSTRAINT_OPTIONS, multiple=True, label="Constraints"),
                                "outlined dense use-chips clearable",
                            )

                            with ui.card().classes("mce-subcard"):
                                subcard_heading("Format and metadata")
                                podcast_format_native = apply_field_props(ui.select(PODCAST_FORMAT_OPTIONS, value=PODCAST_FORMAT_OPTIONS[0], label="Podcast format"))
                                podcast_variant_native = apply_field_props(ui.select(style_choices("podcast"), value=default_style("podcast"), label="Style / variant"))
                                podcast_speakers_native = apply_field_props(ui.number(label="Podcast speaker count", value=2, min=1, format="%.0f"))
                                podcast_roles_native = apply_field_props(ui.input(label="Speaker roles / names", placeholder="e.g. Host: Maya, Guest: Carlos"))
                                podcast_tone_native = apply_field_props(ui.select(PODCAST_TONE_OPTIONS, multiple=True, label="Podcast tone"), "outlined dense use-chips clearable")
                                podcast_length_native = apply_field_props(ui.select(PODCAST_LENGTH_OPTIONS, value=PODCAST_LENGTH_OPTIONS[0], label="Podcast target length"))
                                podcast_show_title_native = apply_field_props(ui.input(label="Podcast show title", placeholder="Series or show name"))
                                podcast_episode_title_native = apply_field_props(ui.input(label="Episode title", placeholder="Specific episode title"))
                                podcast_episode_number_native = apply_field_props(ui.input(label="Episode number", placeholder="For example: 12"))

                            with ui.expansion("Advanced voice settings", icon="record_voice_over").classes("mce-expansion"):
                                with ui.column().classes("mce-stack w-full"):
                                    podcast_model_native = apply_field_props(ui.select(ELEVENLABS_MODEL_OPTIONS, value=ELEVENLABS_MODEL_OPTIONS[0], label="ElevenLabs model"))
                                    # Human-readable casting shown to the user; the raw ID string is kept
                                    # (hidden) because the TTS backend needs the voice IDs (BUG-POD-06).
                                    podcast_voice_casting_display = readonly_input("Voice casting", "Loading voices…")
                                    podcast_voice_ids_native = apply_field_props(
                                        ui.input(label="Resolved voice casting", placeholder="Host=voice_id, Guest=voice_id"),
                                    ).props("readonly outlined dense")
                                    podcast_voice_ids_native.visible = False
                                    voice_load_status = readonly_input("Voice library status", "Loading ElevenLabs voices...", mono=False)
                                    podcast_host_voice = apply_field_props(ui.select({}, label="Host voice"))
                                    podcast_guest_voice = apply_field_props(ui.select({}, label="Guest voice"))
                                    podcast_guest_2_voice = apply_field_props(ui.select({}, label="Guest 2 / co-host voice"))
                                    ui.label("Voice sample preview (not your generated episode)").classes("mce-muted")
                                    selected_voice_preview = ui.audio("", controls=True).classes("mce-audio")
                                    ui.label("Stability").classes("mce-muted")
                                    podcast_stability = ui.slider(min=0, max=1, value=0.5, step=0.05).props("label-always").classes("w-full")
                                    ui.label("Similarity").classes("mce-muted")
                                    podcast_similarity = ui.slider(min=0, max=1, value=0.75, step=0.05).props("label-always").classes("w-full")
                                    ui.label("Style").classes("mce-muted")
                                    podcast_style = ui.slider(min=0, max=1, value=0.0, step=0.05).props("label-always").classes("w-full")
                                    ui.label("Speed").classes("mce-muted")
                                    podcast_speed = ui.slider(min=0.7, max=1.2, value=1.0, step=0.05).props("label-always").classes("w-full")
                                    podcast_speaker_boost = ui.switch("Speaker boost", value=True)

                            podcast_actions = ui.row().classes("mce-actions")

                    with ui.card().classes("mce-card mce-sticky"):
                        section_heading("Podcast Output", "Edit the generated script, preview audio, and download the finished package.")
                        with ui.column().classes("mce-stack w-full"):
                            podcast_status = readonly_input("Status", "No script generated yet.", mono=False)
                            podcast_draft_progress = ui.linear_progress(value=0, show_value=False).classes("w-full")
                            podcast_draft_progress.visible = False
                            podcast_draft_progress_label = ui.label("Ready").classes("mce-muted")
                            podcast_generated_output = ui.textarea(label="Script editor", value="").props("outlined autogrow").classes("w-full mce-script-textarea")
                            attach_podcast_script_toolbox(podcast_generated_output, "mce-podcast-script-studio")
                            podcast_script_meta = ui.label("0 words · ~0.0 min").classes("mce-muted")

                            def update_podcast_script_meta() -> None:
                                words = len(str(podcast_generated_output.value or "").split())
                                minutes = words / 150
                                podcast_script_meta.set_text(f"{words} words · ~{minutes:.1f} min estimated runtime")

                            podcast_generated_output.on_value_change(lambda _event: update_podcast_script_meta())
                            update_podcast_script_meta()
                            podcast_result_actions = ui.row().classes("mce-actions")
                            podcast_audio_status = readonly_input("Audio status / errors", "Generate a preview or full podcast when the script is ready.")
                            podcast_audio_progress = ui.linear_progress(value=0, show_value=False).classes("w-full")
                            podcast_audio_progress.visible = False
                            podcast_audio_progress_label = ui.label("Ready").classes("mce-muted")
                            podcast_audio_player = ui.audio("", controls=True).classes("mce-audio")
                            podcast_audio_path = readonly_input("Podcast MP3 path", "Not generated yet")
                            podcast_audio_package_path = readonly_input("Audio package path", "Not generated yet")
                            podcast_audio_saved_draft_path = readonly_input("Saved audio draft", "Not generated yet")
                            with ui.expansion("Draft artifacts", icon="inventory_2").classes("mce-expansion"):
                                with ui.column().classes("mce-stack w-full"):
                                    podcast_filtered_context_path = readonly_input("Filtered context", "")
                                    podcast_prompt_path = readonly_input("Generation prompt", "")
                                    podcast_draft_path = readonly_input("Draft output", "")
                                    podcast_saved_draft_path = readonly_input("Saved podcast draft", "")
                            podcast_audio_actions = ui.row().classes("mce-actions")

                    podcast_voices_state: list[dict] = []

                    def sync_voice_casting() -> None:
                        podcast_voice_ids_native.value = build_voice_casting_from_selects(
                            podcast_host_voice.value,
                            podcast_guest_voice.value,
                            podcast_guest_2_voice.value,
                        )
                        names = voice_options(podcast_voices_state)

                        def name_for(voice_id):
                            return names.get(voice_id, "—") if voice_id else "—"

                        podcast_voice_casting_display.value = (
                            f"Host: {name_for(podcast_host_voice.value)}  ·  "
                            f"Guest: {name_for(podcast_guest_voice.value)}  ·  "
                            f"Guest 2: {name_for(podcast_guest_2_voice.value)}"
                        )

                    def update_voice_preview(voice_id: str | None) -> None:
                        url = voice_preview_url(voice_id or "", podcast_voices_state)
                        if url:
                            selected_voice_preview.set_source(url)
                            return
                        selected_voice_preview.set_source("")

                    async def load_podcast_voices() -> None:
                        nonlocal podcast_voices_state
                        try:
                            voices = await asyncio.to_thread(list_voices)
                        except Exception as exc:
                            voice_load_status.value = f"Could not load voices: {exc}"
                            ui.notify(voice_load_status.value, type="warning")
                            return
                        podcast_voices_state = voices
                        options = voice_options(voices)
                        default_ids = list(options.keys())
                        podcast_host_voice.set_options(options, value=default_ids[0] if default_ids else None)
                        podcast_guest_voice.set_options(options, value=default_ids[1] if len(default_ids) > 1 else (default_ids[0] if default_ids else None))
                        podcast_guest_2_voice.set_options(options, value=default_ids[2] if len(default_ids) > 2 else (default_ids[0] if default_ids else None))
                        sync_voice_casting()
                        update_voice_preview(podcast_host_voice.value)
                        voice_load_status.value = f"Loaded {len(voices)} ElevenLabs voice(s)."

                    podcast_host_voice.on_value_change(lambda event: (sync_voice_casting(), update_voice_preview(event.value)))
                    podcast_guest_voice.on_value_change(lambda event: (sync_voice_casting(), update_voice_preview(event.value)))
                    podcast_guest_2_voice.on_value_change(lambda event: (sync_voice_casting(), update_voice_preview(event.value)))
                    ui.timer(0.25, load_podcast_voices, once=True)

                    async def generate_podcast_content() -> None:
                        started_at = datetime.now()
                        podcast_draft_progress.visible = True
                        podcast_draft_progress.value = 0
                        podcast_status.value = "Generating podcast draft..."
                        try:
                            await set_generation_progress(
                                progress=podcast_draft_progress,
                                label=podcast_draft_progress_label,
                                status=podcast_status,
                                started_at=started_at,
                                completed_steps=0,
                                total_steps=4,
                                completed_assets=0,
                                total_assets=1,
                                message="Preparing podcast brief...",
                            )
                            await set_generation_progress(
                                progress=podcast_draft_progress,
                                label=podcast_draft_progress_label,
                                status=podcast_status,
                                started_at=started_at,
                                completed_steps=1,
                                total_steps=4,
                                completed_assets=0,
                                total_assets=1,
                                message="Generating podcast script...",
                            )
                            await generate_draft_from_fields(
                                content_type="podcast",
                                topic=(f"Podcast style / variant: {podcast_variant_native.value}\n{podcast_topic.value}" if podcast_variant_native.value else podcast_topic.value),
                                related_book=podcast_related_book.value,
                                platform=podcast_platform.value,
                                social_objectives=podcast_social_objectives.value,
                                audience=podcast_audience.value,
                                cta=podcast_cta.value,
                                constraints=podcast_constraints.value,
                                quote_book="Not applicable",
                                quote_moods="Not applicable",
                                character_tags="Not applicable",
                                podcast_format=podcast_format_native.value,
                                podcast_speakers=(str(int(podcast_speakers_native.value)) if podcast_speakers_native.value else ""),
                                podcast_roles=podcast_roles_native.value,
                                podcast_tone=podcast_tone_native.value,
                                podcast_length=podcast_length_native.value,
                                elevenlabs_model=podcast_model_native.value,
                                elevenlabs_voice_ids=podcast_voice_ids_native.value,
                                podcast_show_title=podcast_show_title_native.value,
                                podcast_episode_title=podcast_episode_title_native.value,
                                podcast_episode_number=podcast_episode_number_native.value,
                                blog_length="",
                                blog_format="",
                                blog_sections="",
                                blog_structure_options="",
                                blog_seo_keywords="",
                                blog_image_mode="",
                                blog_character="",
                                instagram_formats="",
                                instagram_hashtags="",
                                instagram_hook="",
                                cs_character="",
                                cs_focus="",
                                cs_platform_format="",
                                cs_image_mode="",
                                nl_subject="",
                                nl_preview="",
                                nl_structure="",
                                pr_timing="",
                                pr_embargo_date="",
                                pr_city="",
                                pr_state="",
                                pr_release_date="",
                                pr_contact_name="",
                                pr_contact_title="",
                                pr_organization="",
                                pr_contact_email="",
                                pr_contact_phone="",
                                pr_website="",
                                pr_news_angle="",
                                pr_release_goal="",
                                pr_primary_announcement="",
                                pr_supporting_proof="",
                                pr_quote_source="",
                                pr_target_media="",
                                pr_required_assets="",
                                status=podcast_status,
                                filtered_context_path=podcast_filtered_context_path,
                                prompt_path=podcast_prompt_path,
                                draft_path=podcast_draft_path,
                                generated_output=podcast_generated_output,
                                visual_format_labels=None,
                                visual_theme_name="",
                                saved_draft_path=podcast_saved_draft_path,
                            )
                            if str(podcast_generated_output.value or "").strip():
                                record_session_draft("podcast", str(podcast_topic.value or ""))
                            await set_generation_progress(
                                progress=podcast_draft_progress,
                                label=podcast_draft_progress_label,
                                status=podcast_status,
                                started_at=started_at,
                                completed_steps=3,
                                total_steps=4,
                                completed_assets=1,
                                total_assets=1,
                                message="Saving podcast draft...",
                            )
                            sync_voice_casting()
                            await finish_generation_progress(
                                progress=podcast_draft_progress,
                                label=podcast_draft_progress_label,
                                status=podcast_status,
                                total_assets=1,
                                message="Podcast draft ready.",
                            )
                        finally:
                            pass

                    async def generate_audio_preview() -> None:
                        podcast_audio_progress.value = 0
                        podcast_audio_progress.visible = True
                        podcast_audio_status.value = "Generating audio preview..."
                        try:
                            await render_podcast_audio_native(
                                script=podcast_generated_output.value,
                                document_name=podcast_episode_title_native.value or podcast_show_title_native.value or "podcast_preview",
                                voices=podcast_voices_state,
                                host_voice=podcast_host_voice.value,
                                guest_voice=podcast_guest_voice.value,
                                guest_2_voice=podcast_guest_2_voice.value,
                                model_id=podcast_model_native.value,
                                stability=podcast_stability.value or 0.5,
                                similarity_boost=podcast_similarity.value or 0.75,
                                style=podcast_style.value or 0.0,
                                speed=podcast_speed.value or 1.0,
                                speaker_boost=bool(podcast_speaker_boost.value),
                                preview_only=True,
                                status=podcast_audio_status,
                                audio_player=podcast_audio_player,
                                full_audio_path=podcast_audio_path,
                                package_path=podcast_audio_package_path,
                                saved_draft_path=podcast_audio_saved_draft_path,
                                progress=podcast_audio_progress,
                                progress_label=podcast_audio_progress_label,
                            )
                        finally:
                            pass

                    async def generate_full_podcast_audio() -> None:
                        podcast_audio_progress.value = 0
                        podcast_audio_progress.visible = True
                        podcast_audio_status.value = "Generating full podcast audio..."
                        try:
                            await render_podcast_audio_native(
                                script=podcast_generated_output.value,
                                document_name=podcast_episode_title_native.value or podcast_show_title_native.value or "podcast",
                                voices=podcast_voices_state,
                                host_voice=podcast_host_voice.value,
                                guest_voice=podcast_guest_voice.value,
                                guest_2_voice=podcast_guest_2_voice.value,
                                model_id=podcast_model_native.value,
                                stability=podcast_stability.value or 0.5,
                                similarity_boost=podcast_similarity.value or 0.75,
                                style=podcast_style.value or 0.0,
                                speed=podcast_speed.value or 1.0,
                                speaker_boost=bool(podcast_speaker_boost.value),
                                preview_only=False,
                                status=podcast_audio_status,
                                audio_player=podcast_audio_player,
                                full_audio_path=podcast_audio_path,
                                package_path=podcast_audio_package_path,
                                saved_draft_path=podcast_audio_saved_draft_path,
                                progress=podcast_audio_progress,
                                progress_label=podcast_audio_progress_label,
                            )
                        finally:
                            pass

                    def download_podcast_script() -> None:
                        if not str(podcast_generated_output.value or "").strip():
                            ui.notify("Generate a script first.", type="warning")
                            return
                        download_docx(
                            podcast_generated_output.value,
                            (str(podcast_episode_title_native.value or "").strip() or str(podcast_show_title_native.value or "").strip() or "podcast_script"),
                        )

                    with podcast_actions:
                        make_primary_button("1. Generate Script", generate_podcast_content)
                        ui.label("Writes the script into the editor → review/edit → then use the Audio buttons below.").classes("mce-muted")
                    with podcast_result_actions:
                        make_secondary_button("Retry / Regenerate Script", generate_podcast_content)
                        make_secondary_button("Download Script", download_podcast_script)
                    with podcast_audio_actions:
                        ui.label(
                            "2. Render audio from the script above. Preview renders a short sample; "
                            "Full Podcast renders the whole script (slower, more ElevenLabs credits)."
                        ).classes("mce-muted w-full")
                        make_secondary_button("Generate Audio Preview", generate_audio_preview)
                        make_secondary_button("Generate Full Podcast", generate_full_podcast_audio)
                        make_secondary_button("Retry Audio Preview", generate_audio_preview)
                        make_secondary_button("Retry Full Audio", generate_full_podcast_audio)
                        make_secondary_button(
                            "Download MP3",
                            lambda: ui.download(podcast_audio_path.value) if podcast_audio_path.value else ui.notify("Generate audio first.", type="warning"),
                        )
                        make_secondary_button(
                            "Download Package",
                            lambda: ui.download(podcast_audio_package_path.value) if podcast_audio_package_path.value else ui.notify("Generate audio first.", type="warning"),
                        )

                    # Hide the script/audio action rows until a script actually exists.
                    _has_script = lambda v: bool(str(v or "").strip())
                    podcast_result_actions.bind_visibility_from(podcast_generated_output, "value", backward=_has_script)
                    podcast_audio_actions.bind_visibility_from(podcast_generated_output, "value", backward=_has_script)

            with ui.tab_panel(saved_drafts_tab).classes("mce-panel"):
                with ui.card().classes("mce-card"):
                    section_heading(
                        "Saved Drafts",
                        "Open, edit, and keep generated topics, campaign bundles, and podcast drafts inside the workspace.",
                    )

                    saved_draft_lookup: dict[str, str] = {}

                    def saved_draft_options() -> list[str]:
                        saved_draft_lookup.clear()
                        options: list[str] = []
                        for record in list_saved_drafts(100):
                            title = record.get("title") or "Untitled draft"
                            content_type_label = record.get("content_type") or "content"
                            updated = record.get("updated_at") or record.get("created_at") or ""
                            label = f"{title} · {content_type_label} · {updated}"
                            if label in saved_draft_lookup:
                                label = f"{label} · {str(record.get('id'))[-8:]}"
                            saved_draft_lookup[label] = str(record.get("id") or "")
                            options.append(label)
                        return options

                    with ui.column().classes("mce-stack w-full"):
                        with ui.row().classes("items-center justify-between w-full"):
                            saved_search = apply_field_props(
                                ui.input(placeholder="Search drafts by title or type...").props("clearable"),
                            )
                            saved_draft_count = ui.label("").classes("mce-muted")
                        saved_draft_list = ui.column().classes("w-full mce-saved-list")

                        def render_saved_draft_list() -> None:
                            query = str(saved_search.value or "").strip().lower()
                            records = list_saved_drafts(100)
                            filtered = [
                                r for r in records
                                if not query
                                or query in str(r.get("title") or "").lower()
                                or query in str(r.get("content_type") or "").lower()
                            ]
                            saved_draft_count.set_text(f"{len(filtered)} of {len(records)} drafts")
                            saved_draft_list.clear()
                            with saved_draft_list:
                                if not filtered:
                                    ui.label("No matching drafts." if records else "No saved drafts yet.").classes("mce-muted")
                                    return
                                with ui.row().classes("mce-saved-row mce-saved-head"):
                                    ui.label("Title").classes("mce-saved-col-title")
                                    ui.label("Type").classes("mce-saved-col-type")
                                    ui.label("Last saved").classes("mce-saved-col-date")
                                for record in filtered:
                                    rid = str(record.get("id") or "")
                                    rtitle = record.get("title") or "Untitled draft"
                                    rtype = content_type_label(record.get("content_type") or "content")
                                    rwhen = str(record.get("updated_at") or record.get("created_at") or "").replace("T", " ")
                                    row = ui.row().classes("mce-saved-row")
                                    row.on("click", lambda _e, draft_id=rid: load_saved_draft(draft_id))
                                    with row:
                                        ui.label(rtitle).classes("mce-saved-col-title")
                                        ui.label(rtype).classes("mce-saved-col-type")
                                        ui.label(rwhen).classes("mce-saved-col-date")

                        saved_search.on_value_change(lambda _e: render_saved_draft_list())

                        saved_draft_select = apply_field_props(
                            ui.select(saved_draft_options(), label="Open saved draft"),
                        )
                        with ui.element("div").classes("mce-two-col"):
                            saved_editor_title = apply_field_props(ui.input(label="Draft title"))
                            saved_editor_status = apply_field_props(
                                ui.select(["Draft", "Edited", "Approved", "Needs revision"], value="Draft", label="Revision status")
                            )
                        with ui.element("div").classes("mce-two-col"):
                            saved_editor_type = readonly_input("Content type", "")
                            saved_editor_updated = readonly_input("Last saved", "")
                        saved_editor_draft_id = readonly_input("Draft ID", "")
                        saved_editor_topic = apply_field_props(
                            ui.textarea(label="Original topic / query", placeholder="The prompt or campaign topic used to generate this draft."),
                            "outlined autogrow",
                        )
                        saved_editor_content = apply_field_props(
                            ui.textarea(label="Editable generated result", placeholder="Open a saved draft to edit it here."),
                            "outlined autogrow",
                        )
                        saved_editor_path = readonly_input("Draft file", "")
                        saved_editor_message = readonly_input("Editor status", "Select a draft to open it.")
                        with ui.expansion("Stored brief / generation context", icon="notes").classes("mce-expansion"):
                            saved_editor_brief = readonly_textarea("Brief", "")

                    def load_saved_draft(draft_id: str | None = None) -> None:
                        selected_value = draft_id or saved_draft_select.value
                        selected_id = saved_draft_lookup.get(str(selected_value), str(selected_value or ""))
                        if not selected_id:
                            saved_editor_message.value = "Select a draft to open it."
                            return
                        record = get_saved_draft(str(selected_id))
                        if not record:
                            saved_editor_message.value = "That saved draft could not be found."
                            ui.notify(saved_editor_message.value, type="warning")
                            return
                        metadata = record.get("metadata") or {}
                        saved_editor_title.value = record.get("title") or ""
                        saved_editor_status.value = record.get("status") or "Draft"
                        saved_editor_type.value = record.get("content_type") or ""
                        saved_editor_updated.value = record.get("updated_at") or record.get("created_at") or ""
                        saved_editor_draft_id.value = str(selected_id)
                        saved_editor_topic.value = str(metadata.get("topic") or metadata.get("query") or "")
                        saved_editor_brief.value = str(metadata.get("brief") or "")
                        saved_editor_content.value = read_saved_draft_content(str(selected_id))
                        saved_editor_path.value = record.get("path") or ""
                        saved_editor_message.value = "Draft opened. Edits stay here when you save revisions."

                    def open_draft_in_builder(draft_id: str | None = None) -> None:
                        """Re-open a saved draft inside the Generator (or Campaign) form with
                        its fields and content restored, so it can be regenerated or modified."""
                        selected_value = draft_id or saved_editor_draft_id.value or saved_draft_select.value
                        selected_id = saved_draft_lookup.get(str(selected_value), str(selected_value or ""))
                        record = get_saved_draft(str(selected_id)) if selected_id else None
                        if not record:
                            saved_editor_message.value = "Open a saved draft first, then reopen it in the builder."
                            ui.notify(saved_editor_message.value, type="warning")
                            return
                        metadata = record.get("metadata") or {}
                        form_fields = metadata.get("form_fields") or {}
                        content = read_saved_draft_content(str(selected_id))
                        draft_type = record.get("content_type") or ""

                        if draft_type == "chapter_promo":
                            # Chapter promos have their own tab; restore the book/chapter/platforms
                            # selection and the generated content so it can be edited or regenerated.
                            open_chapter_promos()
                            book_value = metadata.get("book")
                            if book_value:
                                try:
                                    chapter_book.value = book_value  # triggers chapter option refresh
                                except Exception:
                                    pass
                            chapter_label = metadata.get("chapter")
                            if book_value and chapter_label:
                                options = chapter_reader.chapter_options(book_value)
                                target_id = next((cid for cid, lbl in options.items() if lbl == chapter_label), None)
                                if target_id:
                                    try:
                                        chapter_select.value = target_id
                                    except Exception:
                                        pass
                            saved_platforms = metadata.get("platforms")
                            if saved_platforms:
                                try:
                                    chapter_platforms.value = saved_platforms
                                except Exception:
                                    pass
                            chapter_output.value = content
                            ui.notify("Chapter promos loaded — edit or regenerate.", type="positive")
                            return

                        if draft_type == "campaign_mode":
                            open_campaign()
                            # Suppress cadence/duration auto-recompute so the draft's saved
                            # per-type quantities are restored, not overwritten by the formula.
                            campaign_qty_autofill["suppress"] = True
                            try:
                                shared = form_fields.get("shared") or {}
                                for name, widget in campaign_shared_fields.items():
                                    if name in shared:
                                        try:
                                            widget.value = shared[name]
                                        except Exception:
                                            pass
                                # Fallbacks for drafts saved before form_fields was captured.
                                if "campaign_name" not in shared and metadata.get("campaign_name"):
                                    campaign_name.value = metadata.get("campaign_name")
                                if "campaign_topic" not in shared and metadata.get("topic"):
                                    campaign_topic.value = metadata.get("topic")
                                if "campaign_formats" not in shared and metadata.get("formats"):
                                    campaign_formats.value = metadata.get("formats")
                                for ct, fields in (form_fields.get("widgets") or {}).items():
                                    target = campaign_widgets.get(ct) or {}
                                    for field_name, value in fields.items():
                                        widget = target.get(field_name)
                                        if widget is not None:
                                            try:
                                                widget.value = value
                                            except Exception:
                                                pass
                            finally:
                                campaign_qty_autofill["suppress"] = False
                            apply_campaign_visibility(campaign_formats.value)
                            campaign_output.value = content
                            ui.notify("Campaign loaded into the builder — edit or regenerate.", type="positive")
                            return

                        if draft_type == "podcast":
                            # Podcast scripts live in the Podcast Studio tab, not the Generator.
                            open_podcast()
                            book_value = metadata.get("related_book")
                            if book_value and book_value != "Not specified":
                                try:
                                    podcast_related_book.value = book_value
                                except Exception:
                                    pass
                            topic_value = str(metadata.get("topic") or metadata.get("query") or "")
                            # The saved topic may carry a "Podcast style / variant: …" prefix line.
                            if topic_value.startswith("Podcast style / variant:"):
                                topic_value = topic_value.split("\n", 1)[1] if "\n" in topic_value else ""
                            if topic_value:
                                try:
                                    podcast_topic.value = topic_value
                                except Exception:
                                    pass
                            podcast_generated_output.value = content
                            ui.notify("Podcast script loaded into the studio — edit or regenerate.", type="positive")
                            return

                        open_generator(draft_type or None)
                        for name, widget in generator_fields.items():
                            if name in form_fields:
                                try:
                                    widget.value = form_fields[name]
                                except Exception:
                                    pass
                        # Fallbacks for drafts saved before form_fields was captured.
                        if "topic" not in form_fields:
                            topic.value = str(metadata.get("topic") or metadata.get("query") or "")
                        book_fallback = metadata.get("related_book")
                        if "related_book" not in form_fields and book_fallback and book_fallback != "Not specified":
                            try:
                                related_book.value = book_fallback
                            except Exception:
                                pass
                        generated_output.value = content
                        ui.notify("Draft loaded into the generator — edit or regenerate.", type="positive")

                    def refresh_saved_drafts() -> None:
                        current_label = saved_draft_select.value
                        current_id = saved_editor_draft_id.value
                        options = saved_draft_options()
                        next_value = current_label if current_label in options else None
                        if not next_value and current_id:
                            for label, draft_id in saved_draft_lookup.items():
                                if draft_id == current_id:
                                    next_value = label
                                    break
                        saved_draft_select.set_options(options, value=next_value)
                        if saved_draft_select.value:
                            load_saved_draft(saved_draft_select.value)
                        else:
                            saved_editor_message.value = "No saved drafts found yet."
                        render_saved_draft_list()
                        ui.notify("Saved drafts refreshed.", type="positive")

                    def save_saved_draft_revision() -> None:
                        selected_id = saved_editor_draft_id.value or saved_draft_lookup.get(str(saved_draft_select.value), "")
                        if not selected_id:
                            saved_editor_message.value = "Select a draft before saving edits."
                            ui.notify(saved_editor_message.value, type="warning")
                            return
                        updated = update_saved_draft(
                            draft_id=str(selected_id),
                            title=saved_editor_title.value,
                            content=saved_editor_content.value,
                            status=saved_editor_status.value,
                            metadata={
                                "topic": saved_editor_topic.value,
                                "edited_in_platform": True,
                            },
                        )
                        saved_editor_path.value = updated.get("path") or ""
                        saved_editor_updated.value = updated.get("updated_at") or ""
                        options = saved_draft_options()
                        selected_label = None
                        for label, draft_id in saved_draft_lookup.items():
                            if draft_id == updated.get("id"):
                                selected_label = label
                                break
                        saved_draft_select.set_options(options, value=selected_label)
                        saved_editor_draft_id.value = updated.get("id") or ""
                        saved_editor_message.value = "Revision saved."
                        render_saved_draft_list()
                        ui.notify("Draft revision saved.", type="positive")

                    def download_saved_draft() -> None:
                        download_docx(
                            saved_editor_content.value,
                            (str(saved_editor_title.value or "").strip() or "saved_draft"),
                        )

                    def clear_saved_editor() -> None:
                        for field in (
                            saved_editor_title, saved_editor_type, saved_editor_updated,
                            saved_editor_draft_id, saved_editor_topic, saved_editor_brief,
                            saved_editor_content, saved_editor_path,
                        ):
                            field.value = ""
                        saved_editor_message.value = "Select a draft to open it."

                    def archive_current_draft() -> None:
                        selected_id = saved_editor_draft_id.value or saved_draft_lookup.get(str(saved_draft_select.value), "")
                        if not selected_id:
                            ui.notify("Open a draft before archiving.", type="warning")
                            return
                        if archive_saved_draft(str(selected_id)):
                            refresh_saved_drafts()
                            ui.notify("Draft archived.", type="positive")
                        else:
                            ui.notify("That draft could not be archived.", type="warning")

                    def delete_current_draft() -> None:
                        selected_id = saved_editor_draft_id.value or saved_draft_lookup.get(str(saved_draft_select.value), "")
                        if not selected_id:
                            ui.notify("Open a draft before deleting.", type="warning")
                            return
                        with ui.dialog() as confirm_dialog, ui.card():
                            ui.label("Delete this saved draft permanently?").classes("mce-section-title")
                            ui.label(f"{saved_editor_title.value or selected_id}").classes("mce-muted")

                            def do_delete() -> None:
                                confirm_dialog.close()
                                if delete_saved_draft(str(selected_id)):
                                    clear_saved_editor()
                                    refresh_saved_drafts()
                                    ui.notify("Draft deleted.", type="positive")
                                else:
                                    ui.notify("That draft could not be deleted.", type="warning")

                            with ui.row().classes("mce-actions"):
                                make_secondary_button("Cancel", confirm_dialog.close)
                                make_primary_button("Delete", do_delete)
                        confirm_dialog.open()

                    saved_draft_select.on_value_change(lambda event: load_saved_draft(event.value))

                    with ui.row().classes("mce-actions"):
                        make_primary_button("Open in builder", open_draft_in_builder)
                        make_primary_button("Save Revision", save_saved_draft_revision)
                        make_secondary_button("Refresh Saved Drafts", refresh_saved_drafts)
                        make_secondary_button("Download Draft", download_saved_draft)
                        make_secondary_button("Archive", archive_current_draft)
                        make_secondary_button("Delete", delete_current_draft)

                    render_saved_draft_list()

        # NTH-01: keyboard shortcuts (Cmd/Ctrl+1..6 tabs, Cmd/Ctrl+Enter generate, Cmd/Ctrl+S save).
        ui.add_body_html(
            """
            <script>
            document.addEventListener('keydown', (e) => {
                const mod = e.metaKey || e.ctrlKey;
                if (mod && (e.key === 's' || (e.key >= '1' && e.key <= '6'))) {
                    e.preventDefault();
                }
            }, true);
            </script>
            """
        )

        def handle_shortcut(event) -> None:
            if not event.action.keydown:
                return
            if not (event.modifiers.meta or event.modifiers.ctrl):
                return
            key_name = str(getattr(event.key, "name", "") or event.key)
            tab_map = {
                "1": dashboard_tab,
                "2": generator_tab,
                "3": saved_drafts_tab,
                "4": campaign_tab,
                "5": podcast_tab,
                "6": chapter_promos_tab,
            }
            if key_name in tab_map:
                tabs.value = tab_map[key_name]
                refresh_dashboard()
            elif key_name == "Enter":
                current = str(tabs.value)
                if current == "Campaign Mode":
                    asyncio.create_task(generate_campaign())
                elif current == "Generator":
                    asyncio.create_task(generate_content())
            elif key_name in ("s", "S"):
                save_current_draft()

        ui.keyboard(on_key=handle_shortcut)


if __name__ in {"__main__", "__mp_main__"}:
    Path("outputs").mkdir(exist_ok=True)
    ui.run(
        host="127.0.0.1",
        port=int(os.getenv("NICEGUI_SERVER_PORT", "7860")),
        title=APP_TITLE,
        reload=False,
    )
