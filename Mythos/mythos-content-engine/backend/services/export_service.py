from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document

from backend.models import Artifact, ExportRequest
from backend.repositories.artifact_repository import artifact_for_path


PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXPORT_DIR = PROJECT_ROOT / "outputs" / "downloads"


def _slug(value: str, fallback: str = "export") -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "_", value or "").strip("_").lower()
    return cleaned[:80] or fallback


def export_content(request: ExportRequest) -> Artifact:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stem = f"{_slug(request.title)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    if request.format == "docx":
        path = EXPORT_DIR / f"{stem}.docx"
        doc = Document()
        doc.add_heading(request.title or "TellTales Ink Export", 0)
        for block in request.content.split("\n\n"):
            line = block.strip()
            if not line:
                continue
            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            else:
                doc.add_paragraph(line)
        doc.save(path)
        return artifact_for_path(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    suffix = "md" if request.format == "markdown" else "txt"
    path = EXPORT_DIR / f"{stem}.{suffix}"
    path.write_text(request.content, encoding="utf-8")
    media_type = "text/markdown" if suffix == "md" else "text/plain"
    return artifact_for_path(path, media_type)
